"""Scheda squadra e lista rinnovi, nel formato che la lega si passa in PDF.

Due documenti, ricostruiti dai dati dell'app:

* la **scheda squadra** ("SLAVIA VIAFONDA 2019"): intestazione con logo,
  presidente, allenatore, stadio, palmarès e maglie; la rosa divisa per
  reparto con squadra di Serie A, spesa e anni di contratto, colorata secondo
  la legenda; le operazioni temporanee (prestiti) e la lista dei ceduti
  temporanei;
* la **lista rinnovi** ("RINNOVO CONTRATTI 2026-2027"): i contratti scaduti di
  tutta la lega per reparto, con le colonne svincolo / rinnovo sì-no / anni
  che si riempiono alla sessione dei dadi.

Qui ci sono solo i builder puri (dati e byte dell'Excel): le viste li
avvolgono in una risposta HTTP, l'import in ``providers.team_sheet_import``
legge gli stessi file all'indietro.
"""
import io
from datetime import date

from .models import ContractEvent, Participant, Player

ROLE_ORDER = ["P", "D", "C", "A"]
ROLE_TITLES = {"P": "PORTIERI", "D": "DIFENSORI", "C": "CENTROCAMPISTI", "A": "ATTACCANTI"}

DEFAULT_HONOURS = [
    "SCUDETTI", "COPPE ITALIA", "CHAMPIONS LEAGUE", "EUROPA LEAGUE",
    "SUPERCOPPE ITALIANE", "SUPERCOPPE EUROPEE",
]

# Colori della legenda, gli stessi del foglio originale.
FILL_EXPIRING = "FFF2CC"   # in scadenza a fine anno
FILL_RENEWAL = "D9E2F3"    # rinnovo obbligatorio (contratto scaduto)
FILL_PENDING = "C5E0B3"    # contratto ancora da assegnare
FILL_OUT = "F8CBAD"        # fuori dalla Serie A
FILL_HEAD = "E7E6E6"
FILL_SECTION = "E2EFDA"
STATUS_FILLS = {"expiring": FILL_EXPIRING, "renewal": FILL_RENEWAL,
                "pending": FILL_PENDING, "out": FILL_OUT}

# Righe fisse delle due tabelle in fondo alla scheda (se ce ne sono di più,
# la tabella si allunga).
LOAN_ROWS = 6
RESERVED_ROWS = 3
# Righe vuote in coda a ogni reparto della lista rinnovi, per aggiunte a mano.
RENEWAL_SPARE_ROWS = 2

FONT = "Century Gothic"


# --- Stagione ----------------------------------------------------------------

def season_start(today=None):
    """Anno d'inizio della stagione in corso: da luglio si è nella nuova."""
    today = today or date.today()
    return today.year if today.month >= 7 else today.year - 1


def season_label(start):
    return f"{start}-{start + 1}"


def season_short(start):
    return f"{start % 100:02d}.{(start + 1) % 100:02d}"


def legend(today=None):
    start = season_start(today)
    return [
        ("expiring", "IN SCADENZA A FINE ANNO", FILL_EXPIRING),
        ("renewal", "RINNOVO OBBLIGATORIO", FILL_RENEWAL),
        ("pending", f"RINNOVO {season_short(start - 1)}", FILL_PENDING),
        ("out", "FUORI DALLA SERIE A", FILL_OUT),
    ]


# --- Dati della scheda -------------------------------------------------------

def team_label(participant):
    """Come la squadra compare negli elenchi di lega: la sigla, se c'è."""
    if participant is None:
        return ""
    return (participant.short_name or participant.display_name or "").upper()


def sheet_title(participant):
    name = (participant.display_name or "").upper()
    if participant.founded and str(participant.founded) not in name:
        name = f"{name} {participant.founded}"
    return name


def league_honour_labels(league):
    """Le voci del palmarès della lega: quelle di una squadra che le ha, o le standard."""
    if league is not None:
        for honours in Participant.objects.filter(league=league).values_list("honours", flat=True):
            labels = [str(h[0]) for h in (honours or []) if h and h[0]]
            if labels:
                return labels
    return list(DEFAULT_HONOURS)


def honours_for(participant, labels=None):
    rows = [(str(h[0]), _int(h[1])) for h in (participant.honours or []) if h and h[0]]
    if rows:
        return rows
    return [(label, 0) for label in (labels or league_honour_labels(participant.league))]


def years_text(player, contracts_on=True):
    if not contracts_on:
        return ""
    if player.contract_years is None:
        return "-"
    if player.contract_years == 0:
        return "RIN."
    return str(player.contract_years)


def player_status(player, contracts_on=True):
    """Chiave della legenda per la riga del giocatore ('' = riga bianca)."""
    if player.left_serie_a_at is not None:
        return "out"
    if not contracts_on:
        return ""
    if player.contract_years == 0:
        return "renewal"
    if player.contract_years is None:
        return "pending"
    if player.contract_years == 1:
        return "expiring"
    return ""


def contract_end(years, start=None):
    """``2`` anni rimasti → "FINE 27.28" (la stagione in cui scade)."""
    if years is None:
        return "-"
    if years <= 0:
        return "RIN."
    start = season_start() if start is None else start
    return f"FINE {season_short(start + years - 1)}"


def _sessions(n):
    if not n:
        return ""
    return f"{n} {'SESSIONE' if n == 1 else 'SESSIONI'}"


def team_sheet(participant, *, labels=None, today=None):
    """Tutto quello che va sulla scheda di una squadra, già pronto da scrivere."""
    league = participant.league
    contracts_on = bool(league and league.contracts_enabled)
    start = season_start(today)
    owned = list(Player.objects.filter(owner=participant).select_related("loan_from").order_by("name"))
    roster = [p for p in owned if not p.abroad_list]

    blocks = []
    for role in ROLE_ORDER:
        players = [p for p in roster if p.role == role]
        rows = [{
            "player": p,
            "name": (p.name or "").upper(),
            "club": ((p.left_club if p.left_serie_a_at else "") or p.team or "").upper(),
            "cost": _int(p.cost),
            "years": years_text(p, contracts_on),
            "status": player_status(p, contracts_on),
            "ext_id": p.ext_id,
        } for p in players]
        slots = league.slots_for(role) if league is not None and not league.is_mantra else 0
        blocks.append({"role": role, "title": ROLE_TITLES[role], "rows": rows,
                       "size": max(len(rows), slots or 0, 1)})

    loans = []
    for p in owned:
        if p.loan_from_id:
            loans.append({"out": "", "in": p.name.upper(), "team": team_label(p.loan_from),
                          "expiry": _sessions(p.loan_sessions_left)})
    for p in Player.objects.filter(loan_from=participant).select_related("owner").order_by("name"):
        loans.append({"out": p.name.upper(), "in": "", "team": team_label(p.owner),
                      "expiry": _sessions(p.loan_sessions_left)})

    reserved = [{
        "name": p.name.upper(),
        "club": (p.left_club or p.team or "").upper(),
        "cost": _int(p.cost),
        "expiry": contract_end(p.contract_years, start) if contracts_on else "",
    } for p in owned if p.abroad_list]

    # Righe già numerate e completate con quelle vuote, per la pagina di stampa.
    for block in blocks:
        block["lines"] = [(i + 1, block["rows"][i] if i < len(block["rows"]) else None)
                          for i in range(block["size"])]
    loan_rows = max(LOAN_ROWS, len(loans))
    reserved_rows = max(RESERVED_ROWS, len(reserved))
    marks = legend(today)
    return {
        "loan_lines": [(i + 1, loans[i] if i < len(loans) else None) for i in range(loan_rows)],
        "reserved_lines": [(i + 1, reserved[i] if i < len(reserved) else None,
                            marks[i + 1] if i + 1 < len(marks) else None) for i in range(reserved_rows)],
        "legend_first": marks[0],
        "participant": participant,
        "title": sheet_title(participant),
        "president": (participant.president_name or "").upper(),
        "coach": (participant.coach_name or "").upper(),
        "stadium": (participant.stadium or "").upper(),
        "capacity": participant.stadium_capacity,
        "honours": honours_for(participant, labels),
        "blocks": blocks,
        "loans": loans,
        "loan_rows": loan_rows,
        "reserved": reserved,
        "reserved_rows": reserved_rows,
        "legend": marks,
        "contracts_on": contracts_on,
    }


def team_sheets(league, participant_ids=None, today=None):
    qs = Participant.objects.all()
    qs = qs.filter(league=league) if league is not None else qs
    if participant_ids:
        qs = qs.filter(pk__in=participant_ids)
    labels = league_honour_labels(league)
    return [team_sheet(p, labels=labels, today=today) for p in qs.order_by("display_name")]


# --- Lista rinnovi -----------------------------------------------------------

def renewal_rows(league):
    """I contratti scaduti della stagione, per reparto, con gli esiti già noti.

    Ci sono sia quelli ancora da decidere (``contract_years == 0``) sia quelli
    già passati dalla sessione dei rinnovi di questa stagione, letti dal
    registro dei contratti: svincolati, rinnovati (con gli anni) e rescissi.
    Chi è fuori dalla Serie A o nella lista ceduti non si rinnova e non c'è.
    """
    rows = {}
    pending = (Player.objects
               .filter(owner__league=league, contract_years=0, abroad_list=False,
                       left_serie_a_at__isnull=True)
               .select_related("owner", "loan_from"))
    for p in pending:
        rows[p.pk] = _renewal_row(p, p.loan_from or p.owner)

    done = (ContractEvent.objects
            .filter(league=league, season=league.season_number, player__isnull=False,
                    kind__in=[ContractEvent.Kind.RENEWED, ContractEvent.Kind.RESCINDED,
                              ContractEvent.Kind.NOT_RENEWED])
            .select_related("player", "participant")
            .order_by("created_at"))
    for ev in done:
        row = rows.get(ev.player_id) or _renewal_row(ev.player, ev.participant, fallback=ev.participant_name)
        row["released"] = ev.kind == ContractEvent.Kind.NOT_RENEWED
        row["renewed"] = ev.kind == ContractEvent.Kind.RENEWED
        row["rescinded"] = ev.kind == ContractEvent.Kind.RESCINDED
        row["years"] = ev.years if row["renewed"] else None
        rows[ev.player_id] = row

    blocks = []
    for role in ROLE_ORDER:
        items = sorted((r for r in rows.values() if r["role"] == role),
                       key=lambda r: (r["team"], r["name"]))
        size = len(items) + RENEWAL_SPARE_ROWS
        blocks.append({"role": role, "title": ROLE_TITLES[role], "rows": items, "size": size,
                       "lines": [(i + 1, items[i] if i < len(items) else None) for i in range(size)]})
    return blocks


def _renewal_row(player, team, fallback=""):
    return {
        "player": player, "id": player.pk, "role": player.role,
        "name": (player.name or "").upper(), "club": (player.team or "").upper(),
        "team": team_label(team) or (fallback or "").upper(),
        "released": False, "renewed": False, "rescinded": False, "years": None,
    }


# --- Excel -------------------------------------------------------------------

def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="A5A5A5")
    black = Side(style="thin", color="000000")
    return {
        "font": Font(name=FONT, size=10),
        "bold": Font(name=FONT, size=10, bold=True),
        "head": Font(name=FONT, size=9, bold=True),
        "italic": Font(name=FONT, size=10, italic=True),
        "center": Alignment(horizontal="center", vertical="center"),
        "left": Alignment(horizontal="left", vertical="center", indent=1, shrink_to_fit=True),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "border_black": Border(left=black, right=black, top=black, bottom=black),
        "fill": lambda rgb: PatternFill("solid", fgColor=rgb),
    }


def _sheet_name(name, used):
    clean = "".join(c for c in name if c not in '\\/?*[]:').strip() or "Squadra"
    base, n = clean[:31], 2
    candidate = base
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _place_image(ws, field, col, row, box_w, box_h, off_x=4, off_y=4):
    """Mette un'immagine della squadra nel riquadro, in scala e senza deformarla."""
    if not field:
        return
    try:
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
        from openpyxl.drawing.xdr import XDRPositiveSize2D
        from openpyxl.utils.units import pixels_to_EMU
        from PIL import Image as PILImage

        field.open("rb")
        try:
            raw = field.read()
        finally:
            field.close()
        pil = PILImage.open(io.BytesIO(raw))
        pil.load()
        if pil.mode not in ("RGB", "RGBA"):
            pil = pil.convert("RGBA")
        buf = io.BytesIO()
        pil.save(buf, format="PNG")
        buf.seek(0)
        scale = min(box_w / pil.width, box_h / pil.height)
        w, h = max(1, int(pil.width * scale)), max(1, int(pil.height * scale))
        img = XLImage(buf)
        img.width, img.height = w, h
        img.anchor = OneCellAnchor(
            _from=AnchorMarker(col=col, colOff=pixels_to_EMU(off_x + (box_w - w) // 2),
                               row=row, rowOff=pixels_to_EMU(off_y + (box_h - h) // 2)),
            ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h)),
        )
        ws.add_image(img)
    except Exception:
        # Un'immagine illeggibile non deve far saltare l'export della rosa.
        return


# Colonne della scheda: reparto, n°, calciatore, squadra, spesa, anni | n°,
# sostituto, squadra, spesa, anni. La L (nascosta) porta l'Id ufficiale del
# giocatore, per riconoscerlo con certezza quando il file torna indietro.
_SHEET_WIDTHS = {"A": 4.3, "B": 4, "C": 22, "D": 14, "E": 8, "F": 7,
                 "G": 4.5, "H": 18, "I": 12, "J": 7.5, "K": 6.5, "L": 10}
HEADER_ROWS = 14


def write_team_sheet(ws, sheet):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Alignment, Border, Font, Side

    st = _styles()
    for col, width in _SHEET_WIDTHS.items():
        ws.column_dimensions[col].width = width
    ws.column_dimensions["L"].hidden = True
    ws.sheet_view.showGridLines = False

    # Intestazione: logo a sinistra, dati al centro, maglie a destra.
    p = sheet["participant"]
    for r in range(1, HEADER_ROWS + 1):
        ws.row_dimensions[r].height = 13.5
    lines = [(sheet["title"], Font(name=FONT, size=12, bold=True)), ("", None)]
    for label, value in (("PRESIDENTE", sheet["president"]), ("ALLENATORE", sheet["coach"])):
        lines.append((f"{label}: {value}", Font(name=FONT, size=9)))
    stadium = sheet["stadium"]
    if sheet["capacity"]:
        stadium = f"{stadium} ({sheet['capacity']})".strip()
    lines.append((f"STADIO: {stadium}", Font(name=FONT, size=9)))
    lines.append(("", None))
    for i, (text, font) in enumerate(lines, start=1):
        ws.merge_cells(start_row=i, start_column=4, end_row=i, end_column=7)
        cell = ws.cell(row=i, column=4, value=text or None)
        cell.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
        if font is not None:
            cell.font = font
    row = len(lines) + 1
    for label, count in sheet["honours"]:
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=7)
        cell = ws.cell(row=row, column=4)
        cell.value = CellRichText(
            TextBlock(InlineFont(rFont=FONT, sz=10), f"{label}: "),
            TextBlock(InlineFont(rFont=FONT, sz=10, b=True), str(count)),
        )
        cell.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
        row += 1
    header_end = max(HEADER_ROWS, row)
    sep = Side(style="thin", color="000000")
    for r in range(1, header_end + 1):
        ws.cell(row=r, column=4).border = Border(left=sep)
        ws.cell(row=r, column=8).border = Border(left=sep)
    _place_image(ws, p.logo, col=0, row=0, box_w=205, box_h=205)
    _place_image(ws, p.kit_home, col=7, row=0, box_w=150, box_h=190)
    _place_image(ws, p.kit_away, col=8, row=0, box_w=150, box_h=190, off_x=90)

    # Tabella della rosa.
    row = header_end + 2
    head = ["", "CALCIATORE", None, "SQUADRA", "SPESA", "ANNI", "SOSTITUTO", None, "SQUADRA", "SPESA", "ANNI", "ID"]
    for c, text in enumerate(head, start=1):
        cell = ws.cell(row=row, column=c, value=text)
        cell.font = st["head"]
        cell.alignment = st["center"]
        if c <= 11:
            cell.fill = st["fill"](FILL_HEAD)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
    ws.merge_cells(start_row=row, start_column=7, end_row=row, end_column=8)
    ws.row_dimensions[row].height = 15
    row += 1

    thick = Side(style="medium", color="000000")
    for block in sheet["blocks"]:
        top = row
        ws.row_dimensions[row].height = 5
        for c in range(1, 12):
            ws.cell(row=row, column=c).fill = st["fill"](FILL_HEAD)
        row += 1
        for i in range(block["size"]):
            data = block["rows"][i] if i < len(block["rows"]) else None
            ws.row_dimensions[row].height = 14
            num = ws.cell(row=row, column=2, value=i + 1)
            num.font, num.alignment = st["bold"], st["center"]
            sub = ws.cell(row=row, column=7, value=i + 1)
            sub.font, sub.alignment = st["bold"], st["center"]
            values = {}
            if data:
                values = {3: data["name"], 4: data["club"], 5: data["cost"],
                          6: int(data["years"]) if data["years"].isdigit() else data["years"]}
                if data["ext_id"]:
                    ws.cell(row=row, column=12, value=data["ext_id"])
            fill = STATUS_FILLS.get(data["status"]) if data else None
            for c in range(2, 12):
                cell = ws.cell(row=row, column=c)
                if c in values:
                    cell.value = values[c]
                if c not in (2, 7):
                    cell.font = st["font"]
                    cell.alignment = st["left"] if c in (3, 4, 8, 9) else st["center"]
                cell.border = st["border"]
                if fill and 3 <= c <= 6:
                    cell.fill = st["fill"](fill)
            ws.cell(row=row, column=7).border = Border(
                left=thick, right=st["border"].right, top=st["border"].top, bottom=st["border"].bottom)
            row += 1
        ws.merge_cells(start_row=top, start_column=1, end_row=row - 1, end_column=1)
        label = ws.cell(row=top, column=1, value=block["title"])
        label.font = Font(name=FONT, size=10, bold=True)
        label.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
        label.fill = st["fill"](FILL_HEAD)

    # Operazioni temporanee (prestiti).
    row += 1
    _section_title(ws, row, 1, 11, "OPERAZIONI TEMPORANEE", st)
    row += 1
    for (c1, c2), text in (((2, 3), "CALCIATORE CEDUTO"), ((4, 7), "CALCIATORE RICEVUTO"),
                           ((8, 9), "SOCIETA’ IMPEGNATA"), ((10, 11), "SCADENZA")):
        _merged(ws, row, c1, c2, text, st, font=st["head"])
    ws.cell(row=row, column=1).fill = st["fill"](FILL_HEAD)
    row += 1
    for i in range(sheet["loan_rows"]):
        loan = sheet["loans"][i] if i < len(sheet["loans"]) else {}
        n = ws.cell(row=row, column=1, value=i + 1)
        n.font, n.alignment, n.border = st["bold"], st["center"], st["border"]
        for (c1, c2), key in (((2, 3), "out"), ((4, 7), "in"), ((8, 9), "team"), ((10, 11), "expiry")):
            _merged(ws, row, c1, c2, loan.get(key) or None, st, align=st["left"])
        row += 1

    # Prelazione ceduti temporanei + legenda, affiancate.
    _section_title(ws, row, 1, 7, "PRELAZIONE CEDUTI TEMPORANEI", st)
    _section_title(ws, row, 8, 11, "LEGENDA", st)
    row += 1
    ws.cell(row=row, column=1).fill = st["fill"](FILL_HEAD)
    for (c1, c2), text in (((2, 3), "CALCIATORE PERSO"), ((4, 4), "SQUADRA"),
                           ((5, 5), "SPESA"), ((6, 7), "SCADENZA")):
        _merged(ws, row, c1, c2, text, st, font=st["head"])
    legend_rows = sheet["legend"]
    first_legend = row
    row += 1
    for i in range(sheet["reserved_rows"]):
        item = sheet["reserved"][i] if i < len(sheet["reserved"]) else {}
        n = ws.cell(row=row, column=1, value=i + 1)
        n.font, n.alignment, n.border = st["bold"], st["center"], st["border"]
        for (c1, c2), key in (((2, 3), "name"), ((4, 4), "club"), ((5, 5), "cost"), ((6, 7), "expiry")):
            value = item.get(key)
            _merged(ws, row, c1, c2, value if value not in ("", None) else None, st,
                    align=st["left"] if key in ("name", "club") else st["center"])
        row += 1
    for i, (_key, text, rgb) in enumerate(legend_rows):
        _merged(ws, first_legend + i, 8, 11, text, st, font=st["italic"], fill=rgb)

    # Stampa: una pagina A4 verticale, come il PDF.
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.print_area = f"A1:K{max(row - 1, first_legend + len(legend_rows) - 1)}"


def _section_title(ws, row, c1, c2, text, st):
    from openpyxl.styles import Font

    _merged(ws, row, c1, c2, text, st, font=Font(name=FONT, size=11, bold=True), fill=FILL_HEAD)
    ws.row_dimensions[row].height = 16


def _merged(ws, row, c1, c2, value, st, *, font=None, fill=None, align=None):
    if c2 > c1:
        ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    cell = ws.cell(row=row, column=c1, value=value)
    cell.font = font or st["font"]
    cell.alignment = align or st["center"]
    for c in range(c1, c2 + 1):
        ws.cell(row=row, column=c).border = st["border"]
        if fill:
            ws.cell(row=row, column=c).fill = st["fill"](fill)
    return cell


def build_team_sheets_xlsx(league, participant_ids=None, today=None):
    """Una scheda per squadra, un foglio per scheda."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    sheets = team_sheets(league, participant_ids, today=today)
    for sheet in sheets:
        ws = wb.create_sheet(_sheet_name(sheet["participant"].display_name, used))
        write_team_sheet(ws, sheet)
    if not sheets:
        ws = wb.create_sheet("Squadre")
        ws["A1"] = "Nessuna squadra in questa lega."
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


_RENEWAL_WIDTHS = {"A": 4.2, "B": 15.7, "C": 12, "D": 17.4, "E": 12.2, "F": 4.6, "G": 5.7, "H": 13.7, "I": 8}


def build_renewals_xlsx(league, today=None):
    """La lista dei contratti in scadenza, come il foglio «RINNOVO CONTRATTI»."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    st = _styles()
    border = st["border_black"]
    start = season_start(today)
    wb = Workbook()
    ws = wb.active
    ws.title = "Rinnovi"
    ws.sheet_view.showGridLines = False
    for col, width in _RENEWAL_WIDTHS.items():
        ws.column_dimensions[col].width = width
    ws.column_dimensions["I"].hidden = True

    ws.merge_cells("A1:H1")
    title = ws["A1"]
    title.value = f"RINNOVO CONTRATTI {season_label(start)}"
    title.font = Font(name=FONT, size=20, bold=True)
    title.alignment = st["center"]
    ws.row_dimensions[1].height = 42
    for c in range(1, 9):
        ws.cell(row=1, column=c).border = border

    # Intestazione delle colonne, ripetuta in cima a ogni pagina stampata.
    ws.merge_cells("A2:D3")
    ws.merge_cells("E2:E3")
    ws.merge_cells("F2:G2")
    ws.merge_cells("H2:H3")
    top = Alignment(horizontal="center", vertical="top", wrap_text=True)
    for ref, text in (("E2", "SVINCOLO"), ("F2", "RINNOVO"), ("H2", "ANNI DI\nCONTRATTO"),
                      ("F3", "SI"), ("G3", "NO"), ("I2", "ID")):
        cell = ws[ref]
        cell.value = text
        cell.font = st["font"]
        cell.alignment = top if ref in ("E2", "H2") else st["center"]
    ws["F3"].fill = st["fill"](FILL_PENDING)
    ws["G3"].fill = st["fill"](FILL_OUT)
    for r in (2, 3):
        ws.row_dimensions[r].height = 15
        for c in range(1, 9):
            ws.cell(row=r, column=c).border = border
    ws.print_title_rows = "2:3"

    row = 4
    for block in renewal_rows(league):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws.cell(row=row, column=1, value=block["title"])
        cell.font = Font(name=FONT, size=14, bold=True)
        cell.alignment = st["center"]
        for c in range(1, 9):
            ws.cell(row=row, column=c).fill = st["fill"](FILL_SECTION)
            ws.cell(row=row, column=c).border = border
        ws.row_dimensions[row].height = 20
        row += 1
        for i in range(block["size"]):
            data = block["rows"][i] if i < len(block["rows"]) else None
            values = {1: i + 1}
            if data:
                values.update({2: data["name"], 3: data["club"], 4: data["team"], 9: data["id"]})
                if data["released"]:
                    values[5] = "X"
                if data["renewed"]:
                    values[6] = "X"
                    values[8] = data["years"]
                if data["rescinded"]:
                    values[7] = "X"
            for c in range(1, 10):
                cell = ws.cell(row=row, column=c, value=values.get(c))
                cell.font = st["font"]
                cell.alignment = st["left"] if c in (2, 3, 4) else st["center"]
                if c <= 8:
                    cell.border = border
            ws.row_dimensions[row].height = 14.5
            row += 1

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.6
    ws.print_area = f"A1:H{row - 1}"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _int(value):
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0
