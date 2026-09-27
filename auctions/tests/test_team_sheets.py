"""Scheda squadra e lista rinnovi: export Excel/stampa, import da Excel e PDF."""
import io
import json
import shutil
import tempfile
from datetime import date
from decimal import Decimal
from unittest import skipUnless

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.utils import timezone

from .. import team_sheets
from ..models import ContractEvent, League, Participant, Player
from ..providers import team_sheet_import as tsi

try:
    import pdfplumber  # noqa: F401
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TODAY = date(2026, 9, 27)


def _png():
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (40, 60), (200, 20, 20)).save(buf, format="PNG")
    return buf.getvalue()


class SheetFixture(TestCase):
    def setUp(self):
        self.media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.media, True)
        override = override_settings(MEDIA_ROOT=self.media)
        override.enable()
        self.addCleanup(override.disable)

        self.league = League.objects.create(name="Lugnanese", contracts_enabled=True,
                                            slots_p=2, slots_d=3, slots_c=2, slots_a=2)
        self.team = Participant.objects.create(
            league=self.league, display_name="Slavia Viafonda", short_name="S. Viafonda", founded=2019,
            president_name="Michele", coach_name="Michele", stadium="Gabbiani Emirates",
            stadium_capacity=51300,
            honours=[["SCUDETTI", 0], ["COPPE LUGNANESI", 2], ["CHAMPIONS LEAGUE", 1]],
        )
        self.other = Participant.objects.create(league=self.league, display_name="Maccabi Defeo")

        def player(name, role, club, **kw):
            return Player.objects.create(league=self.league, name=name, role=role, team=club,
                                         initial_price=Decimal("1"), **kw)

        self.p_ok = player("Butez", "P", "Como", owner=self.team, cost=Decimal("12"), contract_years=2, ext_id="11")
        self.p_out = player("Leali", "P", "Genoa", owner=self.team, cost=Decimal("1"), contract_years=0,
                            left_serie_a_at=timezone.now(), left_club="Verona")
        self.d_exp = player("Marusic", "D", "Lazio", owner=self.team, cost=Decimal("10"), contract_years=0)
        self.d_new = player("Zanoli", "D", "Udinese", owner=self.team, cost=Decimal("1"), contract_years=None)
        self.c_one = player("Pulisic", "C", "Milan", owner=self.team, cost=Decimal("530"), contract_years=1)
        self.a_dyb = player("Dybala", "A", "Roma", owner=self.team, cost=Decimal("1070"), contract_years=0)
        self.listed = player("Malinovskyi", "C", "Genoa", owner=self.team, cost=Decimal("150"),
                             contract_years=2, abroad_list=True, left_club="Trabzonspor")
        self.loaned_in = player("Gilmour", "C", "Napoli", owner=self.team, cost=Decimal("60"),
                                contract_years=1, loan_from=self.other, loan_sessions_left=2)
        self.loaned_out = player("Frendrup", "C", "Genoa", owner=self.other, cost=Decimal("3"),
                                 contract_years=1, loan_from=self.team, loan_sessions_left=1)
        self.free = player("Svincolato Libero", "A", "Pisa")
        self.other_exp = player("Kone", "C", "Roma", owner=self.other, cost=Decimal("40"), contract_years=0)


class TeamSheetDataTests(SheetFixture):
    def test_sheet_rows_follow_the_legend(self):
        sheet = team_sheets.team_sheet(self.team, today=TODAY)
        self.assertEqual(sheet["title"], "SLAVIA VIAFONDA 2019")
        rows = {r["name"]: r for b in sheet["blocks"] for r in b["rows"]}
        self.assertEqual(rows["LEALI"]["status"], "out")
        self.assertEqual(rows["LEALI"]["club"], "VERONA")        # la squadra dove è andato
        self.assertEqual(rows["MARUSIC"]["years"], "RIN.")
        self.assertEqual(rows["MARUSIC"]["status"], "renewal")
        self.assertEqual(rows["ZANOLI"]["years"], "-")
        self.assertEqual(rows["ZANOLI"]["status"], "pending")
        self.assertEqual(rows["PULISIC"]["status"], "")
        self.assertEqual(rows["BUTEZ"]["years"], "2")
        # Il ceduto temporaneo non occupa uno slot: sta nella sua tabella.
        self.assertNotIn("MALINOVSKYI", rows)
        self.assertEqual(sheet["reserved"][0]["name"], "MALINOVSKYI")
        self.assertEqual(sheet["reserved"][0]["expiry"], "FINE 27.28")
        loans = {(l["in"], l["out"]): l for l in sheet["loans"]}
        self.assertEqual(loans[("GILMOUR", "")]["team"], "MACCABI DEFEO")
        self.assertEqual(loans[("", "FRENDRUP")]["expiry"], "1 SESSIONE")
        # Blocchi lunghi almeno quanto gli slot della lega, numerati.
        defenders = sheet["blocks"][1]
        self.assertEqual(defenders["size"], 3)
        self.assertEqual([n for n, _r in defenders["lines"]], [1, 2, 3])
        self.assertEqual(sheet["legend"][2][1], "RINNOVO 25.26")

    def test_honours_default_to_the_league_labels(self):
        honours = team_sheets.honours_for(self.other)
        self.assertEqual([h[0] for h in honours], ["SCUDETTI", "COPPE LUGNANESI", "CHAMPIONS LEAGUE"])
        self.assertTrue(all(h[1] == 0 for h in honours))

    def test_xlsx_has_one_sheet_per_team_with_colours(self):
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(team_sheets.build_team_sheets_xlsx(self.league, today=TODAY)))
        self.assertEqual(wb.sheetnames, ["Maccabi Defeo", "Slavia Viafonda"])
        ws = wb["Slavia Viafonda"]
        cells = {c.value: c for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
        self.assertIn("SLAVIA VIAFONDA 2019", cells)
        self.assertIn("PRESIDENTE: MICHELE", cells)
        self.assertIn("STADIO: GABBIANI EMIRATES (51300)", cells)
        self.assertTrue(cells["LEALI"].fill.fgColor.rgb.endswith(team_sheets.FILL_OUT))
        self.assertTrue(cells["MARUSIC"].fill.fgColor.rgb.endswith(team_sheets.FILL_RENEWAL))
        self.assertTrue(cells["ZANOLI"].fill.fgColor.rgb.endswith(team_sheets.FILL_PENDING))
        butez = cells["BUTEZ"]
        self.assertEqual(ws.cell(row=butez.row, column=12).value, "11")   # Id nascosto
        self.assertTrue(ws.column_dimensions["L"].hidden)


class RenewalListTests(SheetFixture):
    def test_expired_contracts_and_outcomes_by_role(self):
        season = self.league.season_number
        gone = Player.objects.create(league=self.league, name="Celik", role="D", team="Roma")
        renewed = Player.objects.create(league=self.league, name="Musah", role="C", team="Milan",
                                        owner=self.team, contract_years=3)
        for kind, p, years in ((ContractEvent.Kind.NOT_RENEWED, gone, None),
                               (ContractEvent.Kind.RENEWED, renewed, 3)):
            ContractEvent.objects.create(league=self.league, player=p, player_name=p.name,
                                         participant=self.team, participant_name=self.team.display_name,
                                         kind=kind, years=years, season=season)
        blocks = {b["role"]: b for b in team_sheets.renewal_rows(self.league)}
        names = {r["name"]: r for b in blocks.values() for r in b["rows"]}
        self.assertEqual(set(names), {"MARUSIC", "DYBALA", "KONE", "CELIK", "MUSAH"})
        self.assertNotIn("LEALI", names)            # fuori dalla Serie A: non si rinnova
        self.assertEqual(names["MARUSIC"]["team"], "S. VIAFONDA")
        self.assertTrue(names["CELIK"]["released"])
        self.assertTrue(names["MUSAH"]["renewed"])
        self.assertEqual(names["MUSAH"]["years"], 3)
        self.assertEqual(blocks["D"]["size"], 2 + team_sheets.RENEWAL_SPARE_ROWS)

    def test_xlsx_title_and_marks(self):
        import openpyxl

        ws = openpyxl.load_workbook(io.BytesIO(team_sheets.build_renewals_xlsx(self.league, today=TODAY))).active
        self.assertEqual(ws["A1"].value, "RINNOVO CONTRATTI 2026-2027")
        values = [c.value for row in ws.iter_rows() for c in row]
        for text in ("PORTIERI", "DIFENSORI", "CENTROCAMPISTI", "ATTACCANTI", "MARUSIC", "S. VIAFONDA"):
            self.assertIn(text, values)
        self.assertEqual(ws.print_title_rows.replace("$", ""), "2:3")


class ImportHelpersTests(TestCase):
    def test_years_and_costs(self):
        self.assertEqual(tsi.parse_years("RIN."), 0)
        self.assertIsNone(tsi.parse_years("-"))
        self.assertEqual(tsi.parse_years("3"), 3)
        self.assertEqual(tsi.parse_years(2.0), 2)
        self.assertEqual(tsi.parse_cost("1.070"), Decimal("1070"))
        self.assertIsNone(tsi.parse_cost("SCB"))
        self.assertEqual(tsi.parse_expiry_years("FINE 27.28", 2026), 2)
        self.assertEqual(tsi.parse_expiry_years("FINE 26.27", 2026), 1)

    def test_team_titles_match_names_and_initials(self):
        slavia = Participant(pk=1, display_name="Slavia Viafonda")
        maccabi = Participant(pk=2, display_name="Maccabi Defeo", short_name="Maccabi D.")
        teams = [slavia, maccabi]
        self.assertIs(tsi.match_team("SLAVIA VIAFONDA 2019", teams), slavia)
        self.assertIs(tsi.match_team("S. VIAFONDA", teams), slavia)
        self.assertIs(tsi.match_team("MACCABI D.", teams), maccabi)
        self.assertIsNone(tsi.match_team("GELSI UNITED", teams))


class XlsxRoundTripTests(SheetFixture):
    def test_export_then_import_restores_the_roster(self):
        data = team_sheets.build_team_sheets_xlsx(self.league, [self.team.pk], today=TODAY)
        sheets = tsi.parse_team_sheet_xlsx(data, "scheda.xlsx")
        self.assertEqual(len(sheets), 1)
        sheet = sheets[0]
        self.assertEqual(sheet["title"], "SLAVIA VIAFONDA 2019")
        self.assertEqual(sheet["capacity"], 51300)
        self.assertEqual(sheet["honours"][1], ["COPPE LUGNANESI", 2])
        rows = {r["name"]: r for r in sheet["players"]}
        self.assertEqual(rows["LEALI"]["status"], "out")
        self.assertEqual(rows["MARUSIC"]["years"], 0)
        self.assertIsNone(rows["ZANOLI"]["years"])
        self.assertEqual(rows["PULISIC"]["role"], "C")
        self.assertEqual(rows["BUTEZ"]["ext_id"], "11")
        self.assertEqual(sheet["reserved"][0]["name"], "MALINOVSKYI")

        # Rosa scombinata: la scheda la rimette com'era.
        Player.objects.filter(pk=self.c_one.pk).update(owner=None, cost=0, contract_years=None)
        Player.objects.filter(pk=self.d_exp.pk).update(cost=Decimal("99"), contract_years=3)
        stray = Player.objects.create(league=self.league, name="Intruso", role="A", team="Roma",
                                      owner=self.team, cost=Decimal("5"))
        Participant.objects.filter(pk=self.team.pk).update(president_name="", honours=[])

        report = tsi.apply_team_sheets(sheets, self.league, today=TODAY)
        entry = report["sheets"][0]
        self.assertEqual(entry["team"], "Slavia Viafonda")
        self.assertEqual(entry["released"], ["Intruso"])
        self.assertEqual(entry["created_players"], [])
        self.c_one.refresh_from_db()
        self.d_exp.refresh_from_db()
        stray.refresh_from_db()
        self.assertEqual(self.c_one.owner, self.team)
        self.assertEqual(self.c_one.cost, Decimal("530"))
        self.assertEqual(self.d_exp.contract_years, 0)
        self.assertEqual(self.d_exp.cost, Decimal("10"))
        self.assertIsNone(stray.owner)
        self.team.refresh_from_db()
        self.assertEqual(self.team.president_name, "MICHELE")
        self.assertEqual(self.team.honours[2], ["CHAMPIONS LEAGUE", 1])
        self.listed.refresh_from_db()
        self.assertTrue(self.listed.abroad_list)
        self.assertEqual(self.listed.contract_years, 2)
        self.loaned_in.refresh_from_db()
        self.assertEqual(self.loaned_in.loan_from, self.other)
        self.assertEqual(self.loaned_in.loan_sessions_left, 2)

    def test_preview_changes_nothing(self):
        data = team_sheets.build_team_sheets_xlsx(self.league, [self.team.pk], today=TODAY)
        sheets = tsi.parse_team_sheet_xlsx(data, "scheda.xlsx")
        Player.objects.filter(pk=self.c_one.pk).update(owner=None)
        report = tsi.apply_team_sheets(sheets, self.league, dry_run=True, today=TODAY)
        self.assertEqual(report["sheets"][0]["players"], 7)
        self.c_one.refresh_from_db()
        self.assertIsNone(self.c_one.owner)

    def test_unknown_title_creates_the_team_and_flags_departures(self):
        sheet = {"source": "x", "title": "GELSI UNITED 2001", "president": "Anna", "coach": "", "stadium": "",
                 "capacity": None, "honours": [], "loans": [], "reserved": [], "substitutes": [], "images": {},
                 "players": [
                     {"role": "A", "name": "SVINCOLATO LIBERO", "club": "PISA", "cost": Decimal("7"),
                      "cost_raw": "7", "years": 1, "years_raw": "1", "status": "", "ext_id": ""},
                     {"role": "D", "name": "ZEMURA", "club": "WATFORD", "cost": Decimal("1"),
                      "cost_raw": "1", "years": None, "years_raw": "-", "status": "out", "ext_id": ""},
                     {"role": "D", "name": "DARMIAN", "club": "SVINCOLA.", "cost": Decimal("33"),
                      "cost_raw": "33", "years": 0, "years_raw": "RIN.", "status": "out", "ext_id": ""},
                 ]}
        report = tsi.apply_team_sheets([sheet], self.league)
        entry = report["sheets"][0]
        self.assertTrue(entry["created_team"])
        team = Participant.objects.get(league=self.league, display_name="Gelsi United")
        self.assertEqual(team.founded, 2001)
        self.assertEqual(team.president_name, "Anna")
        self.free.refresh_from_db()
        self.assertEqual(self.free.owner, team)
        zemura = Player.objects.get(league=self.league, name="Zemura")
        self.assertEqual(zemura.owner, team)
        self.assertIsNotNone(zemura.left_serie_a_at)
        self.assertEqual(zemura.left_club, "Watford")
        darmian = Player.objects.get(league=self.league, name="Darmian")
        self.assertEqual(darmian.left_rank_kind, "free")
        self.assertEqual(sorted(entry["flagged_out"]), ["Darmian", "Zemura"])

    def test_mapping_can_skip_or_force_a_team(self):
        data = team_sheets.build_team_sheets_xlsx(self.league, [self.team.pk], today=TODAY)
        sheets = tsi.parse_team_sheet_xlsx(data, "scheda.xlsx")
        report = tsi.apply_team_sheets(sheets, self.league, mapping={"0": "skip"})
        self.assertTrue(report["sheets"][0]["skipped"])
        self.assertEqual(Player.objects.filter(owner=self.other).count(), 2)
        tsi.apply_team_sheets(sheets, self.league, mapping={"0": str(self.other.pk)}, replace=False)
        self.c_one.refresh_from_db()
        self.assertEqual(self.c_one.owner, self.other)


# --- PDF sintetico -----------------------------------------------------------
# Una scheda disegnata a mano con gli stessi elementi dei PDF della lega:
# testo, bordi delle celle (rettangoli sottili), righe colorate, etichette dei
# reparti scritte in verticale una lettera per volta.

class _Pdf:
    H = 842

    def __init__(self):
        self.ops = []

    def text(self, x, top, s, size=9, bold=False):
        y = self.H - top - size
        s = s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        self.ops.append(f"BT /{'F2' if bold else 'F1'} {size} Tf {x} {y} Td ({s}) Tj ET")

    def fill(self, x0, top, x1, bottom, rgb):
        r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
        self.ops.append(f"{r:.4f} {g:.4f} {b:.4f} rg {x0} {self.H - bottom} {x1 - x0} {bottom - top} re f 0 g")

    def vline(self, x, top, bottom):
        self.ops.append(f"0 g {x - 0.25} {self.H - bottom} 0.5 {bottom - top} re f")

    def build(self):
        content = "\n".join(self.ops).encode("latin-1")
        objs = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R /F2 6 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>",
        ]
        out, offsets = b"%PDF-1.4\n", []
        for i, obj in enumerate(objs, 1):
            offsets.append(len(out))
            out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
        xref = len(out)
        out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
        for off in offsets:
            out += b"%010d 00000 n \n" % off
        out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
        return out


def _synthetic_sheet_pdf():
    pdf = _Pdf()
    pdf.text(190, 25, "PROVA UNITED 2020", size=12, bold=True)
    pdf.text(190, 50, "PRESIDENTE: MARIO ROSSI")
    pdf.text(190, 61, "ALLENATORE: LUIGI BIANCHI")
    pdf.text(190, 72, "STADIO: COMUNALE (1.200)")
    pdf.text(200, 95, "SCUDETTI: 2", size=10)
    pdf.text(200, 107, "COPPE LUGNANESI: 1", size=10)

    for x, word in ((85, "CALCIATORE"), (198, "SQUADRA"), (260, "SPESA"), (300, "ANNI"),
                    (369, "SOSTITUTO"), (459, "SQUADRA"), (516, "SPESA"), (554, "ANNI")):
        pdf.text(x, 200, word, bold=True)
    grid = [39, 61, 186, 253, 293, 330, 357, 452, 509, 548]
    fills = {"out": "F8CBAD", "renewal": "D9E2F3", "pending": "C5E0B3"}
    blocks = [
        ("PORTIERI", [("PORTIERE UNO", "COMO", "10", "2", None), ("PORTIERE DUE", "VERONA", "5", "RIN.", "out")]),
        ("DIFENSORI", [("DIFENSORE A", "INTER", "30", "RIN.", "renewal"), None]),
        ("CENTROCAMPISTI", [("CENTRO C", "MILAN", "SCB", "-", "pending")]),
        ("ATTACCANTI", [("PUNTA", "ROMA", "100", "1", None)]),
    ]
    top = 215
    for label, rows in blocks:
        top += 6                                             # riga grigia che apre il reparto
        block_top = top
        for n, row in enumerate(rows, 1):
            bottom = top + 13
            if row and row[4]:
                pdf.fill(61, top, 330, bottom, fills[row[4]])
            for x in grid:
                pdf.vline(x, top, bottom)
            pdf.text(44, top + 2, str(n), bold=True)
            pdf.text(341, top + 2, str(n), bold=True)
            if row:
                name, club, cost, years, _st = row
                pdf.text(66, top + 2, name)
                pdf.text(192, top + 2, club)
                pdf.text(262, top + 2, cost)
                pdf.text(302, top + 2, years)
            top = bottom
        # Etichetta verticale: lettere dal basso verso l'alto.
        for i, ch in enumerate(label[:max(1, (top - block_top) // 7)]):
            pdf.text(22, top - 8 - i * 7, ch, size=6, bold=True)

    pdf.text(250, 600, "OPERAZIONI TEMPORANEE", size=10, bold=True)
    for x, word in ((58, "CALCIATORE CEDUTO"), (212, "CALCIATORE RICEVUTO"), (352, "SOCIETA' IMPEGNATA"),
                    (498, "SCADENZA")):
        pdf.text(x, 614, word, bold=True)
    for x in (38, 179, 193, 341, 467):
        pdf.vline(x, 612, 640)
    pdf.text(18, 628, "1", bold=True)
    pdf.text(60, 628, "PRESTATO")
    pdf.text(352, 628, "ALTRA SQUADRA")
    pdf.text(498, 628, "2 SESSIONI")

    pdf.text(117, 700, "PRELAZIONE CEDUTI TEMPORANEI", size=10, bold=True)
    pdf.text(463, 700, "LEGENDA", size=10, bold=True)
    for x, word in ((62, "CALCIATORE PERSO"), (191, "SQUADRA"), (257, "SPESA"), (317, "SCADENZA")):
        pdf.text(x, 714, word, bold=True)
    for x in (38, 179, 250, 292, 395):
        pdf.vline(x, 712, 770)
    legend = [("IN SCADENZA A FINE ANNO", "FFF2CC"), ("RINNOVO OBBLIGATORIO", "D9E2F3"),
              ("RINNOVO 25.26", "E2EFD9"), ("FUORI DALLA SERIE A", "F8CBAD")]
    for i, (text, rgb) in enumerate(legend):
        top = 712 + i * 13
        pdf.fill(396, top, 582, top + 13, rgb)
        pdf.text(420, top + 2, text)
    pdf.text(18, 727, "1", bold=True)
    pdf.text(62, 727, "EMIGRATO")
    pdf.text(191, 727, "PORTO")
    pdf.text(262, 727, "40")
    pdf.text(317, 727, "FINE 27.28")
    return pdf.build()


@skipUnless(HAS_PDF, "pdfplumber non installato")
class PdfImportTests(SheetFixture):
    def test_pdf_sheet_is_read_like_the_excel_one(self):
        sheets = tsi.parse_team_sheet_pdf(_synthetic_sheet_pdf(), "prova.pdf")
        self.assertEqual(len(sheets), 1)
        s = sheets[0]
        self.assertEqual(s["title"], "PROVA UNITED 2020")
        self.assertEqual(s["president"], "MARIO ROSSI")
        self.assertEqual(s["stadium"], "COMUNALE")
        self.assertEqual(s["capacity"], 1200)
        self.assertEqual(s["honours"], [["SCUDETTI", 2], ["COPPE LUGNANESI", 1]])
        got = [(r["role"], r["name"], r["club"], r["cost_raw"], r["years"], r["status"]) for r in s["players"]]
        self.assertEqual(got, [
            ("P", "PORTIERE UNO", "COMO", "10", 2, ""),
            ("P", "PORTIERE DUE", "VERONA", "5", 0, "out"),
            ("D", "DIFENSORE A", "INTER", "30", 0, "renewal"),
            ("C", "CENTRO C", "MILAN", "SCB", None, "pending"),
            ("A", "PUNTA", "ROMA", "100", 1, ""),
        ])
        self.assertEqual(s["loans"], [{"out": "PRESTATO", "in": "", "team": "ALTRA SQUADRA", "expiry": "2 SESSIONI"}])
        self.assertEqual(s["reserved"], [{"name": "EMIGRATO", "club": "PORTO", "cost": "40", "expiry": "FINE 27.28"}])

    def test_pdf_import_through_the_view(self):
        admin = User.objects.create_superuser("root", "r@x.local", "pw")
        client = Client()
        client.force_login(admin)
        Player.objects.create(league=self.league, name="Portiere Uno", role="P", team="Como")
        upload = SimpleUploadedFile("prova.pdf", _synthetic_sheet_pdf(), content_type="application/pdf")
        resp = client.post("/admin-auction/rose/schede/import/",
                           {"league_id": self.league.id, "sheet_files": upload, "action": "preview"})
        self.assertEqual(resp.status_code, 200, resp.content)
        data = resp.json()
        self.assertTrue(data["ok"] and data["preview"])
        entry = data["sheets"][0]
        self.assertTrue(entry["created_team"])
        self.assertEqual(entry["players"], 5)
        self.assertFalse(Participant.objects.filter(display_name="Prova United").exists())

        upload.seek(0)
        resp = client.post("/admin-auction/rose/schede/import/",
                           {"league_id": self.league.id, "sheet_files": upload,
                            "mapping": json.dumps({"0": "new"})})
        self.assertEqual(resp.status_code, 200, resp.content)
        team = Participant.objects.get(league=self.league, display_name="Prova United")
        uno = Player.objects.get(league=self.league, name="Portiere Uno")
        self.assertEqual((uno.owner, uno.cost, uno.contract_years), (team, Decimal("10"), 2))
        self.assertEqual(Player.objects.get(league=self.league, name="Portiere Due").left_club, "Verona")


class ViewTests(SheetFixture):
    def setUp(self):
        super().setUp()
        self.owner = User.objects.create_user("presidente", password="pw")
        self.league.owner = self.owner
        self.league.save()
        self.stranger = User.objects.create_user("altro", password="pw")
        League.objects.create(name="Altra", owner=self.stranger)
        self.client = Client()

    def test_downloads_are_for_the_league_admin_only(self):
        urls = ["/admin-auction/export/schede/", "/admin-auction/export/schede/stampa/",
                "/admin-auction/export/rinnovi/", "/admin-auction/export/rinnovi/stampa/"]
        self.client.force_login(self.stranger)
        for url in urls:
            self.assertEqual(self.client.get(f"{url}?league={self.league.id}").status_code, 403, url)
        self.client.force_login(self.owner)
        for url in urls:
            resp = self.client.get(f"{url}?league={self.league.id}")
            self.assertEqual(resp.status_code, 200, url)
        resp = self.client.get(f"/admin-auction/export/schede/?league={self.league.id}&team={self.team.id}")
        self.assertEqual(resp["Content-Type"], XLSX)
        self.assertIn("Slavia_Viafonda_scheda.xlsx", resp["Content-Disposition"])
        page = self.client.get(f"/admin-auction/export/schede/stampa/?league={self.league.id}").content.decode()
        self.assertIn("SLAVIA VIAFONDA 2019", page)
        self.assertIn("st-out", page)
        page = self.client.get(f"/admin-auction/export/rinnovi/stampa/?league={self.league.id}").content.decode()
        self.assertIn("RINNOVO CONTRATTI", page)
        self.assertIn("MARUSIC", page)

    def test_import_is_refused_outside_the_league(self):
        self.client.force_login(self.stranger)
        data = team_sheets.build_team_sheets_xlsx(self.league, [self.team.pk])
        upload = SimpleUploadedFile("s.xlsx", data, content_type=XLSX)
        resp = self.client.post("/admin-auction/rose/schede/import/",
                                {"league_id": self.league.id, "sheet_files": upload})
        self.assertEqual(resp.status_code, 403)

    def test_profile_form_saves_the_header_and_images(self):
        self.client.force_login(self.owner)
        resp = self.client.post(f"/admin-auction/participants/{self.other.id}/scheda/", {
            "short_name": "Maccabi D.", "president_name": "Anna", "coach_name": "Bea",
            "stadium": "Arena", "stadium_capacity": "12.000", "founded": "1999",
            "honour_label": ["Scudetti", "", "Coppe"], "honour_count": ["3", "", "x"],
            "kit_home": SimpleUploadedFile("home.png", _png(), content_type="image/png"),
        })
        self.assertEqual(resp.status_code, 302)
        self.other.refresh_from_db()
        self.assertEqual(self.other.short_name, "Maccabi D.")
        self.assertEqual(self.other.stadium_capacity, 12000)
        self.assertEqual(self.other.founded, 1999)
        self.assertEqual(self.other.honours, [["SCUDETTI", 3], ["COPPE", 0]])
        self.assertTrue(self.other.kit_home)
        # Chi non gestisce la lega non tocca la scheda.
        self.client.force_login(self.stranger)
        resp = self.client.post(f"/admin-auction/participants/{self.other.id}/scheda/", {"president_name": "X"})
        self.assertEqual(resp.status_code, 403)

    def test_team_page_offers_the_sheet_form(self):
        self.client.force_login(self.owner)
        page = self.client.get(f"/admin-auction/participants/?league={self.league.id}").content.decode()
        self.assertIn(f'id="sheet-dlg-{self.team.id}"', page)
        self.assertIn('value="Gabbiani Emirates"', page)
        self.assertIn("COPPE LUGNANESI", page)

    def test_exported_images_travel_in_the_excel(self):
        from django.core.files.base import ContentFile

        self.team.logo.save("logo.png", ContentFile(_png()))
        data = team_sheets.build_team_sheets_xlsx(self.league, [self.team.pk])
        sheet = tsi.parse_team_sheet_xlsx(data, "s.xlsx")[0]
        self.assertIn("logo", sheet["images"])
