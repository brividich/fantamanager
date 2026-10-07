import asyncio
import json
import io
import os
import re
import sqlite3
import tempfile
import threading
from datetime import timedelta
from decimal import Decimal
from pathlib import Path
from unittest import mock

from channels.routing import URLRouter
from channels.testing import WebsocketCommunicator
from django.contrib.auth.models import User
from django.conf import settings
from django.urls import reverse
from django.test import (RequestFactory, SimpleTestCase, TestCase, TransactionTestCase,
                         override_settings)
from django.utils import timezone

from .. import mantra, services
from ..providers import importers
from ..views import participant_join_url
from ..routing import websocket_urlpatterns
from ..models import (Auction, AuctionCycleResult, AuctionQueueItem, Bid, Formation,
                     League, Participant, Player, RosterLog, SealedBid, Watch)
from .common import make_live_auction

class PublicTokenTests(TestCase):
    """Fase F — strong public tokens for join + read-only TV screen."""

    def test_tokens_assigned_on_creation(self):
        auction = make_live_auction()
        p = Participant.objects.create(display_name="Tok")
        self.assertTrue(auction.public_token)
        self.assertTrue(p.public_token)
        self.assertGreaterEqual(len(p.public_token), 16)

    def test_hot_bid_path_preserves_auction_token(self):
        """save(update_fields=...) on the bid path must not regenerate the token."""
        auction = make_live_auction()
        original = auction.public_token
        auction.current_price = Decimal("999")
        auction.save(update_fields=["current_price"])
        auction.refresh_from_db()
        self.assertEqual(auction.public_token, original)

    def test_join_via_public_token(self):
        auction = make_live_auction()
        p = Participant.objects.create(display_name="Linked", is_active=True)
        resp = self.client.post("/join/", {
            "public_token": p.public_token, "auction_id": auction.id,
        })
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(self.client.session["participant_id"], p.id)

    def test_join_bad_token_rejected(self):
        resp = self.client.post("/join/", {"public_token": "nope-not-real"})
        self.assertEqual(resp.status_code, 200)  # re-render with error
        self.assertNotIn("participant_id", self.client.session)

    @override_settings(PUBLIC_TOKENS_REQUIRED=True)
    def test_anonymous_join_blocked_when_tokens_required(self):
        before = Participant.objects.count()
        resp = self.client.post("/join/", {"display_name": "Stranger"})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(Participant.objects.count(), before)  # none minted
        self.assertNotIn("participant_id", self.client.session)

    @override_settings(PUBLIC_TOKENS_REQUIRED=True)
    def test_screen_requires_token_when_enabled(self):
        auction = make_live_auction()
        forbidden = self.client.get(f"/screen/{auction.id}/")
        self.assertEqual(forbidden.status_code, 403)
        ok = self.client.get(f"/screen/{auction.id}/?t={auction.public_token}")
        self.assertEqual(ok.status_code, 200)

    @override_settings(PUBLIC_TOKENS_REQUIRED=True)
    def test_staff_bypasses_screen_token(self):
        user = User.objects.create_superuser("boss", "b@b.c", "pass12345")
        self.client.force_login(user)
        auction = make_live_auction()
        ok = self.client.get(f"/screen/{auction.id}/")  # no token, but staff
        self.assertEqual(ok.status_code, 200)


class ViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser("admin", "a@b.c", "pass12345")

    def test_admin_dashboard_requires_login(self):
        # The console is behind the login (and, through the tunnel, the PIN too).
        resp = self.client.get("/admin-auction/")
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(resp["Location"].startswith("/login/"))

    def test_admin_dashboard_loads_for_staff(self):
        self.client.force_login(self.user)
        resp = self.client.get("/admin-auction/")
        self.assertEqual(resp.status_code, 200)

    def test_classifica_partial_reflects_a_just_concluded_round(self):
        """Regression: the classifica table used to be rendered once at page
        load and never refreshed — a round's charge was invisible until a
        manual browser reload. The partial the dashboard's JS re-fetches on
        every cycle change must show the fresh numbers."""
        self.client.force_login(self.user)
        league = League.objects.create(name="L", budget=Decimal("500"))
        team = Participant.objects.create(display_name="Squadra Uno", league=league,
                                          credits=Decimal("500"))
        player = Player.objects.create(name="Kvara", role="A", league=league,
                                       owner=team, cost=Decimal("30"))
        team.spent_credits = Decimal("30")
        team.save(update_fields=["spent_credits"])
        auction = make_live_auction(league=league)

        r = self.client.get(f"/admin-auction/{auction.id}/classifica/")
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertTrue(body["ok"])
        self.assertIn("Squadra Uno", body["html"])
        self.assertIn("30", body["html"])   # spesi, freshly recomputed

    def test_storico_partial_lists_lots_with_action_buttons(self):
        """Regression: the Storico lotti panel used to render once at page
        load and only refresh via a full location.reload() — and had no way
        to act on a row (redo an unsold lot, svincola a sold one). The
        re-fetched partial must carry a per-row action, matched to outcome."""
        self.client.force_login(self.user)
        league = League.objects.create(name="L", budget=Decimal("500"))
        team = Participant.objects.create(display_name="Squadra Uno", league=league,
                                          credits=Decimal("500"))
        sold = Player.objects.create(name="Kvara", role="A", league=league,
                                     owner=team, cost=Decimal("30"))
        unsold = Player.objects.create(name="Osimhen", role="A", league=league)
        auction = make_live_auction(league=league)
        AuctionCycleResult.objects.create(
            auction=auction, cycle=1, player=sold, player_name="Kvara", player_role="A",
            assigned=True, winner=team, winner_name="Squadra Uno", amount=Decimal("30"),
        )
        AuctionCycleResult.objects.create(
            auction=auction, cycle=2, player=unsold, player_name="Osimhen", player_role="A",
            assigned=False,
        )

        r = self.client.get(f"/admin-auction/{auction.id}/storico/")
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertTrue(body["ok"])
        html = body["html"]
        self.assertIn("Kvara", html)
        self.assertIn(f"releaseFromStorico({sold.id}", html)
        self.assertIn("Osimhen", html)
        self.assertIn(f"retryFromStorico({unsold.id}", html)

    def test_participants_page_lists_teams_with_join_links(self):
        self.client.force_login(self.user)
        league = League.objects.create(name="L", budget=Decimal("500"))
        p = Participant.objects.create(display_name="Squadra Uno", league=league,
                                       credits=Decimal("500"))
        Player.objects.create(name="Kva", role="A", league=league, owner=p,
                              cost=Decimal("30"), initial_price=Decimal("10"))
        resp = self.client.get("/admin-auction/participants/")
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Squadra Uno")
        self.assertContains(resp, p.public_token)   # personal join link is shown

    def test_dashboard_renders_for_each_flow_mode(self):
        self.client.force_login(self.user)
        Player.objects.create(name="Free", role="D", initial_price=Decimal("5"), owner=None)
        for flow, order in (
            (Auction.FlowMode.CALL, Auction.CallOrder.PDCA),
            (Auction.FlowMode.CONTINUOUS, Auction.CallOrder.ALPHA),
            (Auction.FlowMode.MANUAL, Auction.CallOrder.ACDP),
        ):
            a = make_live_auction(flow_mode=flow, call_order=order,
                                  status=Auction.Status.READY)
            services.build_queue(a)
            resp = self.client.get(f"/admin-auction/?auction={a.id}")
            self.assertEqual(resp.status_code, 200, f"flow={flow}")

    def test_dashboard_hides_the_queue_in_random_order(self):
        """A random draw that lists who is next is not a random draw any more."""
        self.client.force_login(self.user)
        Player.objects.create(name="Zibidibo", role="D", initial_price=Decimal("5"), owner=None)

        a = make_live_auction(flow_mode=Auction.FlowMode.MANUAL,
                              call_order=Auction.CallOrder.RANDOM,
                              status=Auction.Status.READY)
        services.build_queue(a)
        resp = self.client.get(f"/admin-auction/?auction={a.id}")
        self.assertNotContains(resp, "Zibidibo")
        self.assertContains(resp, "Ordine casuale")
        # The live state goes to every screen and phone too, so it must not
        # carry the next name either — only the count.
        state = services.serialize_state(a)
        self.assertIsNone(state["queue_next"])
        self.assertEqual(state["queue_pending"], 1)

        # Any other order still shows the running order.
        b = make_live_auction(flow_mode=Auction.FlowMode.MANUAL,
                              call_order=Auction.CallOrder.ALPHA,
                              status=Auction.Status.READY)
        services.build_queue(b)
        resp = self.client.get(f"/admin-auction/?auction={b.id}")
        self.assertContains(resp, "Zibidibo")

    def test_screen_sizes_saved_from_settings_and_published_in_state(self):
        """The maxischermo reads both sizes off the live state, so no reload."""
        self.client.force_login(self.user)
        a = make_live_auction()
        self.assertEqual(a.screen_timer_size, Auction.ScreenSize.MEDIUM)

        r = self.client.post(f"/admin-auction/{a.id}/edit/", {
            "title": a.title, "screen_timer_size": "l", "screen_name_size": "s",
        })
        self.assertEqual(r.status_code, 200)
        a.refresh_from_db()
        self.assertEqual(a.screen_timer_size, "l")
        self.assertEqual(a.screen_name_size, "s")
        state = r.json()["state"]
        self.assertEqual(state["screen_timer_size"], "l")
        self.assertEqual(state["screen_name_size"], "s")

        # A bogus value leaves the saved size alone rather than breaking the CSS class.
        self.client.post(f"/admin-auction/{a.id}/edit/",
                         {"title": a.title, "screen_timer_size": "xxl"})
        a.refresh_from_db()
        self.assertEqual(a.screen_timer_size, "l")

    def test_wizard_page_renders(self):
        self.client.force_login(self.user)
        resp = self.client.get("/admin-auction/wizard/")
        self.assertEqual(resp.status_code, 200)

    def test_join_creates_participant_session(self):
        auction = make_live_auction()
        resp = self.client.post(
            "/join/", {"display_name": "Dora", "auction_id": auction.id}
        )
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(Participant.objects.filter(display_name="Dora").exists())
        self.assertIn("participant_id", self.client.session)

    def test_bid_page_redirects_without_session(self):
        auction = make_live_auction()
        resp = self.client.get(f"/bid/{auction.id}/")
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/join/", resp["Location"])

    def test_screen_page_loads(self):
        auction = make_live_auction()
        resp = self.client.get(f"/screen/{auction.id}/")
        self.assertEqual(resp.status_code, 200)


class CreateLeagueFlowTests(TestCase):
    """Fase J — 'Nuova lega': create the season-long container (name + teams +
    platform). Distinct from creating an auction, which runs *on* a league."""

    def setUp(self):
        self.user = User.objects.create_superuser("admin", "a@b.c", "pass12345")
        self.client.force_login(self.user)

    def test_league_form_loads(self):
        resp = self.client.get("/admin-auction/league/new/")
        self.assertEqual(resp.status_code, 200)

    def test_create_league_builds_league_and_teams(self):
        orphan = Participant.objects.create(display_name="Senza Lega")  # league=None
        resp = self.client.post("/admin-auction/league/new/", {
            "name": "Lega Test", "budget": "300",
            "source_site": "fantapazz", "external_id": "332175",
            "slots_p": "3", "slots_d": "8", "slots_c": "8", "slots_a": "6",
            "participants_json": json.dumps(
                [{"name": "Alfa", "credits": "250"}, {"name": "Beta", "credits": ""}]
            ),
            "attach_ids": [str(orphan.id)],
        })
        self.assertEqual(resp.status_code, 302)

        league = League.objects.get(name="Lega Test")
        self.assertEqual(league.budget, Decimal("300"))
        self.assertEqual(league.source_site, "fantapazz")
        self.assertEqual(league.external_id, "332175")

        alfa = Participant.objects.get(display_name="Alfa")
        self.assertEqual(alfa.credits, Decimal("250"))
        self.assertEqual(alfa.league_id, league.id)
        beta = Participant.objects.get(display_name="Beta")
        self.assertEqual(beta.credits, Decimal("300"))   # blank → league budget

        orphan.refresh_from_db()
        self.assertEqual(orphan.league_id, league.id)     # folded in via attach

        # No auction is created by the league flow.
        self.assertEqual(Auction.objects.filter(league=league).count(), 0)

    def test_create_league_remembers_config(self):
        from ..models import LeagueConfig
        self.client.post("/admin-auction/league/new/", {
            "name": "Lega Memoria", "budget": "420",
            "slots_p": "2", "slots_d": "7", "slots_c": "7", "slots_a": "5",
            "participants_json": "[]",
        })
        cfg = LeagueConfig.get()
        self.assertEqual(cfg.name, "Lega Memoria")
        self.assertEqual(cfg.budget, Decimal("420"))
        self.assertEqual((cfg.slots_p, cfg.slots_d, cfg.slots_c, cfg.slots_a), (2, 7, 7, 5))
        # The form prefills those values next time.
        resp = self.client.get("/admin-auction/league/new/")
        self.assertContains(resp, "Lega Memoria")


class WizardTests(TestCase):
    """Fase J — 'Nuova asta': the wizard creates an auction ON an existing
    league. It never creates a league or teams (that is 'Nuova lega')."""

    def setUp(self):
        self.user = User.objects.create_superuser("admin", "a@b.c", "pass12345")
        self.client.force_login(self.user)
        self.league = League.objects.create(
            name="Lega Base", budget=Decimal("300"),
            source_site="fantapazz", external_id="332175",
        )
        # An auction requires its league's listone to be already imported.
        Player.objects.create(name="Base", role="A", league=self.league,
                              initial_price=Decimal("1"), owner=None)

    def test_wizard_page_loads(self):
        resp = self.client.get("/admin-auction/wizard/")
        self.assertEqual(resp.status_code, 200)

    def test_wizard_creates_auction_on_existing_league(self):
        league_count = League.objects.count()
        resp = self.client.post("/admin-auction/wizard/create/", {
            "league_id": str(self.league.id),
            "mode": "REPAIR_AUCTION",
            "flow_mode": "continuous", "call_order": "alpha",
            "min_increment": "2", "quick_increments": "1,2,5",
            "duration_seconds": "45", "antisnipe_seconds": "5",
        })
        self.assertEqual(resp.status_code, 302)
        # No new league/teams were created.
        self.assertEqual(League.objects.count(), league_count)

        auction = Auction.objects.latest("id")
        self.assertEqual(auction.league_id, self.league.id)
        self.assertEqual(auction.mode, Auction.Mode.REPAIR_AUCTION)
        self.assertEqual(auction.flow_mode, Auction.FlowMode.CONTINUOUS)
        self.assertEqual(auction.status, Auction.Status.READY)
        self.assertEqual(auction.duration_seconds, 45)
        # Provenance is inherited from the league.
        self.assertEqual(auction.source_site, "fantapazz")

    def test_wizard_without_league_redirects_to_create_league(self):
        resp = self.client.post("/admin-auction/wizard/create/", {"mode": "NEW_FROM_ZERO"})
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/league/new/", resp["Location"])
        self.assertFalse(Auction.objects.exists())

    def test_wizard_create_start_now_goes_live(self):
        # The listone (player pool) is mandatory to go LIVE.
        Player.objects.create(name="Tizio", role="A", league=self.league,
                              initial_price=Decimal("1"), owner=None)
        resp = self.client.post("/admin-auction/wizard/create/", {
            "league_id": str(self.league.id),
            "mode": "NEW_FROM_ZERO",
            "start_now": "1",
        })
        self.assertEqual(resp.status_code, 302)
        auction = Auction.objects.latest("id")
        self.assertEqual(auction.status, Auction.Status.LIVE)

    def test_wizard_refuses_a_league_without_a_listone(self):
        # The listone is mandatory: a league with an empty pool cannot host an
        # auction at all — the admin is sent to the import page instead.
        empty = League.objects.create(name="Lega Vuota", budget=Decimal("300"))
        resp = self.client.post("/admin-auction/wizard/create/", {
            "league_id": str(empty.id),
            "mode": "NEW_FROM_ZERO",
            "start_now": "1",
        })
        self.assertEqual(resp.status_code, 302)
        self.assertIn("need_listone=1", resp["Location"])
        self.assertFalse(Auction.objects.filter(league=empty).exists())


class SetupWizardTests(TestCase):
    """Unified 'Nuova lega' wizard: lega + (optional) import + asta in one POST."""

    def setUp(self):
        self.user = User.objects.create_superuser("admin", "a@b.c", "pass12345")
        self.client.force_login(self.user)

    def _listone(self, rows=b"Nome,R,Squadra,Qt.A\nVlahovic,A,Juventus,30\n"):
        """The mandatory Quotazioni upload every creation POST must carry."""
        from django.core.files.uploadedfile import SimpleUploadedFile
        return SimpleUploadedFile("Quotazioni.csv", rows, content_type="text/csv")

    def test_setup_page_loads(self):
        resp = self.client.get("/admin-auction/setup/")
        self.assertEqual(resp.status_code, 200)

    def test_setup_creates_league_teams_and_auction(self):
        resp = self.client.post("/admin-auction/setup/create/", {
            "name": "Lega Unica", "budget": "300",
            "source_site": "", "import_choice": "none",
            "slots_p": "3", "slots_d": "8", "slots_c": "8", "slots_a": "6",
            "participants_json": json.dumps([{"name": "Alfa", "credits": "250"},
                                             {"name": "Beta", "credits": ""}]),
            "mode": "NEW_FROM_ZERO",
            "flow_mode": "call",
            "min_increment": "1", "quick_increments": "1,2,5,10",
            "duration_seconds": "60", "antisnipe_seconds": "10",
            "enforce_limits": "1", "release_refund_mode": "purchase",
            "listone_file": self._listone(),
        })
        self.assertEqual(resp.status_code, 302)
        league = League.objects.get(name="Lega Unica")
        self.assertEqual(Participant.objects.filter(league=league).count(), 2)
        beta = Participant.objects.get(display_name="Beta")
        self.assertEqual(beta.credits, Decimal("300"))  # blank → league budget
        auction = Auction.objects.get(league=league)
        self.assertEqual(auction.mode, Auction.Mode.NEW_FROM_ZERO)
        self.assertEqual(auction.status, Auction.Status.READY)

    def test_setup_imports_fantacalcio_listone_into_league(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        csv = b"Nome,R,Squadra,Qt.A\nVlahovic,A,Juventus,30\nMaignan,P,Milan,18\n"
        upload = SimpleUploadedFile("Quotazioni.csv", csv, content_type="text/csv")
        resp = self.client.post("/admin-auction/setup/create/", {
            "name": "Lega FC", "budget": "500",
            "source_site": "fantacalcio", "import_choice": "fantacalcio",
            "slots_p": "3", "slots_d": "8", "slots_c": "8", "slots_a": "6",
            "participants_json": "[]",
            "mode": "NEW_FROM_ZERO", "flow_mode": "call",
            "listone_file": upload,
        })
        self.assertEqual(resp.status_code, 302)
        league = League.objects.get(name="Lega FC")
        pool = Player.objects.filter(league=league)
        self.assertEqual(pool.count(), 2)
        self.assertTrue(pool.filter(name="Vlahovic", role="A", owner__isnull=True).exists())

    def test_setup_start_now_goes_live(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        csv = b"Nome,R,Squadra,Qt.A\nVlahovic,A,Juventus,30\n"
        upload = SimpleUploadedFile("Quotazioni.csv", csv, content_type="text/csv")
        resp = self.client.post("/admin-auction/setup/create/", {
            "name": "Lega Live", "budget": "300", "import_choice": "none",
            "participants_json": "[]", "mode": "NEW_FROM_ZERO", "flow_mode": "call",
            "listone_file": upload, "start_now": "1",
        })
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(Auction.objects.latest("id").status, Auction.Status.LIVE)

    def test_setup_refuses_to_create_anything_without_a_listone(self):
        # The listone is mandatory: the POST is rejected *before* anything is
        # written, so no half-built league is left behind.
        resp = self.client.post("/admin-auction/setup/create/", {
            "name": "Lega NoListone", "budget": "300", "import_choice": "none",
            "participants_json": json.dumps([{"name": "Alfa", "credits": "250"}]),
            "mode": "NEW_FROM_ZERO", "flow_mode": "call", "start_now": "1",
        })
        self.assertEqual(resp.status_code, 400)
        self.assertFalse(League.objects.filter(name="Lega NoListone").exists())
        self.assertFalse(Participant.objects.filter(display_name="Alfa").exists())
        self.assertFalse(Auction.objects.exists())

    def test_setup_refuses_an_unreadable_listone(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        junk = SimpleUploadedFile("Quotazioni.csv", b"", content_type="text/csv")
        resp = self.client.post("/admin-auction/setup/create/", {
            "name": "Lega Rotta", "budget": "300", "import_choice": "none",
            "participants_json": "[]", "mode": "NEW_FROM_ZERO", "flow_mode": "call",
            "listone_file": junk,
        })
        self.assertEqual(resp.status_code, 400)
        self.assertFalse(League.objects.filter(name="Lega Rotta").exists())

    def test_setup_analyze_previews_id_based_rose_structurally(self):
        """No league exists yet at this point in the wizard, so the preview
        can only report team names + player counts straight from the file's
        own block structure — not resolve official Ids to real players."""
        from django.core.files.uploadedfile import SimpleUploadedFile
        csv_data = (
            "$,$,$\nDinamo Losca,7212,6\nDinamo Losca,7155,5\n"
            "$,$,$\nSPORTING PASSOA,6908,10\n"
        ).encode("utf-8")
        upload = SimpleUploadedFile("rose.csv", csv_data, content_type="text/csv")
        resp = self.client.post("/admin-auction/setup/analyze/", {"rose_file": upload})
        self.assertEqual(resp.status_code, 200)
        body = resp.json()
        self.assertTrue(body["ok"])
        self.assertTrue(body["id_based"])
        self.assertEqual(body["total_players"], 3)
        self.assertEqual(
            {t["name"]: t["n_players"] for t in body["teams"]},
            {"Dinamo Losca": 2, "SPORTING PASSOA": 1},
        )

    def test_setup_creates_league_from_leghe_fantacalcio_rose(self):
        """The listone must land in the DB *before* the roster for this
        source (reversed order vs. Fantapazz/Excel) since the roster has no
        player names of its own — only official Ids to match against it."""
        from django.core.files.uploadedfile import SimpleUploadedFile
        listone_csv = (
            "Id,Nome,R,Squadra,Qt.A\n"
            "7212,Lautaro,A,Inter,34\n"
            "7155,Falcone,P,Lecce,5\n"
        ).encode("utf-8")
        rose_csv = "$,$,$\nDinamo Losca,7212,6\nDinamo Losca,7155,5\n".encode("utf-8")
        resp = self.client.post("/admin-auction/setup/create/", {
            "name": "Lega Leghe", "budget": "500",
            "source_site": "leghe", "import_choice": "leghe",
            "participants_json": "[]", "mode": "NEW_FROM_ZERO", "flow_mode": "call",
            "listone_file": SimpleUploadedFile("Quotazioni.csv", listone_csv, content_type="text/csv"),
            "rose_file": SimpleUploadedFile("rose.csv", rose_csv, content_type="text/csv"),
        })
        self.assertEqual(resp.status_code, 302)
        league = League.objects.get(name="Lega Leghe")
        team = Participant.objects.get(league=league, display_name="Dinamo Losca")
        lautaro = Player.objects.get(league=league, name="Lautaro")
        self.assertEqual(lautaro.owner_id, team.id)
        self.assertEqual(lautaro.cost, Decimal("6"))
        falcone = Player.objects.get(league=league, name="Falcone")
        self.assertEqual(falcone.owner_id, team.id)
        self.assertEqual(falcone.cost, Decimal("5"))
        self.assertEqual(team.spent_credits, Decimal("11"))

    def test_setup_stores_the_no_slot_limit_choice(self):
        resp = self.client.post("/admin-auction/setup/create/", {
            "name": "Lega Libera", "budget": "300", "import_choice": "none",
            "participants_json": "[]", "mode": "NEW_FROM_ZERO", "flow_mode": "call",
            "slot_limits": "0", "listone_file": self._listone(),
        })
        self.assertEqual(resp.status_code, 302)
        league = League.objects.get(name="Lega Libera")
        self.assertFalse(league.slot_limits)
        self.assertEqual(league.total_slots, 0)
        self.assertEqual(league.slots_for("A"), 0)

    def test_setup_auto_saves_session(self):
        from ..models import AuctionSession
        resp = self.client.post("/admin-auction/setup/create/", {
            "name": "Lega Auto", "budget": "300", "import_choice": "none",
            "participants_json": json.dumps([{"name": "Alfa", "credits": "250"}]),
            "mode": "NEW_FROM_ZERO", "flow_mode": "call",
            "listone_file": self._listone(),
        })
        self.assertEqual(resp.status_code, 302)
        league = League.objects.get(name="Lega Auto")
        session = AuctionSession.objects.get(league=league)
        self.assertEqual(session.name, "Lega Auto")
        self.assertEqual(session.source_auction, Auction.objects.get(league=league))
        # The auto-snapshot captured the wizard's team.
        names = [p["display_name"] for p in session.data["participants"]]
        self.assertIn("Alfa", names)

    def _post_setup(self, **extra):
        data = {
            "name": "Lega Pronta", "budget": "300", "import_choice": "none",
            "participants_json": json.dumps([{"name": "Alfa", "credits": "", "email": "alfa@x.it"},
                                             {"name": "Beta", "credits": "", "email": "non-una-mail"}]),
            "mode": "NEW_FROM_ZERO", "flow_mode": "call", "listone_file": self._listone(),
        }
        data.update(extra)
        return self.client.post("/admin-auction/setup/create/", data)

    def test_setup_lands_on_the_ready_page_with_next_steps(self):
        resp = self._post_setup()
        league = League.objects.get(name="Lega Pronta")
        self.assertRedirects(resp, reverse("admin_setup_done", args=[league.id]))
        page = self.client.get(resp["Location"])
        self.assertContains(page, "Lega Pronta è pronta")
        self.assertContains(page, "Vai alla regia")
        self.assertContains(page, "Squadre: link, QR ed email")
        # The typed email is kept, a malformed one is left out.
        self.assertEqual(Participant.objects.get(display_name="Alfa").email, "alfa@x.it")
        self.assertEqual(Participant.objects.get(display_name="Beta").email, "")

    def test_setup_can_create_only_the_league(self):
        resp = self._post_setup(create_auction="0", start_now="1")
        league = League.objects.get(name="Lega Pronta")
        self.assertFalse(Auction.objects.filter(league=league).exists())
        page = self.client.get(resp["Location"])
        self.assertContains(page, "Crea l'asta")

    @override_settings(EMAIL_HOST="smtp.env.local")
    def test_setup_sends_the_invites_when_asked(self):
        from django.core import mail as outbox
        resp = self._post_setup(send_invites="1")
        self.assertEqual([m.to for m in outbox.outbox], [["alfa@x.it"]])
        page = self.client.get(resp["Location"])
        self.assertContains(page, "Inviti: 1 email inviata")

    def test_ready_page_of_a_foreign_league_is_forbidden(self):
        other = League.objects.create(name="Altrui", owner=User.objects.create_user("o", password="pw"))
        self.client.force_login(User.objects.create_user("x", password="pw", is_staff=True))
        self.assertEqual(self.client.get(reverse("admin_setup_done", args=[other.id])).status_code, 403)

    def test_resume_latest_empty_redirects_to_sessions(self):
        resp = self.client.post("/admin-auction/sessions/resume-latest/")
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/sessions/", resp["Location"])

    def test_dashboard_start_screen_when_no_auction(self):
        resp = self.client.get("/admin-auction/")
        self.assertContains(resp, "Riprendi sessione")
        self.assertContains(resp, "Carica sessione")

    def test_start_screen_hides_the_console_until_an_auction_is_picked(self):
        """With nothing on the table only the ways in are shown.

        The product chrome (topbar + section nav) is always there — it is how
        you reach giocatori/squadre/import — but every live-auction panel and
        the team management below it stay out until an auction is selected.
        """
        resp = self.client.get("/admin-auction/")
        body = resp.content.decode()
        self.assertIn("Da dove vuoi partire?", body)
        for absent in ("rb-activate", "Gestione squadre", "Classifica squadre",
                       "Attività live", "Cosa fare in questa pagina", "Registro offerte"):
            self.assertNotIn(absent, body)

    def test_console_comes_back_with_an_auction_selected(self):
        league = League.objects.create(name="L", budget=Decimal("500"))
        auction = Auction.objects.create(title="Asta", league=league,
                                         status=Auction.Status.READY)
        body = self.client.get(f"/admin-auction/?auction={auction.id}").content.decode()
        self.assertIn("Accesso da internet", body)
        self.assertIn("Gestione squadre", body)


class WatchlistTests(TestCase):
    """Participant watchlist + budget/slot planner (private planning aid)."""

    def setUp(self):
        self.league = League.objects.create(name="L", budget=Decimal("500"),
                                             slots_p=3, slots_d=8, slots_c=8, slots_a=6)
        self.p = Participant.objects.create(display_name="Me", league=self.league,
                                            credits=Decimal("500"))
        self.auction = make_live_auction(league=self.league)

    def _login(self):
        s = self.client.session
        s["participant_id"] = self.p.id
        s.save()

    def _free(self, name, role="A"):
        return Player.objects.create(name=name, role=role, team="Inter",
                                     league=self.league, owner=None, initial_price=Decimal("10"))

    def test_roster_plan_free_slots_and_max_bid(self):
        owned = self._free("Owned", "A")
        owned.owner = self.p; owned.cost = Decimal("40"); owned.save()
        self.p.spent_credits = Decimal("40"); self.p.save()
        plan = services.roster_plan(self.p)
        self.assertEqual(plan["remaining_credits"], Decimal("460"))
        self.assertEqual(plan["free_slots"], 24)          # 25 total − 1 owned
        self.assertEqual(plan["max_bid"], Decimal("437"))  # 460 − (24−1) reserve

    def test_search_ignores_accents(self):
        """Mid-auction nobody types the accented key: 'montipo' must find 'Montipò'."""
        self._login()
        self._free("Montipò", "P")
        r = self.client.get(f"/watch/{self.auction.id}/search/?q=montipo")
        names = [x["name"] for x in r.json()["results"]]
        self.assertIn("Montipò", names)

    def test_search_still_matches_the_accented_spelling(self):
        self._login()
        self._free("Laurientè", "A")
        r = self.client.get(f"/watch/{self.auction.id}/search/?q=Laurientè")
        self.assertEqual([x["name"] for x in r.json()["results"]], ["Laurientè"])

    def test_toggle_add_then_remove(self):
        self._login()
        pl = self._free("Target")
        r = self.client.post(f"/watch/{self.auction.id}/toggle/", {"player_id": pl.id})
        d = r.json()
        self.assertTrue(d["watching"])
        self.assertIn("watch_id", d)
        self.assertTrue(Watch.objects.filter(participant=self.p, player=pl).exists())
        r2 = self.client.post(f"/watch/{self.auction.id}/toggle/", {"player_id": pl.id})
        self.assertFalse(r2.json()["watching"])
        self.assertFalse(Watch.objects.filter(participant=self.p, player=pl).exists())

    def test_toggle_requires_session(self):
        pl = self._free("Target")
        r = self.client.post(f"/watch/{self.auction.id}/toggle/", {"player_id": pl.id})
        self.assertEqual(r.status_code, 403)

    def test_update_and_clear_max_price(self):
        self._login()
        w = Watch.objects.create(participant=self.p, player=self._free("T"))
        self.client.post(f"/watch/item/{w.id}/", {"max_price": "55"})
        w.refresh_from_db(); self.assertEqual(w.max_price, Decimal("55"))
        self.client.post(f"/watch/item/{w.id}/", {"max_price": ""})
        w.refresh_from_db(); self.assertIsNone(w.max_price)

    def test_cannot_edit_another_managers_watch(self):
        self._login()
        other = Participant.objects.create(display_name="X", league=self.league)
        w = Watch.objects.create(participant=other, player=self._free("T"))
        r = self.client.post(f"/watch/item/{w.id}/", {"max_price": "10"})
        self.assertEqual(r.status_code, 404)

    def test_search_only_free_agents_in_league(self):
        self._login()
        self._free("Lautaro Martinez", "A")
        owned = self._free("Lautaro Owned", "A"); owned.owner = self.p; owned.save()
        r = self.client.get(f"/watch/{self.auction.id}/search/?q=lauta")
        names = [x["name"] for x in r.json()["results"]]
        self.assertIn("Lautaro Martinez", names)
        self.assertNotIn("Lautaro Owned", names)

    def test_bid_page_exposes_plan_and_watches(self):
        self._login()
        Watch.objects.create(participant=self.p, player=self._free("Kvara"))
        resp = self.client.get(f"/bid/{self.auction.id}/")
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "I miei obiettivi")
        self.assertContains(resp, "Max su questo")

    def test_bid_page_exposes_watch_max_prices_for_the_live_target_chip(self):
        """The "sopra/sotto il tuo obiettivo" chip (bid.html) needs each
        watched player's noted ceiling client-side — only where one was
        actually set, a bare star with no number isn't a target."""
        self._login()
        priced = self._free("Kvara")
        unpriced = self._free("Osimhen", "A")
        Watch.objects.create(participant=self.p, player=priced, max_price=Decimal("42"))
        Watch.objects.create(participant=self.p, player=unpriced, max_price=None)
        resp = self.client.get(f"/bid/{self.auction.id}/")
        prices = json.loads(resp.context["watch_max_prices_json"])
        self.assertEqual(prices, {str(priced.id): "42.00"})


class JoinPathTests(TestCase):
    """Every supported way of identifying a bidder at /join/."""

    def test_join_via_legacy_access_code(self):
        auction = make_live_auction()
        p = Participant.objects.create(
            display_name="Coded", access_code="DRAGO23", is_active=True,
        )
        resp = self.client.post("/join/", {"access_code": "DRAGO23", "auction_id": auction.id})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(self.client.session["participant_id"], p.id)

    def test_join_via_token_in_query_string(self):
        """The shareable link form: POST to /join/?t=<token>."""
        auction = make_live_auction()
        p = Participant.objects.create(display_name="Linked", is_active=True)
        resp = self.client.post(f"/join/?t={p.public_token}", {"auction_id": auction.id})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(self.client.session["participant_id"], p.id)

    def test_participant_hot_path_preserves_token(self):
        """save(update_fields=...) must not regenerate a participant's token."""
        p = Participant.objects.create(display_name="Tok")
        original = p.public_token
        p.spent_credits = Decimal("50")
        p.save(update_fields=["spent_credits"])
        p.refresh_from_db()
        self.assertEqual(p.public_token, original)


class WizardEnforceLimitsTests(TestCase):
    """Fase C+D integration: the wizard's enforce-limits checkbox reaches the
    Auction and defaults to on when omitted."""

    def setUp(self):
        self.user = User.objects.create_superuser("admin", "a@b.c", "pass12345")
        self.client.force_login(self.user)
        self.league = League.objects.create(name="Lega", budget=Decimal("500"))
        Player.objects.create(name="Base", role="A", league=self.league,
                              initial_price=Decimal("1"), owner=None)

    def _create(self, extra):
        payload = {"league_id": str(self.league.id), "mode": "NEW_FROM_ZERO"}
        payload.update(extra)
        resp = self.client.post("/admin-auction/wizard/create/", payload)
        self.assertEqual(resp.status_code, 302)
        return Auction.objects.latest("id")

    def test_wizard_create_honors_enforce_limits_off(self):
        auction = self._create({"enforce_limits": "0"})
        self.assertFalse(auction.enforce_limits)

    def test_wizard_create_defaults_enforce_limits_on(self):
        auction = self._create({})
        self.assertTrue(auction.enforce_limits)


class MultiLeagueIsolationTests(TestCase):
    """Two leagues on the same install must never bleed into each other."""

    def setUp(self):
        self.user = User.objects.create_superuser("admin", "a@b.c", "pass12345")
        self.client.force_login(self.user)
        self.a = League.objects.create(name="Lega A", budget=Decimal("300"))
        self.b = League.objects.create(name="Lega B", budget=Decimal("700"))
        Player.objects.create(name="A1", role="A", league=self.a, initial_price=Decimal("1"))
        Player.objects.create(name="B1", role="A", league=self.b, initial_price=Decimal("1"))

    def test_a_team_is_created_inside_the_league_it_was_added_from(self):
        self.client.post("/admin-auction/participants/create/", {
            "display_name": "Squadra B1", "access_code": "B1",
            "league_id": str(self.b.id), "next": "/admin-auction/",
        })
        team = Participant.objects.get(display_name="Squadra B1")
        self.assertEqual(team.league_id, self.b.id)          # never None
        self.assertEqual(Participant.objects.filter(league=self.a).count(), 0)

    def test_credits_default_to_that_league_budget(self):
        self.client.post("/admin-auction/participants/create/", {
            "display_name": "Senza crediti", "access_code": "X",
            "league_id": str(self.b.id),
        })
        self.assertEqual(Participant.objects.get(display_name="Senza crediti").credits,
                         Decimal("700"))

    def test_without_a_league_the_team_is_refused_rather_than_orphaned(self):
        resp = self.client.post("/admin-auction/participants/create/", {
            "display_name": "Orfana", "access_code": "O", "next": "/admin-auction/",
        }, follow=True)
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(Participant.objects.filter(display_name="Orfana").exists())

    def test_the_console_only_shows_its_own_league(self):
        Participant.objects.create(display_name="Solo A", league=self.a, credits=Decimal("300"))
        Participant.objects.create(display_name="Solo B", league=self.b, credits=Decimal("700"))
        auction = Auction.objects.create(title="Asta A", league=self.a,
                                         status=Auction.Status.READY)
        body = self.client.get(f"/admin-auction/?auction={auction.id}").content.decode()
        self.assertIn("Solo A", body)
        self.assertNotIn("Solo B", body)

    def test_the_listone_page_only_shows_its_own_league(self):
        body = self.client.get(f"/admin-auction/players/?league={self.b.id}").content.decode()
        self.assertIn("B1", body)
        self.assertNotIn(">A1<", body)


class ListoneSearchAndPagingTests(TestCase):
    """The listone must not travel to the browser in one piece."""

    def setUp(self):
        self.user = User.objects.create_superuser("admin", "a@b.c", "pass12345")
        self.client.force_login(self.user)
        self.league = League.objects.create(name="Lega", budget=Decimal("500"))
        self.other = League.objects.create(name="Altra", budget=Decimal("500"))
        for i in range(120):
            Player.objects.create(name=f"Giocatore {i:03d}", role="C", team="Inter",
                                  league=self.league, initial_price=Decimal(i + 1))
        Player.objects.create(name="Lautaro Martinez", role="A", team="Inter",
                              league=self.league, initial_price=Decimal("37"))
        Player.objects.create(name="Segreto Altrui", role="A", team="Milan",
                              league=self.other, initial_price=Decimal("10"))

    # --- pagination --------------------------------------------------------

    def test_only_one_page_of_rows_is_rendered(self):
        resp = self.client.get(f"/admin-auction/players/?league={self.league.id}")
        self.assertEqual(len(resp.context["players"]), 50)
        self.assertEqual(resp.context["page_obj"].paginator.count, 121)
        # The headline count is the listone, not the page.
        self.assertEqual(resp.context["total_count"], 121)
        self.assertIn("Totale: <b>121</b>", resp.content.decode())

    def test_filters_survive_the_pagination_links(self):
        resp = self.client.get(f"/admin-auction/players/?league={self.league.id}&q=Giocatore&role=C&page=2")
        self.assertEqual(resp.context["page_obj"].number, 2)
        body = resp.content.decode()
        self.assertIn("q=Giocatore", body)
        self.assertIn("role=C", body)
        self.assertIn(f"league={self.league.id}", body)

    def test_search_and_role_filter_run_in_the_database(self):
        resp = self.client.get(f"/admin-auction/players/?league={self.league.id}&q=lautaro")
        names = [p.name for p in resp.context["players"]]
        self.assertEqual(names, ["Lautaro Martinez"])

    def test_another_league_never_shows_up(self):
        resp = self.client.get(f"/admin-auction/players/?league={self.league.id}&q=segreto")
        self.assertEqual(list(resp.context["players"]), [])

    # --- the console's player search ---------------------------------------

    def test_search_endpoint_returns_at_most_twenty_free_agents(self):
        resp = self.client.get("/admin-auction/players/search/",
                               {"q": "giocatore", "league": self.league.id})
        data = resp.json()
        self.assertTrue(data["ok"])
        self.assertEqual(len(data["results"]), 20)
        self.assertEqual(data["total"], 120)
        self.assertEqual(set(data["results"][0]), {"id", "name", "role", "team", "quotation"})

    def test_search_needs_two_characters(self):
        data = self.client.get("/admin-auction/players/search/", {"q": "l"}).json()
        self.assertEqual(data["results"], [])

    def test_search_is_scoped_to_its_league(self):
        data = self.client.get("/admin-auction/players/search/",
                               {"q": "segreto", "league": self.league.id}).json()
        self.assertEqual(data["results"], [])
        data = self.client.get("/admin-auction/players/search/",
                               {"q": "segreto", "league": self.other.id}).json()
        self.assertEqual([r["name"] for r in data["results"]], ["Segreto Altrui"])

    def test_owned_players_are_not_offered_for_the_block(self):
        team = Participant.objects.create(display_name="Alfa", league=self.league,
                                          credits=Decimal("500"))
        Player.objects.filter(name="Lautaro Martinez").update(owner=team)
        data = self.client.get("/admin-auction/players/search/",
                               {"q": "lautaro", "league": self.league.id}).json()
        self.assertEqual(data["results"], [])

    def test_the_console_no_longer_ships_the_whole_listone(self):
        auction = Auction.objects.create(title="Asta", league=self.league,
                                         flow_mode=Auction.FlowMode.CALL,
                                         status=Auction.Status.READY)
        body = self.client.get(f"/admin-auction/?auction={auction.id}").content.decode()
        self.assertNotIn("Giocatore 099", body)      # no 500-row dump
        self.assertIn("call-search", body)           # the search box instead


class ConfigDeskTests(TestCase):
    """The 'Configurazione' page: the maintenance desk for leagues/auctions."""

    def setUp(self):
        self.user = User.objects.create_superuser("admin", "a@b.c", "pass12345")
        self.client.force_login(self.user)
        self.good = League.objects.create(name="Vera", budget=Decimal("500"))
        self.team = Participant.objects.create(display_name="Alfa", league=self.good,
                                               credits=Decimal("500"))
        Player.objects.create(name="Tizio", role="A", league=self.good,
                              initial_price=Decimal("5"), owner=None)
        self.auction = make_live_auction(league=self.good, status=Auction.Status.READY)
        # The kind of leftovers a failed resume used to leave behind.
        self.empty1 = League.objects.create(name="Scarto 1")
        self.empty2 = League.objects.create(name="Scarto 2")

    def _post(self, **payload):
        return self.client.post("/admin-auction/config/action/", payload, follow=True)

    def test_page_lists_every_league_with_its_size(self):
        resp = self.client.get("/admin-auction/config/")
        self.assertEqual(resp.status_code, 200)
        rows = {r["league"].name: r for r in resp.context["rows"]}
        self.assertEqual(rows["Vera"]["pool"], 1)
        self.assertEqual(rows["Vera"]["teams"], 1)
        self.assertEqual(rows["Vera"]["auctions"], 1)
        self.assertTrue(rows["Vera"]["playable"])
        self.assertFalse(rows["Scarto 1"]["playable"])
        self.assertEqual(set(resp.context["empties"]), {self.empty1.id, self.empty2.id})

    def test_clean_empty_removes_only_the_leftovers(self):
        self._post(action="clean_empty")
        self.assertEqual([l.name for l in League.objects.all()], ["Vera"])
        self.assertTrue(Player.objects.filter(league=self.good).exists())
        self.assertTrue(Auction.objects.filter(league=self.good).exists())

    def test_deleting_a_league_takes_its_whole_tree(self):
        """SET_NULL would leave orphans in the legacy pool — it must not."""
        from ..models import AuctionSession
        services.save_session(self.auction.id, name="Snap")
        self._post(action="delete_league", league_id=self.good.id)
        self.assertFalse(League.objects.filter(pk=self.good.id).exists())
        self.assertEqual(Player.objects.count(), 0)
        self.assertEqual(Participant.objects.count(), 0)
        self.assertEqual(Auction.objects.count(), 0)
        self.assertEqual(AuctionSession.objects.count(), 0)

    def test_deleting_an_auction_keeps_the_league(self):
        self._post(action="delete_auction", auction_id=self.auction.id)
        self.assertFalse(Auction.objects.filter(pk=self.auction.id).exists())
        self.assertTrue(League.objects.filter(pk=self.good.id).exists())
        self.assertEqual(Player.objects.filter(league=self.good).count(), 1)
        self.assertEqual(Participant.objects.filter(league=self.good).count(), 1)

    def test_renaming_a_league_updates_name_budget_and_slot_mode(self):
        self._post(action="rename_league", league_id=self.good.id, name="Rinominata",
                   budget="333", slot_limits="0", slots_p="2", slots_d="7",
                   slots_c="7", slots_a="5")
        lg = League.objects.get(pk=self.good.id)
        self.assertEqual(lg.name, "Rinominata")
        self.assertEqual(lg.budget, Decimal("333"))
        self.assertFalse(lg.slot_limits)
        self.assertEqual((lg.slots_p, lg.slots_d, lg.slots_c, lg.slots_a), (2, 7, 7, 5))


class ExportTests(TestCase):
    """Export rose + classifica (recap page, xlsx, csv)."""

    def setUp(self):
        from auctions import exporters
        self.exporters = exporters
        self.user = User.objects.create_superuser("admin", "a@b.c", "pass12345")
        self.league = League.objects.create(name="Lega Test", budget=Decimal("500"))
        self.t1 = Participant.objects.create(
            display_name="Bravi", league=self.league,
            credits=Decimal("500"), spent_credits=Decimal("120"),
        )
        self.t2 = Participant.objects.create(
            display_name="Scarsi", league=self.league,
            credits=Decimal("500"), spent_credits=Decimal("30"),
        )
        Player.objects.create(name="Vlahovic", role="A", team="Juve",
                              league=self.league, owner=self.t1, cost=Decimal("90"))
        Player.objects.create(name="Bremer", role="D", team="Juve",
                              league=self.league, owner=self.t1, cost=Decimal("30"))
        Player.objects.create(name="Mandas", role="P", team="Lazio",
                              league=self.league, owner=self.t2, cost=Decimal("30"))
        # a free agent — must NOT appear in any roster
        Player.objects.create(name="Libero", role="C", league=self.league)

    def test_standings_sorted_by_spent(self):
        rows = self.exporters.build_standings(self.league)
        self.assertEqual([r["name"] for r in rows], ["Bravi", "Scarsi"])
        self.assertEqual(rows[0]["total"], 2)
        self.assertEqual(rows[0]["by_role"]["A"], 1)
        self.assertEqual(rows[0]["remaining"], Decimal("380"))

    def test_standings_excludes_free_agents(self):
        rows = self.exporters.build_standings(self.league)
        names = [pl.name for r in rows for pl in r["roster"]]
        self.assertNotIn("Libero", names)

    def test_csv_has_owned_players(self):
        data = self.exporters.build_csv(self.league)
        text = data.decode("utf-8")
        self.assertIn("Vlahovic", text)
        self.assertIn("Mandas", text)
        self.assertNotIn("Libero", text)

    def test_xlsx_builds_two_sheets(self):
        import io
        from openpyxl import load_workbook
        data = self.exporters.build_xlsx(self.league)
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["Classifica", "Rose"])

    def test_recap_page_renders(self):
        self.client.force_login(self.user)
        resp = self.client.get(f"/admin-auction/export/?league={self.league.id}")
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Vlahovic")
        self.assertContains(resp, "Classifica")

    def test_xlsx_download(self):
        self.client.force_login(self.user)
        resp = self.client.get(f"/admin-auction/export/xlsx/?league={self.league.id}")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        self.assertIn("attachment", resp["Content-Disposition"])

    def test_csv_download(self):
        self.client.force_login(self.user)
        resp = self.client.get(f"/admin-auction/export/csv/?league={self.league.id}")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("text/csv", resp["Content-Type"])

    def test_leghe_csv_format(self):
        """Reverse-engineered from a real FantaAsta Buzz export for the same
        Leghe Fantacalcio import: no header, comma-separated, a ``$,$,$``
        line opens each fantateam's block, each player row is just
        ``Fantasquadra,Id,Prezzo``."""
        Player.objects.filter(name="Vlahovic").update(ext_id="2170")
        Player.objects.filter(name="Bremer").update(ext_id="1234")
        Player.objects.filter(name="Mandas").update(ext_id="5678")
        data = self.exporters.build_leghe_csv(self.league)
        self.assertNotIn(b"\xef\xbb\xbf", data[:3])  # no BOM
        text = data.decode("utf-8")
        self.assertNotIn("\r\n", text)               # plain \n line endings
        lines = [l for l in text.split("\n") if l]
        self.assertNotIn("Libero", text)              # free agents excluded
        self.assertEqual(lines.count("$,$,$"), 2)      # one block per fantateam
        self.assertIn("Bravi,1234,30", lines)          # Fantasquadra,Id,Prezzo
        self.assertIn("Bravi,2170,90", lines)
        self.assertIn("Scarsi,5678,30", lines)
        # roster ordered P→D→C→A within a team: Bremer (D) before Vlahovic (A)
        self.assertLess(lines.index("Bravi,1234,30"), lines.index("Bravi,2170,90"))

    def test_leghe_csv_skips_players_without_an_official_id(self):
        """No name column to fall back on in this format — a player with no
        ext_id (listone imported without it, or never backfilled) simply
        cannot be represented, so it's left out rather than emitting a row
        Leghe Fantacalcio could never match."""
        # Vlahovic/Bremer/Mandas all keep ext_id="" from setUp.
        data = self.exporters.build_leghe_csv(self.league)
        text = data.decode("utf-8")
        self.assertNotIn("Vlahovic", text)
        self.assertNotIn("Bremer", text)
        self.assertNotIn("Mandas", text)
        self.assertEqual(text, "")  # nothing exportable at all here

    def test_leghe_csv_download(self):
        self.client.force_login(self.user)
        resp = self.client.get(f"/admin-auction/export/leghe/?league={self.league.id}")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("text/csv", resp["Content-Type"])
        self.assertIn("leghe-fantacalcio", resp["Content-Disposition"])


class PhotoTests(TestCase):
    """Player photos: ext_id capture, backfill, and photo_url templating."""

    def setUp(self):
        from auctions.providers import importers
        self.importers = importers
        self.user = User.objects.create_superuser("admin", "a@b.c", "pass12345")
        self.league = League.objects.create(name="L", budget=Decimal("500"))

    def test_sync_players_stores_ext_id(self):
        rows = [{"name": "Carnesecchi", "role": "P", "team": "Atalanta",
                 "price": Decimal("18"), "ext_id": "4431"}]
        self.importers.sync_players(rows, league=self.league)
        p = Player.objects.get(name="Carnesecchi", league=self.league)
        self.assertEqual(p.ext_id, "4431")

    def test_apply_photos_uses_template(self):
        Player.objects.create(name="Svilar", role="P", league=self.league, ext_id="5841")
        rep = self.importers.apply_photos(league=self.league,
                                          template="https://x/{id}.png")
        self.assertEqual(rep["set"], 1)
        p = Player.objects.get(name="Svilar")
        self.assertEqual(p.photo_url, "https://x/5841.png")

    def test_apply_photos_skips_without_ext_id(self):
        Player.objects.create(name="Ignoto", role="A", league=self.league)  # no ext_id
        rep = self.importers.apply_photos(league=self.league, template="https://x/{id}.png")
        self.assertEqual(rep["set"], 0)
        self.assertEqual(rep["without_ext_id"], 1)

    def test_apply_photos_only_missing(self):
        Player.objects.create(name="Tizio", role="A", league=self.league,
                             ext_id="100", photo_url="https://old/x.png")
        rep = self.importers.apply_photos(league=self.league,
                                          template="https://new/{id}.png", only_missing=True)
        self.assertEqual(rep["set"], 0)  # already had a photo
        p = Player.objects.get(name="Tizio")
        self.assertEqual(p.photo_url, "https://old/x.png")

    def test_backfill_ext_ids_by_name(self):
        # player imported before ids were captured
        Player.objects.create(name="Carnesecchi", role="P", team="Atalanta",
                             league=self.league)
        rows = [{"name": "Carnesecchi", "role": "P", "team": "Atalanta",
                 "price": Decimal("18"), "ext_id": "4431"}]
        n = self.importers.backfill_ext_ids(rows, league=self.league)
        self.assertEqual(n, 1)
        p = Player.objects.get(name="Carnesecchi")
        self.assertEqual(p.ext_id, "4431")

    def test_there_is_no_default_third_party_template(self):
        self.assertFalse(hasattr(self.importers, "FANTACALCIO_PHOTO_TEMPLATE"))
        with self.assertRaises(ValueError):
            self.importers.apply_photos(league=self.league, template="")

    def test_migration_drops_third_party_photo_urls(self):
        import importlib
        from django.apps import apps
        from auctions.models import Footballer
        mig = importlib.import_module("auctions.migrations.0058_clear_thirdparty_photo_urls")
        f = Footballer.objects.create(api_id=9, name="Svilar",
                                      photo_url="https://media.api-sports.io/football/players/9.png")
        linked = Player.objects.create(name="Svilar", role="P", league=self.league, footballer=f,
                                       photo_url="https://content.fantacalcio.it/web/campioncini/20/card/1.png")
        alone = Player.objects.create(name="Altro", role="P", league=self.league,
                                      photo_url="https://content.fantacalcio.it/web/campioncini/20/card/2.png")
        mine = Player.objects.create(name="Mio", role="P", league=self.league,
                                     photo_url="https://example.com/mio.png")
        mig.clear_photos(apps, None)
        for p in (linked, alone, mine):
            p.refresh_from_db()
        self.assertEqual(linked.photo_url, f.photo_url)
        self.assertEqual(alone.photo_url, "")
        self.assertEqual(mine.photo_url, "https://example.com/mio.png")

    def test_apply_photos_view_needs_a_template(self):
        Player.objects.create(name="Vlahovic", role="A", league=self.league, ext_id="2702")
        self.client.force_login(self.user)
        resp = self.client.post("/admin-auction/players/photos/",
                                {"league_id": self.league.id})
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(Player.objects.get(name="Vlahovic").photo_url, "")

    def test_apply_photos_view(self):
        Player.objects.create(name="Vlahovic", role="A", league=self.league, ext_id="2702")
        self.client.force_login(self.user)
        resp = self.client.post("/admin-auction/players/photos/",
                                {"league_id": self.league.id, "template": "https://x/{id}.png"})
        self.assertEqual(resp.status_code, 200)
        d = resp.json()
        self.assertTrue(d["ok"])
        self.assertEqual(d["set"], 1)

    def test_apply_photos_view_custom_template(self):
        Player.objects.create(name="Yildiz", role="A", league=self.league, ext_id="555")
        self.client.force_login(self.user)
        resp = self.client.post("/admin-auction/players/photos/", {
            "league_id": self.league.id,
            "template": "https://cdn.example/{team}/{id}.jpg",
        })
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json()["template"], "https://cdn.example/{team}/{id}.jpg")
        p = Player.objects.get(name="Yildiz")
        self.assertEqual(p.photo_url, "https://cdn.example//555.jpg")


class StatsTests(TestCase):
    """Decision-support stats: FVM on the listone + the Statistiche file merge."""

    def setUp(self):
        from auctions.providers import importers
        self.importers = importers
        self.user = User.objects.create_superuser("admin", "a@b.c", "pass12345")
        self.league = League.objects.create(name="L", budget=Decimal("500"))

    def _upload(self, name, text):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return SimpleUploadedFile(name, text.encode("utf-8"), content_type="text/csv")

    def test_listone_reads_the_real_leghe_fantacalcio_svincolati_export(self):
        """Regression: Leghe Fantacalcio's own "Lista calciatori svincolati"
        export — verified against a real downloaded file, not guessed —
        uses headers (#, Sq., R., QUOT., FVM/1000) that plain "id"/"sq"/"r"/
        "quotazione"/"fvm" aliases don't match. Every field silently came
        back empty or wrong (price defaulting to 1 for all 539 players,
        ext_id blank so the export/import features had nothing to key on)
        until these aliases were added.

        The blank "Costo" column matters here too: it exists in this file
        (always empty — it's a free-agents list, nobody has paid anything)
        and must not shadow "QUOT." — see the comment on the price line in
        parse_listone_file.
        """
        csv = (
            "#,Nome,Fuori lista,Sq.,Under,R.,R.MANTRA,PGv,MV,FM,FVM/1000,QUOT.,FantaSquadra,Costo\n"
            "254,Dimarco,,Inter,29,D,E/W,1,6,6,154,30,,\n"
            "5841,Svilar,,Roma,27,P,Por,0,0,0,39,18,,\n"
        )
        rows, errors = self.importers.parse_listone_file(
            self._upload("lista_calciatori_svincolati.csv", csv),
            "lista_calciatori_svincolati.csv",
        )
        self.assertEqual(errors, [])
        dimarco = next(r for r in rows if r["name"] == "Dimarco")
        self.assertEqual(dimarco["ext_id"], "254")
        self.assertEqual(dimarco["team"], "Inter")
        self.assertEqual(dimarco["role"], "D")
        self.assertEqual(dimarco["price"], Decimal("30"))
        self.assertEqual(dimarco["fvm"], Decimal("154"))
        svilar = next(r for r in rows if r["name"] == "Svilar")
        self.assertEqual(svilar["role"], "P")
        self.assertEqual(svilar["price"], Decimal("18"))

    def test_listone_captures_fvm(self):
        csv = "Id,R,Nome,Squadra,Qt.A,FVM\n1,A,Vlahovic,Juventus,20,45\n"
        rows, _ = self.importers.parse_listone_file(self._upload("Quotazioni.csv", csv), "Quotazioni.csv")
        self.assertEqual(rows[0]["fvm"], Decimal("45"))
        self.importers.sync_players(rows, league=self.league)
        self.assertEqual(Player.objects.get(name="Vlahovic").fvm, Decimal("45"))

    def test_parse_stats_file(self):
        csv = "Id,R,Nome,Squadra,Pv,Mv,Fm,Gf,Ass\n7,A,Lautaro,Inter,34,6.8,8.9,24,5\n"
        rows, errors = self.importers.parse_stats_file(self._upload("Stats.csv", csv), "Stats.csv")
        self.assertEqual(errors, [])
        r = rows[0]
        self.assertEqual(r["ext_id"], "7")
        self.assertEqual(r["presences"], 34)
        self.assertEqual(r["avg_vote"], Decimal("6.8"))
        self.assertEqual(r["fanta_avg"], Decimal("8.9"))
        self.assertEqual(r["goals"], 24)
        self.assertEqual(r["assists"], 5)

    def test_import_stats_matches_by_ext_id(self):
        Player.objects.create(name="Lautaro Martinez", role="A", team="Inter",
                              league=self.league, ext_id="7")
        csv = "Id,R,Nome,Squadra,Pv,Mv,Fm,Gf,Ass\n7,A,Lautaro DIVERSO,Inter,30,6.5,8.0,20,4\n"
        rows, _ = self.importers.parse_stats_file(self._upload("s.csv", csv), "s.csv")
        rep = self.importers.import_stats(rows, league=self.league)
        self.assertEqual(rep["matched"], 1)
        p = Player.objects.get(ext_id="7")
        self.assertEqual(p.fanta_avg, Decimal("8.0"))
        self.assertEqual(p.goals, 20)

    def test_import_stats_matches_by_name_when_no_ext_id(self):
        Player.objects.create(name="Bastoni", role="D", team="INT", league=self.league)
        csv = "Nome,Squadra,Pv,Mv,Fm,Gf,Ass\nBastoni A.,Inter,32,6.3,6.9,3,2\n"
        rows, _ = self.importers.parse_stats_file(self._upload("s.csv", csv), "s.csv")
        rep = self.importers.import_stats(rows, league=self.league)
        self.assertEqual(rep["matched"], 1)
        p = Player.objects.get(name="Bastoni")
        self.assertEqual(p.presences, 32)

    def test_import_stats_preserves_owner_and_price(self):
        part = Participant.objects.create(display_name="T", league=self.league)
        p = Player.objects.create(name="Leao", role="A", team="Milan", league=self.league,
                                  ext_id="9", owner=part, cost=Decimal("77"),
                                  initial_price=Decimal("30"))
        csv = "Id,Nome,Squadra,Pv,Fm\n9,Leao,Milan,28,7.5\n"
        rows, _ = self.importers.parse_stats_file(self._upload("s.csv", csv), "s.csv")
        self.importers.import_stats(rows, league=self.league)
        p.refresh_from_db()
        self.assertEqual(p.owner_id, part.id)
        self.assertEqual(p.cost, Decimal("77"))
        self.assertEqual(p.initial_price, Decimal("30"))
        self.assertEqual(p.fanta_avg, Decimal("7.5"))

    def test_serialize_state_exposes_stats(self):
        Player.objects.create(name="Osimhen", role="A", team="Napoli", league=self.league,
                              fanta_avg=Decimal("8.2"), goals=18, fvm=Decimal("60"))
        p = Player.objects.get(name="Osimhen")  # normalise decimals via the DB
        a = make_live_auction(player=p, league=self.league)
        stats = services.serialize_state(a)["player"]["stats"]
        # Trailing zeros trimmed, Italian comma — same shape the template's
        # floatformat renders, so the card doesn't flip separator on first state.
        self.assertEqual(stats["fanta_avg"], "8,2")
        self.assertEqual(stats["fvm"], "60")   # no decimals → no separator
        self.assertEqual(stats["goals"], 18)
        self.assertNotIn("assists", stats)  # never set → omitted

    def test_serialize_state_stats_none_when_empty(self):
        p = Player.objects.create(name="Anon", role="A", team="X", league=self.league)
        a = make_live_auction(player=p, league=self.league)
        self.assertIsNone(services.serialize_state(a)["player"]["stats"])

    def test_import_stats_view(self):
        Player.objects.create(name="Dimarco", role="D", team="Inter", league=self.league, ext_id="12")
        self.client.force_login(self.user)
        csv = "Id,Nome,Squadra,Pv,Fm\n12,Dimarco,Inter,33,7.1\n"
        resp = self.client.post("/admin-auction/players/stats/", {
            "league_id": self.league.id,
            "stats_file": self._upload("Stats.csv", csv),
        })
        self.assertEqual(resp.status_code, 200)
        d = resp.json()
        self.assertTrue(d["ok"])
        self.assertEqual(d["matched"], 1)
        self.assertEqual(Player.objects.get(name="Dimarco").fanta_avg, Decimal("7.1"))


class ServerStatsTests(TestCase):
    """Le statistiche di stagione: l'app non ne include; le dà il server
    (FANTAMANAGER_STATS_FILE) oppure le carica la lega."""

    CSV = (b"Id;R;Nome;Squadra;Pv;Mv;Fm;Gf;Ass\n"
           b"4220;A;Malen;Roma;30;6,8;8,9;9;3\n"
           b"2160;C;Dimarco;Inter;33;6,5;7,1;4;8\n")

    def setUp(self):
        import tempfile
        from django.contrib.auth import get_user_model
        self.user = get_user_model().objects.create_user(
            "regia-stat", password="x", is_staff=True, is_superuser=True)
        self.client.force_login(self.user)
        self.league = League.objects.create(name="Lega Stat")
        tmp = tempfile.NamedTemporaryFile(suffix=".csv", delete=False)
        tmp.write(self.CSV)
        tmp.close()
        self.addCleanup(os.unlink, tmp.name)
        self.stats_path = tmp.name

    def _server_file(self):
        return override_settings(FANTAMANAGER_STATS_FILE=self.stats_path,
                                 FANTAMANAGER_STATS_SEASON="2025/26")

    def test_the_repository_ships_no_third_party_stats(self):
        data = Path(importers.__file__).resolve().parent.parent / "data"
        shipped = [f for f in data.glob("*") if f.suffix.lower() in (".xlsx", ".xls", ".csv")] \
            if data.is_dir() else []
        self.assertEqual(shipped, [])

    @override_settings(FANTAMANAGER_STATS_FILE="")
    def test_without_a_server_file_nothing_is_seeded(self):
        rows, errors = importers.bundled_stats_rows()
        self.assertEqual(rows, [])
        report = importers.sync_players(
            [{"ext_id": "4220", "name": "Malen", "role": "A", "team": "Roma", "initial_price": 1}],
            league=self.league)
        self.assertEqual(report["stats_seeded"], 0)

    def test_the_server_file_parses(self):
        with self._server_file():
            rows, errors = importers.bundled_stats_rows()
        self.assertEqual(errors, [])
        self.assertEqual(len(rows), 2)

    def test_importing_a_listone_fills_the_stats_by_itself(self):
        with self._server_file():
            report = importers.sync_players(
                [{"ext_id": "4220", "name": "Malen", "role": "A", "team": "Roma",
                  "initial_price": 1}],
                league=self.league)
        self.assertEqual(report["stats_seeded"], 1)
        p = Player.objects.get(league=self.league, ext_id="4220")
        self.assertEqual(p.fanta_avg, Decimal("8.9"))
        self.assertEqual(p.presences, 30)

    def test_seeding_never_overwrites_stats_already_there(self):
        Player.objects.create(
            league=self.league, ext_id="4220", name="Malen",
            role="A", initial_price=1, fanta_avg=Decimal("3.5"))
        with self._server_file():
            importers.seed_stats(self.league)
        p = Player.objects.get(league=self.league, ext_id="4220")
        self.assertEqual(p.fanta_avg, Decimal("3.5"))

    def test_the_button_applies_the_server_file_with_no_upload(self):
        Player.objects.create(
            league=self.league, ext_id="4220", name="Malen", role="A", initial_price=1)
        with self._server_file():
            r = self.client.post("/admin-auction/players/stats/",
                                 {"league_id": self.league.id})
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertTrue(body["ok"])
        self.assertEqual(body["matched"], 1)
        self.assertIn("2025/26", body["source"])

    @override_settings(FANTAMANAGER_STATS_FILE="")
    def test_without_a_server_file_the_button_asks_for_one(self):
        r = self.client.post("/admin-auction/players/stats/", {"league_id": self.league.id})
        self.assertEqual(r.status_code, 400)
        self.assertFalse(r.json()["ok"])

    def test_an_uploaded_file_still_wins_over_the_server_one(self):
        Player.objects.create(league=self.league, name="Malen", role="A", initial_price=1)
        from django.core.files.uploadedfile import SimpleUploadedFile
        csv = SimpleUploadedFile(
            "mie.csv", b"Nome;Squadra;Pv;Mv;Fm;Gf;Ass\nMalen;Atalanta;30;7;9;20;5\n",
            content_type="text/csv")
        with self._server_file():
            r = self.client.post("/admin-auction/players/stats/",
                                 {"league_id": self.league.id, "stats_file": csv})
        self.assertEqual(r.json()["source"], "mie.csv")
        p = Player.objects.get(league=self.league, name="Malen")
        self.assertEqual(p.goals, 20)

    def test_the_console_reports_how_many_cards_are_covered(self):
        Player.objects.create(league=self.league, name="Con", role="A",
                              initial_price=1, fanta_avg=Decimal("7.5"))
        Player.objects.create(league=self.league, name="Senza", role="A", initial_price=1)
        with self._server_file():
            r = self.client.get(f"/admin-auction/players/?league={self.league.id}")
        self.assertEqual(r.context["stats_covered"], 1)
        self.assertEqual(r.context["stats_season"], "2025/26")
        self.assertTrue(r.context["stats_server_file"])


class AppShellTests(TestCase):
    """Mobile-first product shell — session-participant identity, 6 tabs."""

    TABS = ["/app/", "/app/rosa/", "/app/live/", "/app/lega/", "/app/mercato/", "/app/altro/"]

    def setUp(self):
        self.league = League.objects.create(name="Lega Test", budget=Decimal("500"),
                                             slots_p=3, slots_d=8, slots_c=8, slots_a=6)
        self.p = Participant.objects.create(display_name="Mister", league=self.league,
                                            credits=Decimal("500"), spent_credits=Decimal("120"))
        self.other = Participant.objects.create(display_name="Rivale", league=self.league,
                                                credits=Decimal("500"))
        Player.objects.create(name="Bomber", role="A", team="Inter",
                              league=self.league, owner=self.p, cost=Decimal("120"))

    def _login(self):
        s = self.client.session
        s["participant_id"] = self.p.id
        s.save()

    def test_all_tabs_redirect_to_join_without_session(self):
        for url in self.TABS:
            resp = self.client.get(url)
            self.assertEqual(resp.status_code, 302, url)
            self.assertIn("/app/login/", resp["Location"], url)

    def test_all_tabs_render_with_session(self):
        self._login()
        for url in self.TABS:
            resp = self.client.get(url)
            self.assertEqual(resp.status_code, 200, url)
            self.assertContains(resp, 'class="app-nav"')
            self.assertContains(resp, "Mister")

    def test_home_shows_roster_and_standings(self):
        self._login()
        resp = self.client.get("/app/")
        self.assertContains(resp, "Classifica lampo")
        self.assertContains(resp, "Rivale")          # other league team in standings
        self.assertContains(resp, "Mister")          # my own team highlighted in standings

    def test_rosa_groups_owned_players_by_role(self):
        self._login()
        resp = self.client.get("/app/rosa/")
        self.assertContains(resp, "Bomber")
        self.assertContains(resp, "Attaccanti")

    def test_active_auction_cta_links_into_bid_page(self):
        auction = make_live_auction(league=self.league)
        self._login()
        resp = self.client.get("/app/")
        self.assertContains(resp, f"/bid/{auction.id}/")
        self.assertContains(resp, "Asta in corso")


class ScanToJoinTests(TestCase):
    """Scanning team X's QR must put you in the auction AS team X.

    That is the whole point of the printed codes on auction night: no typing,
    no menus, and no chance of bidding as somebody else.
    """

    def setUp(self):
        from .. import remote
        self.remote = remote
        remote.stop()
        self.addCleanup(remote.stop)
        self.league = League.objects.create(name="L", budget=Decimal("500"))
        self.alfa = Participant.objects.create(display_name="Alfa", league=self.league,
                                               credits=Decimal("500"))
        self.beta = Participant.objects.create(display_name="Beta", league=self.league,
                                               credits=Decimal("500"))
        self.auction = make_live_auction(league=self.league)

    def _scan(self, participant, auction=None, **extra):
        """Follow the URL a scanned QR contains."""
        url = f"/join/?t={participant.public_token}"
        if auction is not None:
            url += f"&a={auction.id}"
        return self.client.get(url, follow=True, **extra)

    def test_scanning_lands_in_the_auction_as_that_team(self):
        r = self._scan(self.alfa, self.auction)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.redirect_chain[-1][0], f"/bid/{self.auction.id}/")
        self.assertEqual(self.client.session["participant_id"], self.alfa.id)
        self.assertContains(r, "Alfa")

    def test_each_code_signs_in_its_own_team(self):
        self._scan(self.alfa, self.auction)
        self.assertEqual(self.client.session["participant_id"], self.alfa.id)
        self._scan(self.beta, self.auction)
        self.assertEqual(self.client.session["participant_id"], self.beta.id)

    def test_scanning_works_without_the_auction_in_the_code(self):
        """One auction in the team's league: no ambiguity, go straight in."""
        r = self._scan(self.alfa)
        self.assertEqual(r.redirect_chain[-1][0], f"/bid/{self.auction.id}/")

    def test_several_running_auctions_never_show_a_picker(self):
        """A team belongs in its league's current auction - the latest running one."""
        newer = make_live_auction(league=self.league)   # a second LIVE one
        r = self._scan(self.alfa)
        self.assertEqual(r.redirect_chain[-1][0], f"/bid/{newer.id}/")
        self.assertEqual(self.client.session["participant_id"], self.alfa.id)

    def test_a_live_auction_wins_over_one_waiting_to_start(self):
        make_live_auction(league=self.league, status=Auction.Status.READY)
        r = self._scan(self.alfa)
        self.assertEqual(r.redirect_chain[-1][0], f"/bid/{self.auction.id}/")

    def test_another_leagues_auction_is_not_used(self):
        other = League.objects.create(name="Altra", budget=Decimal("500"))
        make_live_auction(league=other)
        r = self._scan(self.alfa)
        self.assertEqual(r.redirect_chain[-1][0], f"/bid/{self.auction.id}/")

    def test_an_invalid_code_signs_nobody_in(self):
        r = self.client.get("/join/?t=nonesiste", follow=True)
        self.assertEqual(r.redirect_chain, [])
        self.assertIsNone(self.client.session.get("participant_id"))

    def test_a_deactivated_team_cannot_be_used(self):
        self.alfa.is_active = False
        self.alfa.save()
        self._scan(self.alfa, self.auction)
        self.assertIsNone(self.client.session.get("participant_id"))

    # --- the same, through the internet tunnel -----------------------------

    def test_scanning_works_from_the_internet_too(self):
        host = "abc-def.trycloudflare.com"
        self.remote._harden(host, f"https://{host}")
        self.remote._set(status="on", url=f"https://{host}", host=host, pin="123456")
        r = self._scan(self.alfa, self.auction, HTTP_HOST=host)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.redirect_chain[-1][0], f"/bid/{self.auction.id}/")
        self.assertEqual(self.client.session["participant_id"], self.alfa.id)

    def test_the_qr_encodes_the_auction_and_the_reachable_address(self):
        host = "abc-def.trycloudflare.com"
        self.remote._harden(host, f"https://{host}")
        self.remote._set(status="on", url=f"https://{host}", host=host, pin="123456")
        request = RequestFactory().get("/", HTTP_HOST="localhost:8123")
        url = participant_join_url(request, self.alfa, self.auction)
        self.assertEqual(
            url, f"https://{host}/join/?t={self.alfa.public_token}&a={self.auction.id}")

    def test_qr_image_is_served_for_a_team(self):
        # The big screen's link: the auction plus its screen token.
        r = self.client.get(f"/participants/{self.alfa.id}/qr.png"
                            f"?a={self.auction.id}&t={self.auction.public_token}")
        if r.status_code == 503:
            self.skipTest("qrcode non installato")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "image/png")
        self.assertTrue(r.content.startswith(b"\x89PNG"))


class TeamBelongsToItsAuctionTests(TestCase):
    """A team plays its own league's auctions, and nobody else's.

    Rosters and credits carry from one auction of a league to the next (the
    repair auction continues where the main one left off), but a team from
    another league has no budget, no slots and no business bidding here.
    """

    def setUp(self):
        self.a_league = League.objects.create(name="Lega A", budget=Decimal("500"))
        self.b_league = League.objects.create(name="Lega B", budget=Decimal("500"))
        self.a_team = Participant.objects.create(display_name="Squadra A",
                                                 league=self.a_league, credits=Decimal("500"))
        self.b_team = Participant.objects.create(display_name="Squadra B",
                                                 league=self.b_league, credits=Decimal("500"))
        self.player = Player.objects.create(name="Tizio", role="A", league=self.a_league,
                                            initial_price=Decimal("10"))
        self.a_auction = make_live_auction(league=self.a_league, player=self.player,
                                           current_price=Decimal("10"),
                                           min_increment=Decimal("1"),
                                           quick_increments="1,5")

    def test_a_foreign_team_cannot_bid(self):
        r = services.place_bid(self.a_auction.id, self.b_team.id, 1)
        self.assertFalse(r.accepted)
        self.assertEqual(r.reason, services.Reject.WRONG_LEAGUE)
        self.a_auction.refresh_from_db()
        self.assertEqual(self.a_auction.current_price, Decimal("10"))   # untouched

    def test_the_leagues_own_team_can_bid(self):
        r = services.place_bid(self.a_auction.id, self.a_team.id, 1)
        self.assertTrue(r.accepted, msg=r.reason)

    def test_the_regia_cannot_bid_on_behalf_of_a_foreign_team(self):
        self.client.force_login(User.objects.create_superuser("admin", "a@b.c", "pw"))
        r = self.client.post(f"/admin-auction/{self.a_auction.id}/bid-for/",
                             {"participant_id": self.b_team.id, "increment": "1"})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["error"], services.Reject.WRONG_LEAGUE)

    def test_legacy_rows_without_a_league_still_work(self):
        """Pre-league installs must not be locked out of their own auctions."""
        old_team = Participant.objects.create(display_name="Storica", credits=Decimal("500"))
        old_player = Player.objects.create(name="Caio", role="A", initial_price=Decimal("10"))
        old_auction = make_live_auction(player=old_player, current_price=Decimal("10"),
                                        min_increment=Decimal("1"), quick_increments="1,5")
        r = services.place_bid(old_auction.id, old_team.id, 1)
        self.assertTrue(r.accepted, msg=r.reason)

    def test_the_bid_page_sends_a_foreign_team_to_its_own_auction(self):
        b_auction = make_live_auction(league=self.b_league)
        session = self.client.session
        session["participant_id"] = self.b_team.id
        session.save()
        r = self.client.get(f"/bid/{self.a_auction.id}/")
        self.assertEqual(r.status_code, 302)
        self.assertEqual(r["Location"], f"/bid/{b_auction.id}/")

    def test_the_bid_page_sends_them_to_join_when_they_have_no_auction(self):
        session = self.client.session
        session["participant_id"] = self.b_team.id
        session.save()
        r = self.client.get(f"/bid/{self.a_auction.id}/")
        self.assertEqual(r.status_code, 302)
        self.assertIn("/join/", r["Location"])

    def test_a_qr_naming_a_foreign_auction_does_not_hand_it_over(self):
        """?a= is a convenience, not an authorisation."""
        r = self.client.get(f"/join/?t={self.b_team.public_token}&a={self.a_auction.id}",
                            follow=True)
        self.assertNotIn(f"/bid/{self.a_auction.id}/",
                         [step[0] for step in r.redirect_chain])


class ConsoleLandingTests(TestCase):
    """Where the console lands: mid-evening it resumes, at launch it doesn't."""

    def setUp(self):
        self.user = User.objects.create_user("regia", password="x", is_staff=True,
                                             is_superuser=True)
        self.client.force_login(self.user)
        self.league = League.objects.create(name="Lega")
        self.auction = make_live_auction(league=self.league)

    def _pin(self):
        session = self.client.session
        session["current_auction_id"] = self.auction.id
        session.save()

    def test_the_console_resumes_the_auction_it_was_running(self):
        self._pin()
        r = self.client.get("/admin-auction/")
        self.assertEqual(r.status_code, 302)
        self.assertEqual(r["Location"], f"/regia/{self.auction.id}/")

    def test_home_opens_the_start_screen_even_with_an_auction_open(self):
        """What the launcher opens: avviare l'app riparte da qui."""
        self._pin()
        r = self.client.get("/admin-auction/?home=1")
        self.assertEqual(r.status_code, 200)
        self.assertIsNone(r.context["selected"])          # the dashboard, not the regia
        self.assertContains(r, "Console Gestionale")



class AdminFormIntTests(TestCase):
    """Numbers typed in the console: junk never 500s, timers stay sane."""

    def setUp(self):
        self.user = User.objects.create_superuser("admin_int", "a@b.c", "pass12345")
        self.client.force_login(self.user)
        self.auction = make_live_auction()

    def test_edit_clamps_a_zero_or_negative_duration(self):
        for raw in ("0", "-20"):
            with self.subTest(raw=raw):
                self.client.post(f"/admin-auction/{self.auction.id}/edit/", {"duration_seconds": raw})
                self.auction.refresh_from_db()
                self.assertGreaterEqual(self.auction.duration_seconds, 3)

    def test_edit_ignores_junk(self):
        r = self.client.post(f"/admin-auction/{self.auction.id}/edit/", {
            "duration_seconds": "boh", "antisnipe_seconds": "1.5"})
        self.assertNotEqual(r.status_code, 500)
        self.auction.refresh_from_db()
        self.assertEqual(self.auction.duration_seconds, 60)
        self.assertEqual(self.auction.antisnipe_seconds, 0)

    def test_queue_preview_with_junk_limit(self):
        r = self.client.get(f"/admin-auction/{self.auction.id}/queue/?limit=tanti")
        self.assertEqual(r.status_code, 200)

    def test_form_int_helper(self):
        from ..views.common import form_int
        self.assertEqual(form_int(None, 7), 7)
        self.assertEqual(form_int("", 7), 7)
        self.assertEqual(form_int(" 12 ", 7), 12)
        self.assertEqual(form_int("x", 7), 7)
        self.assertEqual(form_int("-4", 7, min_value=0), 0)
        self.assertEqual(form_int("9999", 7, max_value=100), 100)


class HealthzTests(TestCase):
    def test_healthz_answers_without_login(self):
        r = self.client.get("/healthz/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["db"], "ok")

    def test_healthz_reports_a_database_down(self):
        from django.db import connection
        with mock.patch.object(connection, "cursor", side_effect=Exception("down")):
            r = self.client.get("/healthz/")
        self.assertEqual(r.status_code, 503)
        self.assertFalse(r.json()["ok"])


class TemplateCommentTests(SimpleTestCase):
    """A ``{# … #}`` comment ends on its own line: across lines Django prints it
    on the page (the dice notes showed up under the Mercato). Longer notes use
    ``{% comment %}``."""

    def test_no_multiline_hash_comments(self):
        import re
        from pathlib import Path
        root = Path(__file__).resolve().parent.parent / "templates"
        bad = []
        for path in root.rglob("*.html"):
            for n, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
                for chunk in re.findall(r"\{#.*", line):
                    if "#}" not in chunk:
                        bad.append(f"{path.relative_to(root)}:{n}")
        self.assertEqual(bad, [])
