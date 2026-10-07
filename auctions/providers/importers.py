"""Persist normalised roster data into the DB.

Provider-agnostic: anything that produces the normalised team shape documented
in ``base`` (Fantapazz scraping, an uploaded Excel/CSV, a saved session) can be
imported through here. Keeps DB logic out of both providers and views.
"""
import re
import secrets
import unicodedata
from decimal import Decimal
from pathlib import Path

from .. import mantra

_SEP = ";"


# --- Player-name reconciliation --------------------------------------------
# The official "Quotazioni" listone and a roster imported from a fantasy site
# rarely spell a player identically: case, accents, apostrophes, hyphens, word
# order and disambiguating initials all drift ("Mctominay" vs "McTominay",
# "Bastoni A." vs "Bastoni", "Anguissa" vs "Zambo Anguissa", "Ndicka" vs
# "N'Dicka"). Matching is therefore algorithmic — never a hardcoded name list —
# so it keeps working as players join/leave Serie A every season.

def _norm(s):
    """Lowercase, strip accents/apostrophes, collapse punctuation to spaces."""
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.lower().replace("'", "").replace("’", "")
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


# Public alias: name matching that ignores accents/punctuation is useful well
# beyond the importers (player search during a live auction, for one).
normalize_name = _norm


def _name_parts(name):
    """Split a name into (surname/main tokens >=3 chars, initial tokens <3)."""
    toks = [t for t in _norm(name).split() if t]
    longs = {t for t in toks if len(t) >= 3}
    shorts = {t for t in toks if len(t) < 3}
    return longs, shorts


def _team_code(team):
    """Serie A club code = first 3 letters of the club name (INT, NAP, ...).

    Holds for every 2025/26 club; rosters store the 3-letter code, the listone
    stores the full name, and ``name[:3].upper()`` reconciles the two.
    """
    return (team or "").strip()[:3].upper()


def _decimal_or_none(s):
    """Parse a decimal cell, returning None for blank/garbage (never a default)."""
    s = str(s or "").replace(",", ".").strip()
    if not s:
        return None
    try:
        return Decimal(s)
    except Exception:
        return None


def _int_or_none(s):
    """Parse an integer cell, tolerating '3', '3.0' and '3,0'; None when blank."""
    d = _decimal_or_none(s)
    if d is None:
        return None
    try:
        return int(d)
    except Exception:
        return None


def _shorts_compatible(a, b):
    """True if two sets of initial-tokens don't contradict (prefix-compatible).

    Distinguishes same-surname, same-club players: "Martinez L." (Lautaro) must
    not match listone "Martinez Jo." (Josep), but "Martinez J." may.
    """
    if not a or not b:
        return True
    small, big = (a, b) if len(a) <= len(b) else (b, a)
    return all(any(o.startswith(s) or s.startswith(o) for o in big) for s in small)


def _find_match(row, existing, claimed):
    """Find the existing Player that best corresponds to a listone ``row``.

    ``row`` is a dict with ``name``/``role``/``team``. Tries, in order: exact
    (case-insensitive) name, accent/punctuation-normalised name, then a fuzzy
    pass keyed on club + shared surname token + compatible initials. Returns the
    matched Player (whose pk is added to ``claimed`` by the caller) or ``None``.
    """
    name = (row.get("name") or "").strip()
    if not name:
        return None

    ext_id = (row.get("ext_id") or "").strip()
    if ext_id:
        for p in existing:
            if p.pk not in claimed and p.ext_id and p.ext_id == ext_id:
                return p

    lname = name.lower()
    for p in existing:
        if p.pk not in claimed and p.name.lower() == lname:
            return p

    nname = _norm(name)
    for p in existing:
        if p.pk not in claimed and _norm(p.name) == nname:
            return p

    code = _team_code(row.get("team"))
    role = (row.get("role") or "").strip()
    l_long, l_short = _name_parts(name)
    best, best_score = None, -1
    for p in existing:
        if p.pk in claimed or _team_code(p.team) != code:
            continue
        p_long, p_short = _name_parts(p.name)
        if not (l_long & p_long) or not _shorts_compatible(l_short, p_short):
            continue
        score = 2 * len(l_long & p_long) + (role == p.role) + (l_short == p_short)
        if score > best_score:
            best, best_score = p, score
    if best is not None:
        return best

    # Se il giocatore ha cambiato club all'interno della Serie A (stesso ruolo, cognome identico, iniziali compatibili)
    if l_long and role:
        for p in existing:
            if p.pk in claimed or p.role != role:
                continue
            p_long, p_short = _name_parts(p.name)
            if (l_long & p_long) and _shorts_compatible(l_short, p_short):
                score = 3 * len(l_long & p_long) + (l_short == p_short)
                if score > best_score:
                    best, best_score = p, score
    return best


def sync_players(rows, *, league=None, replace=False, prune=False):
    """Reconcile a parsed player list (the listone) against the DB pool.

    The pool is scoped to ``league`` (the league's own listone). ``league=None``
    operates on the legacy/global pool, so existing single-league installs and
    the pre-existing tests are unaffected.

    The listone is the master pool. For each incoming ``row`` (dict with
    ``name``/``role``/``team``/``price``):

    * if it matches an existing Player (see :func:`_find_match`), that player is
      *updated* in place — role/team/quotazione refreshed — while its ``owner``
      and paid ``cost`` are preserved, so players already in a roster are never
      duplicated and stay assigned;
    * otherwise a new **free-agent** Player (``owner=None``) is created.

    Season churn:

    * ``replace=True`` wipes the whole pool first (brand-new setup);
    * ``prune=True`` deletes leftover **free agents** absent from the new listone
      (players who left Serie A). Owned players are never pruned — they stay in
      their roster until released — nor is any player currently on an auction
      block.
    """
    from ..models import Player, Auction

    if replace:
        Player.objects.filter(league=league).delete()

    existing = list(Player.objects.filter(league=league))
    claimed = set()
    created = updated = matched_owned = 0

    for row in rows:
        name = (row.get("name") or "").strip()
        if not name or len(name) < 2:
            continue
        role = (row.get("role") or "A").strip() or "A"
        team = (row.get("team") or "").strip()
        try:
            price = Decimal(str(row.get("price", 1) or 1))
        except Exception:
            price = Decimal("1")

        ext_id = (row.get("ext_id") or "").strip()
        fvm = row.get("fvm")
        mantra_roles = (row.get("mantra_roles") or "").strip()
        price_m, fvm_m = row.get("price_m"), row.get("fvm_m")

        match = None if replace else _find_match(row, existing, claimed)
        if match is not None:
            claimed.add(match.pk)
            match.role = role
            match.team = team
            match.initial_price = price
            match.name = name  # canonicalise to the official listone spelling
            fields = ["role", "team", "initial_price", "name"]
            if ext_id:
                match.ext_id = ext_id
                fields.append("ext_id")
            if fvm is not None:
                match.fvm = fvm
                fields.append("fvm")
            # I dati Mantra si aggiornano solo se il file li porta: un listone
            # senza colonna RM non deve cancellare i ruoli di uno che ce l'ha.
            if mantra_roles:
                match.mantra_roles = mantra_roles
                fields.append("mantra_roles")
            if price_m is not None:
                match.price_m = price_m
                fields.append("price_m")
            if fvm_m is not None:
                match.fvm_m = fvm_m
                fields.append("fvm_m")
            match.save(update_fields=fields)
            updated += 1
            if match.owner_id:
                matched_owned += 1
        else:
            Player.objects.create(
                league=league,
                name=name, role=role, team=team,
                initial_price=price, owner=None, ext_id=ext_id, fvm=fvm,
                mantra_roles=mantra_roles, price_m=price_m, fvm_m=fvm_m,
            )
            created += 1

    pruned = 0
    owned_not_in_listone = []
    if not replace:
        on_block = set(
            Auction.objects.exclude(player__isnull=True)
            .values_list("player_id", flat=True)
        )
        for p in existing:
            if p.pk in claimed:
                continue
            if p.owner_id:
                owned_not_in_listone.append(p.name)
            elif prune and p.pk not in on_block:
                p.delete()
                pruned += 1

    flagged_left = 0
    if league is not None and not replace:
        # 5.05: chi è in rosa ma non è più nel listone è (forse) uscito dalla Serie A.
        from ..services.abroad import flag_missing
        flagged_left = flag_missing(league, owned_not_in_listone)

    # Un listone caricato è la lista propria della lega: da qui in poi la
    # lista generale non le arriva più (finché non ci torna lei).
    if league is not None and (created or updated) and not league.own_listone:
        league.own_listone = True
        league.save(update_fields=["own_listone"])

    # I giocatori nuovi del listone si collegano subito all'anagrafica comune.
    from ..services.footballers import link_players
    link_players(Player.objects.filter(league=league) if league else Player.objects.filter(league__isnull=True))

    return {
        "created": created,
        "updated": updated,
        "matched_owned": matched_owned,
        "pruned": pruned,
        "owned_not_in_listone": owned_not_in_listone,
        "flagged_left_serie_a": flagged_left,
        "stats_seeded": seed_stats(league),
    }


# There is no default photo source: the images of a fantasy site's CDN are not
# ours to hotlink. Photos come from the API-Football registry
# (``Player.footballer``) or from a URL pattern the league admin provides for a
# source they are entitled to use. ``{id}`` = ``Player.ext_id``; ``{name}`` and
# ``{team}`` are also available.


def backfill_ext_ids(rows, *, league=None):
    """Set ``Player.ext_id`` on existing players by matching a parsed listone.

    Lets photos be generated for a pool imported before ext_id was captured,
    without touching ownership/price. ``rows`` come from
    :func:`parse_listone_file` (each carrying ``ext_id``). Returns the count of
    players whose ext_id was set/updated.
    """
    from ..models import Player

    existing = list(Player.objects.filter(league=league))
    claimed = set()
    updated = 0
    for row in rows:
        ext_id = (row.get("ext_id") or "").strip()
        if not ext_id:
            continue
        match = _find_match(row, existing, claimed)
        if match is None:
            continue
        claimed.add(match.pk)
        if match.ext_id != ext_id:
            match.ext_id = ext_id
            match.save(update_fields=["ext_id"])
            updated += 1
    return updated


def apply_photos(*, league=None, template, only_missing=True):
    """Fill ``Player.photo_url`` from ``template`` for players that have an ext_id.

    ``template`` may reference ``{id}`` (the ext_id), ``{name}`` and ``{team}``.
    With ``only_missing`` (default) players that already have a photo are left
    untouched, so it is safe to re-run. Returns a report dict.
    """
    from ..models import Player

    if not (template or "").strip():
        raise ValueError("apply_photos needs a URL template")
    qs = Player.objects.all()
    qs = qs.filter(league=league) if league is not None else qs
    total = qs.count()
    with_id = no_id = set_count = 0
    for p in qs:
        if not p.ext_id:
            no_id += 1
            continue
        with_id += 1
        if only_missing and p.photo_url:
            continue
        try:
            url = template.format(id=p.ext_id, name=p.name, team=p.team)
        except (KeyError, IndexError, ValueError):
            continue
        if url != p.photo_url:
            p.photo_url = url
            p.save(update_fields=["photo_url"])
            set_count += 1
    return {"total": total, "with_ext_id": with_id, "without_ext_id": no_id, "set": set_count}


# Header tokens that identify the season-stats sheet (Fantacalcio "Statistiche").
_STATS_HEADER_TOKENS = {"id", "r", "nome", "squadra", "pv", "mv", "fm", "gf", "ass"}

# Stat columns → parser. ``fm``/``mv`` are decimals; the rest counts.
_STATS_FIELDS = {
    "presences": (("pv", "presenze", "partite", "pg", "pgv"), _int_or_none),
    "avg_vote":  (("mv", "media voto", "mediavoto", "media"), _decimal_or_none),
    "fanta_avg": (("fm", "fantamedia", "fanta media", "fmv", "fanta"), _decimal_or_none),
    "goals":     (("gf", "gol", "goal", "reti", "gol fatti"), _int_or_none),
    "assists":   (("ass", "assist", "assists", "a"), _int_or_none),
}


def parse_stats_file(file_obj, filename):
    """Parse a Fantacalcio "Statistiche" (season summary) file into stat rows.

    Companion to :func:`parse_listone_file`: the Quotazioni file gives the pool +
    prices, this one gives per-player season numbers (presences, media voto,
    fantamedia, goals, assists). Returns ``(rows, errors)`` where each row is
    ``{ext_id, name, team, presences, avg_vote, fanta_avg, goals, assists}``.
    Accepts xlsx/xls/csv; like the listone it tolerates a leading title row.
    """
    import csv
    import io

    def _col(row, *names):
        for n in names:
            for k, v in row.items():
                if k and k.strip().lower() == n.lower():
                    return "" if v is None else str(v).strip()
        return ""

    def _iter_rows():
        name = (filename or "").lower()
        if name.endswith(".xlsx") or name.endswith(".xls"):
            rows = list(_tabular_rows(file_obj, filename))
            if not rows:
                return
            header_idx = 0
            for i, row in enumerate(rows[:5]):
                cells = {str(c).strip().lower() for c in row if c not in (None, "")}
                if len(cells & _STATS_HEADER_TOKENS) >= 3:
                    header_idx = i
                    break
            headers = [str(c).strip() if c not in (None, "") else "" for c in rows[header_idx]]
            for row in rows[header_idx + 1:]:
                yield dict(zip(headers, [str(c).strip() if c is not None else "" for c in row]))
        else:
            text = file_obj.read().decode("utf-8-sig", errors="replace")
            sample = text[:2048]
            delim = ";" if sample.count(";") > sample.count(",") else ","
            yield from csv.DictReader(io.StringIO(text), delimiter=delim)

    parsed, errors = [], []
    for i, row in enumerate(_iter_rows(), 1):
        name = _col(row, "nome", "name", "giocatore", "calciatore", "nominativo")
        if not name or name.lower() in ("-", "", "none", "null"):
            continue
        ext_id = _col(row, "id", "id_giocatore", "playerid", "#")
        team = _col(row, "squadra", "team", "club", "sq", "sq.")
        out = {"ext_id": ext_id, "name": name, "team": team}
        for field, (aliases, conv) in _STATS_FIELDS.items():
            out[field] = conv(_col(row, *aliases))
        parsed.append(out)
    return parsed, errors


# --- Statistiche del server -------------------------------------------------
# The app no longer ships a season-stats export: those numbers belong to whoever
# published them, and a public repository cannot redistribute them. A server
# admin who owns or licensed a stats file can point ``FANTAMANAGER_STATS_FILE``
# at it (e.g. a file on the NAS) and every import seeds from it exactly as the
# bundled file used to; without it, leagues upload their own file from the
# Giocatori page. ``FANTAMANAGER_STATS_SEASON`` is just the label shown.


def server_stats_path():
    """Path of the server-provided stats file, or ``None`` when not configured."""
    from django.conf import settings
    raw = (getattr(settings, "FANTAMANAGER_STATS_FILE", "") or "").strip()
    return Path(raw) if raw else None


def server_stats_season():
    from django.conf import settings
    return (getattr(settings, "FANTAMANAGER_STATS_SEASON", "") or "").strip()


def bundled_stats_rows():
    """Parse the server-provided stats file → ``(rows, errors)``.

    Not configured, or missing on disk, comes back as an empty result instead of
    an exception: an auction that opens without stat tiles is a far better
    failure than one that will not open.
    """
    path = server_stats_path()
    if path is None:
        return [], ["Nessun file di statistiche configurato sul server."]
    try:
        with open(path, "rb") as fh:
            return parse_stats_file(fh, path.name)
    except OSError as exc:
        return [], [f"Statistiche del server non disponibili: {exc}"]


def apply_bundled_stats(*, league=None, only_missing=False):
    """Merge the server stats onto a league's pool. Report, or ``None`` if absent."""
    rows, _errors = bundled_stats_rows()
    if not rows:
        return None
    return import_stats(rows, league=league, only_missing=only_missing)


def seed_stats(league=None):
    """Fill in stats for players that have none, straight after an import.

    Every route that puts players into a pool ends here, so on a server with
    ``FANTAMANAGER_STATS_FILE`` a league that never opens the Statistiche page
    still gets full cards. Without it this is a no-op. Deliberately silent about
    failure: an import must not be reported as broken because the extra numbers
    could not be attached.
    """
    try:
        report = apply_bundled_stats(league=league, only_missing=True)
    except Exception:
        return 0
    return (report or {}).get("matched", 0)


def import_stats(rows, *, league=None, only_missing=False):
    """Merge parsed season stats onto the existing player pool (never creates).

    Matches each stat row to a Player first by ``ext_id`` (exact, unambiguous),
    then falls back to the fuzzy name matcher. Only the stat fields are touched —
    ownership, price and photo are left alone — so it is safe to re-run and safe
    to run before or after the roster import. Returns a report dict.

    ``only_missing`` skips players that already carry any stat. That is what the
    automatic seeding after an import uses: filling in blanks is a courtesy,
    overwriting the file a league deliberately uploaded is not.
    """
    from ..models import Player

    existing = list(Player.objects.filter(league=league))
    if only_missing:
        existing = [
            p for p in existing
            if not any(getattr(p, f) not in (None, "") for f in _STATS_FIELDS)
        ]
        if not existing:
            return {"matched": 0, "missing": 0, "total": len(rows)}
    by_ext = {}
    for p in existing:
        if p.ext_id:
            by_ext.setdefault(p.ext_id, p)

    claimed = set()
    matched = missing = 0
    for row in rows:
        ext_id = (row.get("ext_id") or "").strip()
        player = by_ext.get(ext_id) if ext_id else None
        if player is None or player.pk in claimed:
            player = _find_match(row, existing, claimed)
        if player is None:
            missing += 1
            continue
        claimed.add(player.pk)
        fields = []
        for field in _STATS_FIELDS:
            val = row.get(field)
            if val is not None:
                setattr(player, field, val)
                fields.append(field)
        if fields:
            player.save(update_fields=fields)
            matched += 1
    return {"matched": matched, "missing": missing, "total": len(rows)}


def import_rose_data(teams_data, replace=False, league=None, default_budget=None):
    """Assign owners (rose) *onto* the existing listone pool — never wipe it.

    Rose and listone are two complementary files: the listone (Quotazioni) is the
    master pool of every player; a roster file only says which of those players is
    already owned and at what paid cost. The svincolati shown at the auction are
    exactly the pool's free agents (``owner is None``), so importing rosters must
    **only set ownership** — it must never delete the listone, or the auction is
    left with no players to call ("all'asta non ho la lista").

    For each roster player we therefore reconcile against the existing pool with
    the same fuzzy matcher used for the listone (:func:`_find_match`), so "Bastoni
    A." in a roster maps onto listone "Bastoni" instead of creating a duplicate.
    A player absent from the listone (not yet imported) is created as owned.

    ``replace=True`` resets every existing assignment first (owner cleared, paid
    cost zeroed) so a re-import doesn't leave stale owners — but the player pool
    itself, and thus the free-agent listone, is always preserved.

    Each Participant's total budget = remaining credits + sum of player costs,
    so ``remaining_credits`` matches the value shown by the source site. When a
    team carries no remaining-credits figure (e.g. the Fantacalcio.it flat export
    only lists each player's cost, not the budget), ``default_budget`` is used as
    the team's *total* budget so ``remaining_credits = default_budget - spent``.
    """
    from ..models import Participant, Player

    existing = list(Player.objects.filter(league=league))

    if replace:
        Player.objects.filter(league=league).update(owner=None, cost=Decimal("0"))
        for p in existing:
            p.owner_id = None
            p.cost = Decimal("0")

    claimed = set()
    teams_created = 0
    players_created = 0

    for t in teams_data:
        name = (t.get("name") or "").strip()
        if not name:
            continue
        players = t.get("players", [])
        spent   = sum(Decimal(str(p.get("cost", 0) or 0)) for p in players)
        remaining = t.get("credits")
        if remaining is not None:
            total = Decimal(str(remaining)) + spent
        elif default_budget is not None:
            total = Decimal(str(default_budget))
        else:
            total = spent

        defaults = {
            "access_code": secrets.token_hex(3),
            "credits": total,
            "spent_credits": spent,
            "is_active": True,
        }
        if league is not None:
            defaults["league"] = league
        if t.get("external_id"):
            defaults["external_team_id"] = str(t["external_id"])

        participant, created = Participant.objects.get_or_create(
            display_name=name, defaults=defaults,
        )
        if not created:
            participant.credits = total
            participant.spent_credits = spent
            if league is not None:
                participant.league = league
            if t.get("external_id"):
                participant.external_team_id = str(t["external_id"])
            participant.save(update_fields=["credits", "spent_credits", "league", "external_team_id"])
        teams_created += 1 if created else 0

        for p in players:
            pname = (p.get("name") or "").strip()
            if not pname:
                continue
            cost = Decimal(str(p.get("cost", 0) or 0))
            row = {"name": pname, "role": p.get("role", "A"), "team": p.get("club", "")}

            match = _find_match(row, existing, claimed)
            if match is not None:
                # Already in the listone: just assign it. Keep the player's
                # listone quotazione (initial_price) as the auction base price;
                # only record the paid cost and owner.
                claimed.add(match.pk)
                match.owner = participant
                match.cost = cost
                match.save(update_fields=["owner", "cost"])
            else:
                # Not in the listone yet (e.g. rosters imported before the
                # Quotazioni): create it as an owned player.
                new = Player.objects.create(
                    league=league,
                    name=pname,
                    role=p.get("role", "A"),
                    team=p.get("club", ""),
                    cost=cost,
                    initial_price=cost or Decimal("1"),
                    owner=participant,
                )
                existing.append(new)
                claimed.add(new.pk)
            players_created += 1

    return {
        "teams": teams_created,
        "players": players_created,
        "stats_seeded": seed_stats(league),
    }


def import_players_simple(players, replace=False, league=None):
    """Create free-agent Players (no ownership) from a flat parsed list.

    Used by the 'import giocatori' action: only the player pool matters, costs
    and ownership are ignored. Scoped to ``league`` (None = legacy pool).
    """
    from ..models import Player

    if replace:
        Player.objects.filter(league=league).delete()

    created = 0
    for p in players:
        name = (p.get("name") or "").strip()
        if not name or len(name) < 2:
            continue
        Player.objects.get_or_create(
            name=name, league=league,
            defaults={
                "role": p.get("role", "A"),
                "team": p.get("team", ""),
                "initial_price": Decimal("1"),
            },
        )
        created += 1
    seed_stats(league)
    return created


def parse_rose_xls(path_or_bytes):
    """Re-export of the Fantapazz 'Rose Lega' .xls parser (see ``fp_rose``)."""
    from ..fp_rose import parse_rose_xls as _parse
    return _parse(path_or_bytes)


# --- Listone (player pool) file parsing ------------------------------------

_LISTONE_ROLE_MAP = {
    "p": "P", "por": "P", "portiere": "P", "goalkeeper": "P", "gk": "P",
    "d": "D", "dif": "D", "difensore": "D", "defender": "D", "dd": "D", "ds": "D", "dc": "D",
    "c": "C", "cen": "C", "centrocampista": "C", "midfielder": "C", "mid": "C", "m": "C", "e": "C", "t": "C",
    "a": "A", "att": "A", "attaccante": "A", "forward": "A", "fwd": "A", "w": "A", "pc": "A", "tr": "A",
    "1": "P", "2": "D", "3": "C", "4": "A",
}

_LISTONE_HEADER_TOKENS = {
    "id", "r", "rm", "nome", "name", "giocatore", "calciatore",
    "ruolo", "role", "squadra", "team", "club",
}


def parse_listone_file(file_obj, filename):
    """Parse an official "Quotazioni"/listone file (xlsx/xls/csv) into rows.

    Returns ``(rows, errors)`` where each row is ``{name, role, team, price}``.
    Shared by the Giocatori import page and the unified setup wizard (the
    Fantacalcio.it / Excel listone source).
    """
    import csv
    import io

    def _col(row, *names):
        for n in names:
            for k, v in row.items():
                if k and k.strip().lower() == n.lower():
                    return (v or "").strip()
        return ""

    def _price(s):
        try:
            return Decimal(str(s).replace(",", ".").strip() or "1")
        except Exception:
            return Decimal("1")

    def _iter_rows():
        name = (filename or "").lower()
        if name.endswith(".xlsx") or name.endswith(".xls"):
            import openpyxl
            raw = file_obj.read()
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return
            header_idx = 0
            for i, row in enumerate(rows[:5]):
                cells = {str(c).strip().lower() for c in row if c is not None}
                if len(cells & _LISTONE_HEADER_TOKENS) >= 2:
                    header_idx = i
                    break
            headers = [str(c).strip() if c is not None else "" for c in rows[header_idx]]
            for row in rows[header_idx + 1:]:
                yield dict(zip(headers, [str(c).strip() if c is not None else "" for c in row]))
        else:
            text = file_obj.read().decode("utf-8-sig", errors="replace")
            sample = text[:2048]
            delim = ";" if sample.count(";") > sample.count(",") else ","
            yield from csv.DictReader(io.StringIO(text), delimiter=delim)

    parsed, errors = [], []
    for i, row in enumerate(_iter_rows(), 1):
        name = _col(row, "nome", "name", "giocatore", "player", "calciatore", "nominativo", "cognome")
        role_raw = _col(row, "ruolo", "role", "r", "r.", "pos", "rm", "ru").lower()
        role = _LISTONE_ROLE_MAP.get(role_raw, "A")
        team = _col(row, "squadra", "team", "club", "sq", "sq.", "società", "societa")
        # Quotation-specific synonyms first: "costo"/"crediti" mean "amount
        # already paid" in a roster file, and stay blank on a free-agents
        # listone — if tried first, _col stops at that (empty) column the
        # moment it exists, never reaching "quot." where the real price is.
        price = _price(_col(row, "quotazione", "quota", "quot.", "qt.a", "qt.i", "qt",
                            "q", "price", "valore", "fvm", "costo", "crediti") or "1")
        # Official Fantacalcio player id — "Id" in the Quotazioni file, "#" in
        # Leghe Fantacalcio's "Lista calciatori" export — used to build photo
        # URLs and to match the Id-based roster import (see
        # _parse_id_based_roster) against this listone.
        ext_id = _col(row, "id", "id_giocatore", "playerid", "#")
        # Fanta market value (col "FVM", or "FVM/1000" in the Leghe Fantacalcio
        # export) — a decision-support figure shown in the auction card, kept
        # distinct from the quotazione used as the base price.
        fvm = _decimal_or_none(_col(row, "fvm", "fvm/1000", "fantavalore", "valore di mercato"))
        # Mantra: la colonna RM porta i ruoli veri, anche multipli ("B;Dd;E"),
        # e il listone affianca a ogni numero Classic il suo gemello Mantra
        # ("Qt.A M", "FVM M") — in Mantra un esterno che fa anche l'ala vale
        # un'altra cifra. Li leggiamo sempre: costano nulla da tenere e una lega
        # può passare a Mantra dopo aver già importato il listone.
        mantra_roles = _col(row, "rm", "r.mantra", "rm.", "ruolo mantra", "mantra")
        roles = mantra.parse_roles(mantra_roles)
        price_m = _price(_col(row, "qt.a m", "qt.i m", "quotazione mantra") or "")             if _col(row, "qt.a m", "qt.i m", "quotazione mantra") else None
        fvm_m = _decimal_or_none(_col(row, "fvm m", "fvm.m", "fvm mantra"))
        # Un listone di sola colonna RM (capita nei file rifatti a mano) non
        # deve far diventare attaccanti tutti quanti: il ruolo classico si
        # ricava dal primo ruolo Mantra elencato.
        if roles and not _col(row, "ruolo", "role", "r", "r.", "pos", "ru"):
            role = mantra.classic_role(roles) or role
        if not name or name.lower() in ("-", "", "none", "null"):
            errors.append(f"Riga {i}: nome mancante")
            continue
        parsed.append({"name": name, "role": role, "team": team, "price": price,
                       "ext_id": ext_id, "fvm": fvm,
                       "mantra_roles": _SEP.join(roles), "price_m": price_m, "fvm_m": fvm_m})
    return parsed, errors


# --- Roster (rose) file parsing — auto-detecting the source format ---------

# Column-name aliases for the flat "one row per player" roster export
# (Fantacalcio.it "Lista calciatori" / Leghe Fantacalcio, and generic files).
# Order matters only within a group; the first header that matches wins.
_ROSE_NAME_COLS  = ("nome", "calciatore", "giocatore", "player", "nominativo", "cognome")
_ROSE_ROLE_COLS  = ("r.", "ruolo", "role", "r", "pos", "ru")          # NB: never R.MANTRA
_ROSE_CLUB_COLS  = ("sq.", "sq", "squadra", "club", "team", "società", "societa")
_ROSE_COST_COLS  = ("costo", "prezzo", "price", "pagato", "crediti", "cost")
_ROSE_QUOT_COLS  = ("quot.", "quot", "quotazione", "quota", "qt.a", "qt", "fvm", "valore")
# The owner ("which fanta-team holds this player"). Kept distinct from the Serie A
# club above — "Sq." is the real club, "FantaSquadra" is the owner.
_ROSE_OWNER_COLS = (
    "fantasquadra", "fanta squadra", "squadra fantacalcio", "fantallenatore",
    "fantallenatori", "proprietario", "rosa", "manager", "fantateam",
)


def _norm_header(s):
    return str(s or "").strip().lower()


def _pick_col(headers, candidates, *, exclude=()):
    """Return the index of the first header equal to one of ``candidates``.

    Comparison is case-insensitive on the trimmed header text. Headers whose
    text contains any token in ``exclude`` are skipped (e.g. drop 'R.MANTRA'
    when looking for the plain role column 'R.').
    """
    norm = [_norm_header(h) for h in headers]
    for cand in candidates:
        for i, h in enumerate(norm):
            if h == cand and not any(x in h for x in exclude):
                return i
    return None


def _to_decimal(v, default="0"):
    try:
        return Decimal(str(v).replace(",", ".").strip() or default)
    except Exception:
        return Decimal(default)


def _tabular_rows(file_obj, filename):
    """Yield the raw rows (lists of cell values) of an xlsx/xls/csv file."""
    import csv
    import io

    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_obj.read()), data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            yield ["" if c is None else c for c in row]
    elif name.endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(file_contents=file_obj.read(),
                                ignore_workbook_corruption=True)
        sh = wb.sheets()[0]
        for r in range(sh.nrows):
            yield [sh.cell_value(r, c) for c in range(sh.ncols)]
    else:
        text = file_obj.read().decode("utf-8-sig", errors="replace")
        sample = text[:2048]
        delim = ";" if sample.count(";") > sample.count(",") else ","
        for row in csv.reader(io.StringIO(text), delimiter=delim):
            yield row


def _parse_flat_roster(rows):
    """Parse a flat 'one row per player' roster table into normalised teams.

    The owner column (``FantaSquadra`` & friends) groups players into teams;
    a blank owner means the player is a free agent (svincolato) and is kept only
    in ``listone`` so it still shows up at the auction. Returns
    ``(teams, listone, source)`` or ``(None, None, None)`` when the table has no
    recognisable owner column (i.e. it isn't a roster export at all).
    """
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return None, None, None

    # Header = the first of the first 5 rows that carries a name column.
    header_idx = name_idx = None
    for i, row in enumerate(rows[:5]):
        idx = _pick_col(row, _ROSE_NAME_COLS)
        if idx is not None:
            header_idx, name_idx = i, idx
            break
    if header_idx is None:
        return None, None, None

    headers = rows[header_idx]
    owner_idx = _pick_col(headers, _ROSE_OWNER_COLS)
    if owner_idx is None:
        return None, None, None   # a listone, not a roster — let the caller fall back

    role_idx = _pick_col(headers, _ROSE_ROLE_COLS, exclude=("mantra",))
    cost_idx = _pick_col(headers, _ROSE_COST_COLS)
    quot_idx = _pick_col(headers, _ROSE_QUOT_COLS)
    # Club: prefer a dedicated club column that is NOT the owner column.
    club_idx = _pick_col(headers, _ROSE_CLUB_COLS)
    if club_idx == owner_idx:
        club_idx = None

    source = "fantacalcio" if _pick_col(headers, ("fantasquadra", "fanta squadra")) is not None else "generic"

    def cell(row, idx):
        return "" if idx is None or idx >= len(row) else str(row[idx]).strip()

    teams = {}
    listone = []
    for row in rows[header_idx + 1:]:
        name = cell(row, name_idx)
        if not name or name.lower() in ("-", "none", "null"):
            continue
        role = _LISTONE_ROLE_MAP.get(cell(row, role_idx).lower(), "A")
        club = cell(row, club_idx)
        quot = _to_decimal(cell(row, quot_idx) or "1", "1")
        listone.append({"name": name, "role": role, "team": club, "price": quot})

        owner = cell(row, owner_idx)
        if not owner or owner.lower() in ("-", "none", "null", "svincolato", "svincolati"):
            continue   # free agent: stays in the listone, owned by nobody
        cost = _to_decimal(cell(row, cost_idx), "0")
        team = teams.setdefault(owner, {"name": owner, "credits": None, "players": []})
        team["players"].append({"role": role, "name": name, "cost": cost, "club": club})

    return list(teams.values()), listone, source


_ID_BLOCK_MARKER = ["$", "$", "$"]


def preview_id_based_roster(file_obj, filename):
    """Team names + player counts + total spent, straight from the file's
    own block structure — no DB lookup, so this works before a league even
    exists (the "Nuova lega" wizard's rose preview, where there is nothing
    yet to resolve official Ids against). Returns ``None`` if the rows don't
    start with the ``$,$,$`` marker at all.

    :func:`_parse_id_based_roster` is the full, Id-resolved version used
    once the league (and its listone) actually exist.
    """
    try:
        rows = list(_tabular_rows(file_obj, filename))
    except Exception:
        return None
    rows = [[str(c).strip() for c in r] for r in rows if any(str(c).strip() for c in r)]
    if not rows or rows[0][:3] != _ID_BLOCK_MARKER:
        return None
    teams = {}
    for row in rows[1:]:
        if row[:3] == _ID_BLOCK_MARKER or len(row) < 3 or not row[0]:
            continue
        team = teams.setdefault(row[0], {"name": row[0], "n_players": 0, "cost": Decimal("0")})
        team["n_players"] += 1
        team["cost"] += _to_decimal(row[2], "0")
    return list(teams.values())


def _parse_id_based_roster(rows, league):
    """Parse the ``$,$,$``-delimited, Id-only roster export from Leghe
    Fantacalcio (fantacalcio.it) — the exact same shape
    :func:`auctions.exporters.build_leghe_csv` writes *to* it, verified
    against a real downloaded file rather than guessed: no header, a
    ``$,$,$`` row opens each fantateam's block, then one
    ``Fantasquadra,Id,Costo`` row per player.

    There is no player-name column at all in this format — identity is
    entirely the official numeric Id — so a player only resolves if the
    league's listone has already been imported with ``Player.ext_id``
    populated (the ordinary listone import already captures it). Returns
    ``(teams, unmatched)`` in the same normalised shape
    :func:`_parse_flat_roster` produces, so it drops straight into
    :func:`import_rose_data`; or ``None`` if the rows don't start with the
    marker at all, so the caller falls through to the other parsers.
    ``unmatched`` lists ``(team, id, cost)`` for every Id not found in the
    league's pool, so the caller can warn instead of silently dropping a
    purchased player.
    """
    rows = [[str(c).strip() for c in r] for r in rows if any(str(c).strip() for c in r)]
    if not rows or rows[0][:3] != _ID_BLOCK_MARKER:
        return None

    from ..models import Player

    by_id = {
        p["ext_id"]: p
        for p in Player.objects.filter(league=league).exclude(ext_id="")
                                .values("ext_id", "name", "role", "team")
    }

    teams = {}
    unmatched = []
    for row in rows[1:]:
        if row[:3] == _ID_BLOCK_MARKER:
            continue
        if len(row) < 3 or not row[0]:
            continue
        team_name, ext_id, cost_raw = row[0], row[1], row[2]
        cost = _to_decimal(cost_raw, "0")
        info = by_id.get(ext_id)
        if info is None:
            unmatched.append((team_name, ext_id, cost))
            continue
        team = teams.setdefault(team_name, {"name": team_name, "credits": None, "players": []})
        team["players"].append({
            "role": info["role"], "name": info["name"], "cost": cost, "club": info["team"],
        })

    return list(teams.values()), unmatched


def parse_rose_file(file_obj, filename, *, league=None):
    """Auto-detect a roster export's format and parse it. Source-agnostic.

    Recognises, with no per-file configuration:

    * **Fantapazz** "Rose Lega" ``.xls`` — teams laid out in horizontal 4-column
      blocks (delegated to :func:`parse_rose_xls`);
    * **Leghe Fantacalcio** "Rose Mantra/Classic" ``.csv`` — the ``$,$,$``-delimited,
      Id-only export (delegated to :func:`_parse_id_based_roster`); needs
      ``league`` so its official player Ids resolve against that league's
      already-imported listone — falls through to the flat reader without one;
    * **Fantacalcio.it / Leghe Fantacalcio** "Lista calciatori" ``.xlsx`` — a flat
      one-row-per-player table with a ``FantaSquadra`` owner column (and ``QUOT.``
      so the same file rebuilds the listone too);
    * **generic** flat exports (csv/xlsx) that carry any recognised owner column.

    Returns ``(teams, listone, meta)`` where ``teams`` feeds
    :func:`import_rose_data`, ``listone`` (may be ``None``) feeds
    :func:`sync_players` so free agents still appear at the auction, and ``meta``
    is ``{"source": ..., "n_teams": ..., "n_players": ..., "unmatched": [...]}``
    (``unmatched`` only ever populated by the Id-based format).
    """
    name = (filename or "").lower()

    # Fantapazz's binary .xls is unmistakable: its parser yields multi-team blocks.
    # Try it first for .xls, but fall back to the flat reader if it finds nothing
    # (a Fantacalcio export saved as .xls would otherwise be missed).
    if name.endswith(".xls"):
        try:
            data = file_obj.read()
            teams = parse_rose_xls(data)
        except Exception:
            teams = None
        if teams:
            meta = {"source": "fantapazz", "n_teams": len(teams),
                    "n_players": sum(len(t["players"]) for t in teams), "unmatched": []}
            return teams, None, meta
        # Re-wrap the bytes we already consumed so the flat reader can retry.
        import io
        file_obj = io.BytesIO(data if isinstance(data, (bytes, bytearray)) else b"")

    rows = list(_tabular_rows(file_obj, filename))

    if league is not None:
        id_based = _parse_id_based_roster(rows, league)
        if id_based is not None:
            teams, unmatched = id_based
            if not teams:
                raise ValueError(
                    "Il file usa gli Id ufficiali di Leghe Fantacalcio, ma nessuno "
                    "corrisponde a un giocatore di questa lega: importa prima il "
                    "listone (Quotazioni) di questa lega, poi ricarica le rose."
                )
            meta = {
                "source": "leghe_fantacalcio_id",
                "n_teams": len(teams),
                "n_players": sum(len(t["players"]) for t in teams),
                "unmatched": unmatched,
            }
            return teams, None, meta

    teams, listone, source = _parse_flat_roster(rows)
    if teams is None:
        raise ValueError(
            "Formato non riconosciuto: nessuna colonna squadra/proprietario "
            "(es. «FantaSquadra») trovata. Carica l'export rose di Fantapazz "
            "(.xls) o la «Lista calciatori» di Fantacalcio.it (.xlsx)."
        )
    meta = {
        "source": source,
        "n_teams": len(teams),
        "n_players": sum(len(t["players"]) for t in teams),
        "unmatched": [],
    }
    return teams, listone, meta


def sync_budget():
    """Recompute every participant's spent_credits from owned player costs.

    Useful after a manual roster edit. Remaining = credits - spent.
    """
    from django.db.models import Sum
    from ..models import Participant

    updated = 0
    for p in Participant.objects.all():
        spent = p.roster.aggregate(s=Sum("cost"))["s"] or Decimal("0")
        if p.spent_credits != spent:
            p.spent_credits = spent
            p.save(update_fields=["spent_credits"])
            updated += 1
    return {"updated": updated}
