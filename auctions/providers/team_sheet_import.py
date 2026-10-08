"""Import della scheda squadra: il foglio che la lega si passa prima dell'asta.

Legge sia l'Excel che esporta l'app (``team_sheets.build_team_sheets_xlsx``)
sia i PDF della lega (una scheda per pagina), e li riporta alla stessa forma:

    {"title", "president", "coach", "stadium", "capacity", "honours",
     "players": [{"role", "name", "club", "cost", "years", "status", ...}],
     "loans": [...], "reserved": [...], "images": {"logo": png, ...}}

``apply_team_sheets`` poi scrive quella forma sulla lega: aggiorna
l'intestazione della squadra, assegna i giocatori (riconosciuti sul listone
con lo stesso matcher delle altre importazioni), porta spesa e anni di
contratto e segna chi è fuori dalla Serie A. La stessa funzione fa anche
l'anteprima: esegue tutto e poi annulla la transazione, così quello che si
vede in anteprima è esattamente quello che succederà.
"""
import io
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher

from django.db import transaction
from django.utils import timezone

from ..models.participant import generate_access_code
from ..uploads import MAX_IMPORT_ROWS, MAX_PDF_PAGES, UploadRejected, clean_image_bytes
from .importers import _find_match, _name_parts, _norm, _shorts_compatible, _team_code

ROLE_BY_TITLE = {"PORTIERI": "P", "DIFENSORI": "D", "CENTROCAMPISTI": "C", "ATTACCANTI": "A"}
ROLE_ORDER = ["P", "D", "C", "A"]

# Colori della legenda del foglio originale (RGB 0-1): servono quando il file
# non porta una legenda leggibile.
_DEFAULT_LEGEND = [
    ((1.0, 0.949, 0.8), "expiring"),
    ((0.851, 0.886, 0.953), "renewal"),
    ((0.773, 0.878, 0.702), "pending"),
    ((0.973, 0.796, 0.678), "out"),
]
# Abbastanza stretta da non scambiare il grigio delle intestazioni (E7E6E6)
# per l'azzurro del rinnovo obbligatorio (D9E2F3).
_COLOR_TOLERANCE = 0.045


class SheetError(Exception):
    """Il file non si legge come scheda squadra: il messaggio va all'admin."""


# --- Valori delle celle ------------------------------------------------------

def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_years(raw):
    """'RIN.' → 0 (scaduto), '-' o vuoto → None (da tirare), '2' → 2."""
    s = _text(raw).upper().rstrip(".")
    if s.startswith("RIN"):
        return 0
    m = re.fullmatch(r"(\d{1,2})(?:[.,]0+)?", s)
    return int(m.group(1)) if m else None


def parse_cost(raw):
    s = _text(raw).replace(".", "").replace(",", ".") if isinstance(raw, str) else _text(raw)
    if not s:
        return None
    try:
        value = Decimal(s)
    except (InvalidOperation, ValueError):
        return None
    return value if value >= 0 else None


def parse_expiry_years(raw, season_start):
    """Scadenza di un ceduto temporaneo → anni di contratto rimasti.

    Accetta "FINE 27.28" (come la scrive l'export), "RIN." e un numero secco.
    """
    s = _text(raw).upper()
    m = re.search(r"(\d{2})\s*[./-]\s*(\d{2})", s)
    if m:
        start = int(m.group(1))
        return max(0, (start - season_start % 100) % 100 + 1)
    return parse_years(s)


def parse_sessions(raw):
    m = re.search(r"\d+", _text(raw))
    return int(m.group(0)) if m else None


def _split_founded(title):
    """"SLAVIA VIAFONDA 2019" → ("SLAVIA VIAFONDA", 2019)."""
    m = re.fullmatch(r"(.*?)\s+((?:18|19|20)\d{2})", _text(title))
    if m and m.group(1):
        return m.group(1), int(m.group(2))
    return _text(title), None


def _nice_name(s):
    """Maiuscolo della scheda → come scrive i nomi il listone ("Esposito S.")."""
    return " ".join(w[:1].upper() + w[1:].lower() for w in _text(s).split(" "))


def _is_number(value):
    return bool(re.fullmatch(r"\d{1,3}", _text(value)))


def _rgb(color):
    """Colore di riempimento (PDF o Excel) → terna RGB 0-1, o None."""
    if color is None:
        return None
    if isinstance(color, str):
        s = color.strip().lstrip("#")
        if len(s) == 8:
            s = s[2:]
        if len(s) != 6:
            return None
        try:
            return tuple(int(s[i:i + 2], 16) / 255 for i in (0, 2, 4))
        except ValueError:
            return None
    if isinstance(color, (int, float)):
        color = (color,)
    try:
        vals = [float(c) for c in color]
    except (TypeError, ValueError):
        return None
    if len(vals) == 1:
        return (vals[0],) * 3
    if len(vals) == 3:
        return tuple(vals)
    if len(vals) == 4:
        c, m, y, k = vals
        return ((1 - c) * (1 - k), (1 - m) * (1 - k), (1 - y) * (1 - k))
    return None


def _legend_key(text):
    t = _text(text).upper()
    if "FUORI" in t:
        return "out"
    if "OBBLIGATOR" in t:
        return "renewal"
    if "SCADENZA" in t:
        return "expiring"
    if "RINNOV" in t:
        return "pending"
    return None


def _status_for(rgb, legend_map):
    """Stato della riga dal colore: chiave della legenda, '' (bianca) o None (ignoto)."""
    if rgb is None:
        return None
    best, best_d = None, None
    # La legenda del file può avere una tinta più chiara delle righe (il verde
    # della legenda originale lo è): valgono sia i suoi colori sia i nostri.
    for ref, key in list(legend_map or []) + _DEFAULT_LEGEND:
        d = max(abs(a - b) for a, b in zip(rgb, ref))
        if best_d is None or d < best_d:
            best, best_d = key, d
    if best_d is not None and best_d <= _COLOR_TOLERANCE:
        return best
    return ""


def _parse_header_lines(lines, sheet):
    """Intestazione: titolo, presidente, allenatore, stadio e palmarès."""
    honours = []
    for line in lines:
        text = _text(line)
        if not text:
            continue
        m = re.match(r"(PRESIDENTE|ALLENATORE|STADIO)\s*:\s*(.*)", text, re.I)
        if m:
            key, value = m.group(1).upper(), m.group(2).strip()
            if key == "PRESIDENTE":
                sheet["president"] = value
            elif key == "ALLENATORE":
                sheet["coach"] = value
            else:
                cap = re.search(r"\(\s*([\d.\s']+)\s*\)", value)
                if cap:
                    digits = re.sub(r"\D", "", cap.group(1))
                    sheet["capacity"] = int(digits) if digits else None
                    value = value[:cap.start()].strip()
                sheet["stadium"] = value
            continue
        m = re.fullmatch(r"(.+?)\s*:\s*(\d+)", text)
        if m:
            honours.append([m.group(1).strip().upper(), int(m.group(2))])
            continue
        if not sheet.get("title"):
            sheet["title"] = text
    sheet["honours"] = honours


def _empty_sheet(source):
    return {"source": source, "title": "", "president": "", "coach": "", "stadium": "",
            "capacity": None, "honours": [], "players": [], "loans": [], "reserved": [],
            "substitutes": [], "images": {}}


def _assign_roles(players, labels):
    """Reparto di ogni riga: l'etichetta del blocco, o l'ordine P-D-C-A.

    La numerazione ricomincia da 1 a ogni reparto: è questo che separa i
    blocchi quando l'etichetta (scritta in verticale) non si legge.
    """
    block, last = -1, None
    for row in players:
        num = row.pop("num", None)
        if num is not None and (last is None or num <= last):
            block += 1
        if num is not None:
            last = num
        if row.get("role"):
            continue
        idx = max(block, 0)
        row["role"] = labels.get(idx) or (ROLE_ORDER[idx] if idx < len(ROLE_ORDER) else "A")
        row["block"] = idx


def _drop_block(row):
    for key in ("block", "top", "bottom"):
        row.pop(key, None)
    return row


# --- Excel ---------------------------------------------------------------------

def parse_team_sheet_xlsx(data, filename=""):
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise SheetError(f"{filename}: file Excel non leggibile ({exc}).")
    sheets = []
    for ws in wb.worksheets:
        if (ws.max_row or 0) > MAX_IMPORT_ROWS:
            raise SheetError(f"{filename}: foglio «{ws.title}» troppo lungo.")
        sheet = _parse_worksheet(ws, f"{filename} · {ws.title}" if len(wb.worksheets) > 1 else filename)
        if sheet is not None:
            sheets.append(sheet)
    return sheets


def _cell_rgb(cell):
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return (1.0, 1.0, 1.0)
    color = fill.fgColor
    if color is None or color.type != "rgb":
        return None  # colore del tema: non si sa quale sia senza il tema
    return _rgb(color.rgb)


def _find_row(rows, *words, start=0):
    for i in range(start, len(rows)):
        texts = " ".join(_text(c.value).upper() for c in rows[i] if c.value is not None)
        if all(w in texts for w in words):
            return i
    return None


def _cols_with(row, word, lo=0, hi=None):
    return [c.column for c in row if word in _text(c.value).upper()
            and c.column >= lo and (hi is None or c.column < hi)]


def _parse_worksheet(ws, source):
    rows = list(ws.iter_rows())
    head = None
    for i, row in enumerate(rows):
        texts = {_text(c.value).upper() for c in row}
        if {"CALCIATORE", "SPESA", "ANNI"} <= texts:
            head = i
            break
    if head is None:
        return None

    sheet = _empty_sheet(source)
    header_row = rows[head]

    def first(word, after):
        cols = [c.column for c in header_row if _text(c.value).upper() == word and c.column > after]
        return min(cols) if cols else None

    name_c = first("CALCIATORE", 0)
    club_c = first("SQUADRA", name_c)
    cost_c = first("SPESA", club_c or name_c)
    years_c = first("ANNI", cost_c or name_c)
    if None in (club_c, cost_c, years_c):
        return None
    id_c = first("ID", years_c)
    sub_c = first("SOSTITUTO", years_c)

    _parse_header_lines(
        [c.value for r in rows[:head] for c in r if isinstance(c.value, str)], sheet)

    ops = _find_row(rows, "OPERAZIONI", start=head + 1)
    pre = _find_row(rows, "PRELAZIONE", start=head + 1)
    end = ops if ops is not None else (pre if pre is not None else len(rows))

    legend_map = _xlsx_legend(rows, pre)
    labels, block = {}, -1
    last_num = None
    sub_end = id_c if id_c else (sub_c + 5 if sub_c else None)
    for row in rows[head + 1:end]:
        by_col = {c.column: c for c in row}
        num, name, name_cell = None, "", None
        for col in range(1, club_c):
            cell = by_col.get(col)
            value = _text(cell.value) if cell is not None else ""
            if not value:
                continue
            if value.upper() in ROLE_BY_TITLE:
                # L'etichetta sta nella prima riga del blocco (o in quella
                # grigia che lo apre): vale per il blocco che sta per iniziare.
                labels[block + 1] = ROLE_BY_TITLE[value.upper()]
                continue
            if num is None and not name and _is_number(value):
                num = int(value)
            elif not name:
                name, name_cell = value, cell
        if num is not None:
            if last_num is None or num <= last_num:
                block += 1
            last_num = num
        if sub_c:
            extra = [_text(by_col[c].value) for c in range(sub_c, sub_end)
                     if c in by_col and _text(by_col[c].value) and not _is_number(by_col[c].value)]
            if extra:
                sheet["substitutes"].append(extra[0])
        if not name:
            if num is not None:
                sheet["players"].append({"num": num})
            continue

        def val(col):
            return by_col[col].value if col in by_col else None

        sheet["players"].append({
            "num": num, "name": name,
            "club": _text(val(club_c)),
            "cost": parse_cost(val(cost_c)), "cost_raw": _text(val(cost_c)),
            "years": parse_years(val(years_c)), "years_raw": _text(val(years_c)),
            "status": _status_for(_cell_rgb(name_cell), legend_map),
            "ext_id": _text(val(id_c)) if id_c else "",
        })
    _assign_roles(sheet["players"], labels)
    sheet["players"] = [_drop_block(r) for r in sheet["players"] if r.get("name")]

    if ops is not None:
        sub = _find_row(rows, "CEDUTO", start=ops)
        stop = pre if pre is not None else len(rows)
        if sub is not None and sub < stop:
            hr = rows[sub]
            cols = {"out": _cols_with(hr, "CEDUTO"), "in": _cols_with(hr, "RICEVUTO"),
                    "team": _cols_with(hr, "IMPEGNAT") or _cols_with(hr, "SOCIET"),
                    "expiry": _cols_with(hr, "SCADENZA")}
            for row in rows[sub + 1:stop]:
                by_col = {c.column: c for c in row}
                item = {k: _text(by_col[v[0]].value) if v and v[0] in by_col else "" for k, v in cols.items()}
                if item["out"] or item["in"]:
                    sheet["loans"].append(item)

    if pre is not None:
        sub = _find_row(rows, "PERSO", start=pre)
        if sub is not None:
            hr = rows[sub]
            legend_col = min(_cols_with(rows[pre], "LEGENDA") or [10 ** 6])
            perso = _cols_with(hr, "PERSO", hi=legend_col)
            after = perso[0] if perso else 0
            cols = {"name": perso,
                    "club": _cols_with(hr, "SQUADRA", lo=after, hi=legend_col),
                    "cost": _cols_with(hr, "SPESA", lo=after, hi=legend_col),
                    "expiry": _cols_with(hr, "SCADENZA", lo=after, hi=legend_col)}
            for row in rows[sub + 1:]:
                by_col = {c.column: c for c in row}
                item = {k: _text(by_col[v[0]].value) if v and v[0] in by_col else "" for k, v in cols.items()}
                if item["name"]:
                    sheet["reserved"].append(item)

    sheet["images"] = _xlsx_images(ws, club_c)
    return sheet


def _xlsx_legend(rows, pre):
    if pre is None:
        return None
    legend_cols = _cols_with(rows[pre], "LEGENDA")
    if not legend_cols:
        return None
    col = legend_cols[0]
    found = []
    for row in rows[pre + 1:pre + 12]:
        for c in row:
            if c.column == col and c.value:
                key = _legend_key(c.value)
                rgb = _cell_rgb(c)
                if key and rgb:
                    found.append((rgb, key))
    return found or None


def _xlsx_images(ws, club_c):
    """Logo (a sinistra della colonna squadra) e maglie (a destra), in PNG."""
    out = {}
    try:
        from PIL import Image as PILImage
    except ImportError:
        return out
    items = []
    for img in getattr(ws, "_images", []):
        try:
            col = img.anchor._from.col
            raw = img._data()
            pil = PILImage.open(io.BytesIO(raw))
            pil.load()
            items.append((col, pil))
        except Exception:
            continue
    items.sort(key=lambda t: t[0])
    kits = []
    for col, pil in items:
        if col + 1 < club_c and "logo" not in out:
            out["logo"] = _png(pil)
        else:
            kits.append(pil)
    _store_kits(out, kits)
    return out


def _png(pil):
    buf = io.BytesIO()
    if pil.mode not in ("RGB", "RGBA"):
        pil = pil.convert("RGBA")
    pil.save(buf, format="PNG")
    return buf.getvalue()


def _store_kits(out, kits):
    """Due maglie: prima e seconda, da sinistra a destra.

    Nei PDF della lega le maglie sono spesso un'unica immagine ritagliata dal
    foglio Excel, con dentro anche le linee della tabella: ogni immagine si
    divide nei pezzi di disegno veri, e le prime due sagome sono le maglie.
    """
    pieces = [piece for pil in kits for piece in _pieces(pil)]
    for key, pil in zip(("kit_home", "kit_away"), pieces):
        out[key] = _png(pil)


def _pieces(pil, *, min_share=0.15):
    """Le sagome di un'immagine, separate da colonne bianche, senza righe della griglia."""
    rgba = pil.convert("RGBA")
    if max(rgba.size) > 900:
        # Per una scheda bastano e avanzano; e l'analisi pixel per pixel resta rapida.
        rgba.thumbnail((900, 900))
    gray = rgba.convert("L")
    w, h = gray.size
    if w < 4 or h < 4:
        return []
    px, alpha = gray.load(), rgba.getchannel("A").load()

    def raw_ink(x, y):
        return alpha[x, y] > 16 and px[x, y] < 235

    # Le righe orizzontali sottili che attraversano tutto il ritaglio (bordi
    # di cella) non sono disegno: si ignorano, se no uniscono tutto.
    row_cover = [sum(1 for x in range(0, w, 2) if raw_ink(x, y)) * 2 for y in range(h)]
    grid_rows = {y for y in range(h) if row_cover[y] >= 0.85 * w
                 and sum(1 for d in range(-3, 4) if 0 <= y + d < h and row_cover[y + d] >= 0.85 * w) <= 4}

    def ink(x, y):
        return y not in grid_rows and raw_ink(x, y)

    cover = [sum(1 for y in range(h) if ink(x, y)) for x in range(w)]
    # Una riga della tabella: poche colonne piene da cima a fondo, con accanto
    # il vuoto. Dentro una maglia le colonne vicine sono piene anche loro.
    solid = [c >= 0.95 * h for c in cover]
    x = 0
    while x < w:
        if solid[x]:
            end = x
            while end + 1 < w and solid[end + 1]:
                end += 1
            left = cover[x - 3] if x >= 3 else 0
            right = cover[end + 3] if end + 3 < w else 0
            if end - x < 5 and max(left, right) < 0.5 * h:
                for i in range(x, end + 1):
                    cover[i] = 0
            x = end + 1
        else:
            x += 1
    segments, start, gap = [], None, 0
    for x in range(w):
        if cover[x] > 0:
            if start is None:
                start = x
            gap = 0
        elif start is not None:
            gap += 1
            if gap > 3:
                segments.append((start, x - gap + 1))
                start, gap = None, 0
    if start is not None:
        segments.append((start, w - gap))
    out = []
    for x0, x1 in segments:
        if x1 - x0 < min_share * w and x1 - x0 < 0.5 * h:
            continue  # briciole: un pezzo di testo o di bordo rimasto nel ritaglio
        rows = [y for y in range(h) if any(ink(x, y) for x in range(x0, x1, 2))]
        if rows:
            out.append(rgba.crop((x0, rows[0], x1, rows[-1] + 1)))
    return out


# --- PDF -----------------------------------------------------------------------

def parse_team_sheet_pdf(data, filename=""):
    try:
        import pdfplumber
    except ImportError:
        raise SheetError("Per leggere i PDF serve il pacchetto «pdfplumber» (pip install pdfplumber).")
    try:
        import pypdfium2
        images_doc = pypdfium2.PdfDocument(data)
    except Exception:
        images_doc = None
    sheets = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if len(pdf.pages) > MAX_PDF_PAGES:
                raise SheetError(f"{filename}: troppe pagine (massimo {MAX_PDF_PAGES}).")
            for i, page in enumerate(pdf.pages):
                source = f"{filename} · pag. {i + 1}" if len(pdf.pages) > 1 else filename
                pdfium_page = images_doc[i] if images_doc is not None else None
                sheet = _parse_pdf_page(page, source, pdfium_page)
                if sheet is not None:
                    sheets.append(sheet)
    except SheetError:
        raise
    except Exception as exc:
        raise SheetError(f"{filename}: PDF non leggibile ({exc}).")
    finally:
        if images_doc is not None:
            images_doc.close()
    return sheets


class _Line:
    def __init__(self, words):
        self.words = sorted(words, key=lambda w: w["x0"])
        self.top = min(w["top"] for w in words)
        self.bottom = max(w["bottom"] for w in words)
        self.mid = (self.top + self.bottom) / 2

    @property
    def text(self):
        return " ".join(w["text"] for w in self.words)

    def upper_words(self):
        return [w["text"].upper() for w in self.words]


def _group_lines(words, tolerance=3):
    lines = []
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if lines and abs(lines[-1][0]["top"] - w["top"]) <= tolerance:
            lines[-1].append(w)
        else:
            lines.append([w])
    return [_Line(ws) for ws in lines]


def _v_edges(page, top, bottom):
    """Posizioni x dei bordi verticali delle celle fra ``top`` e ``bottom``.

    Contano i bordi veri (righe e rettangoli sottili): i riempimenti colorati
    delle celle hanno lati che non coincidono con la griglia.
    """
    def inside(o):
        return o["bottom"] - o["top"] > 4 and o["top"] >= top - 2 and o["bottom"] <= bottom + 2

    xs = [(r["x0"] + r["x1"]) / 2 for r in page.rects if r["x1"] - r["x0"] <= 2.5 and inside(r)]
    xs += [(ln["x0"] + ln["x1"]) / 2 for ln in page.lines if abs(ln["x1"] - ln["x0"]) <= 2.5 and inside(ln)]
    if len(xs) < 3:
        xs = [e["x0"] for e in page.edges if e["orientation"] == "v" and inside(e)]
    xs.sort()
    merged = []
    for x in xs:
        if merged and x - merged[-1][-1] <= 2.5:
            merged[-1].append(x)
        else:
            merged.append([x])
    return [sum(g) / len(g) for g in merged]


def _col_index(bounds, x):
    for i in range(len(bounds) - 1):
        if bounds[i] <= x < bounds[i + 1]:
            return i
    return None


def _center(w):
    return (w["x0"] + w["x1"]) / 2


def _fill_at(page, x, y):
    """Colore della cella che contiene il punto (la più piccola che lo copre)."""
    best, area = None, None
    for r in page.rects:
        if not r.get("fill"):
            continue
        if r["x0"] - 0.5 <= x <= r["x1"] + 0.5 and r["top"] - 0.5 <= y <= r["bottom"] + 0.5:
            a = (r["x1"] - r["x0"]) * (r["bottom"] - r["top"])
            if a < 1:
                continue
            if area is None or a < area:
                best, area = r.get("non_stroking_color"), a
    rgb = _rgb(best)
    return rgb if rgb is not None else (1.0, 1.0, 1.0)


def _parse_pdf_page(page, source, pdfium_page=None):
    words = page.extract_words(extra_attrs=["size"], x_tolerance=1.5, keep_blank_chars=False)
    if not words:
        return None
    lines = _group_lines(words)
    head = next((l for l in lines if {"CALCIATORE", "SPESA", "ANNI"} <= set(l.upper_words())), None)
    if head is None:
        return None
    sheet = _empty_sheet(source)

    # Intestazione: il titolo è la riga col carattere più grande.
    above = [l for l in lines if l.bottom <= head.top]
    if above:
        title = max(above, key=lambda l: max(w.get("size", 0) for w in l.words))
        rest = [l.text for l in above if l is not title]
        _parse_header_lines([title.text] + rest, sheet)

    ops = next((l for l in lines if l.top > head.bottom and "OPERAZIONI" in l.upper_words()), None)
    pre = next((l for l in lines if l.top > head.bottom and "PRELAZIONE" in l.upper_words()), None)
    roster_bottom = ops.top if ops else (pre.top if pre else page.height)

    hw = head.words

    def header_word(word, after=-1.0):
        for w in hw:
            if w["text"].upper() == word and w["x0"] > after:
                return w
        return None

    calc = header_word("CALCIATORE")
    club_w = header_word("SQUADRA", calc["x1"])
    cost_w = header_word("SPESA", club_w["x1"] if club_w else calc["x1"])
    years_w = header_word("ANNI", cost_w["x1"] if cost_w else calc["x1"])
    sub_w = header_word("SOSTITUTO", years_w["x1"] if years_w else calc["x1"])
    if not (club_w and cost_w and years_w):
        return None

    bounds = _v_edges(page, head.bottom, roster_bottom)
    cols = {}
    if len(bounds) >= 5:
        bounds = [0.0] + bounds + [float(page.width) + 1]
        cols = {k: _col_index(bounds, _center(w)) for k, w in
                (("name", calc), ("club", club_w), ("cost", cost_w), ("years", years_w))}
        cols["num"] = cols["name"] - 1 if cols["name"] else None
        if sub_w is not None:
            cols["sub"] = _col_index(bounds, _center(sub_w))
    if not cols or None in (cols.get("name"), cols.get("club"), cols.get("cost"), cols.get("years")):
        # Niente griglia: confini a metà fra le intestazioni.
        right = (sub_w["x0"] - 2) if sub_w else float(page.width) + 1
        bounds = [0.0, (calc["x1"] + club_w["x0"]) / 2, (club_w["x1"] + cost_w["x0"]) / 2,
                  (cost_w["x1"] + years_w["x0"]) / 2, (years_w["x1"] + right) / 2 if sub_w else right,
                  float(page.width) + 1]
        cols = {"num": None, "name": 0, "club": 1, "cost": 2, "years": 3, "sub": 4 if sub_w else None}
    label_right = bounds[cols["num"]] if cols.get("num") is not None else bounds[cols["name"]]

    legend_map, legend_left = _pdf_legend(page, lines, pre)
    name_x = (bounds[cols["name"]] + bounds[cols["name"] + 1]) / 2
    for line in lines:
        if not (head.bottom < line.mid < roster_bottom):
            continue
        cells = defaultdict(list)
        for w in line.words:
            if w["x1"] <= label_right + 0.5 and cols.get("num") is not None:
                continue  # lettere dell'etichetta verticale del reparto
            cells[_col_index(bounds, _center(w))].append(w["text"])
        get = lambda k: " ".join(cells.get(cols.get(k), [])) if cols.get(k) is not None else ""  # noqa: E731
        num_text, name = get("num"), get("name")
        if cols.get("num") is None:
            m = re.match(r"(\d{1,3})\s+(.*)", name)
            if m:
                num_text, name = m.group(1), m.group(2)
        num = int(num_text) if _is_number(num_text) else None
        if cols.get("sub") is not None:
            extra = [t for c, ts in cells.items() if c is not None and c > cols["sub"] for t in ts]
            if any(not _is_number(t) for t in extra):
                sheet["substitutes"].append(" ".join(extra))
        if num is None and not name:
            continue
        row = {"num": num, "name": name, "top": line.top, "bottom": line.bottom}
        if name:
            row.update({
                "club": get("club"), "cost": parse_cost(get("cost")), "cost_raw": get("cost"),
                "years": parse_years(get("years")), "years_raw": get("years"),
                "status": _status_for(_fill_at(page, name_x, line.mid), legend_map), "ext_id": "",
            })
        sheet["players"].append(row)

    labels = _pdf_role_labels(page, sheet["players"], label_right, head.bottom, roster_bottom)
    _assign_roles(sheet["players"], labels)
    sheet["players"] = [_drop_block(r) for r in sheet["players"] if r.get("name")]

    if ops is not None:
        stop = pre.top if pre else page.height
        sheet["loans"] = _pdf_table(page, lines, ops, stop, {
            "out": ("CEDUTO",), "in": ("RICEVUTO",), "team": ("IMPEGNATA", "SOCIETA’", "SOCIETA'"),
            "expiry": ("SCADENZA",)}, required=("out", "in"))
    if pre is not None:
        sheet["reserved"] = _pdf_table(page, lines, pre, page.height, {
            "name": ("PERSO",), "club": ("SQUADRA",), "cost": ("SPESA",), "expiry": ("SCADENZA",)},
            required=("name",), right_limit=legend_left)

    sheet["images"] = _pdf_images(pdfium_page, head.top)
    return sheet


def _pdf_role_labels(page, rows, label_right, top, bottom):
    """L'etichetta verticale di ogni reparto ("PORTIERI" scritto dal basso)."""
    blocks, block, last = [], -1, None
    for r in rows:
        num = r.get("num")
        if num is not None and (last is None or num <= last):
            block += 1
            blocks.append([r["top"], r["bottom"]])
        if num is not None:
            last = num
        if blocks:
            blocks[-1][1] = r["bottom"]
    chars = [c for c in page.chars if c["x1"] <= label_right + 0.5 and top < c["top"] < bottom and c["text"].strip()]
    labels = {}
    for i, (b_top, b_bottom) in enumerate(blocks):
        start = blocks[i - 1][1] if i else top
        inside = [c for c in chars if start - 1 <= (c["top"] + c["bottom"]) / 2 <= b_bottom + 1]
        for text in ("".join(c["text"] for c in sorted(inside, key=lambda c: -c["top"])),
                     "".join(c["text"] for c in sorted(inside, key=lambda c: c["top"]))):
            role = ROLE_BY_TITLE.get(text.upper())
            if role:
                labels[i] = role
                break
    return labels


def _pdf_legend(page, lines, pre):
    if pre is None:
        return None, None
    title = next((w for w in pre.words if w["text"].upper() == "LEGENDA"), None)
    if title is None:
        return None, None
    bounds = _v_edges(page, pre.top, page.height)
    left = max([x for x in bounds if x < title["x0"]] or [title["x0"] - 40])
    found = []
    for line in lines:
        if line.top <= pre.bottom:
            continue
        ws = [w for w in line.words if w["x0"] >= left - 1]
        if not ws:
            continue
        key = _legend_key(" ".join(w["text"] for w in ws))
        if key:
            x = (min(w["x0"] for w in ws) + max(w["x1"] for w in ws)) / 2
            found.append((_fill_at(page, x, line.mid), key))
    return (found or None), left


def _pdf_table(page, lines, title_line, stop, fields, *, required, right_limit=None):
    """Una tabellina in fondo alla scheda: intestazioni → colonne → righe."""
    sub = next((l for l in lines if title_line.bottom < l.top < stop
                and any(k in l.upper_words() for keys in fields.values() for k in keys)), None)
    if sub is None:
        return []
    limit = right_limit if right_limit is not None else float(page.width) + 1
    # I bordi partono dal filo della cella, un po' sopra il testo dell'intestazione.
    bounds = [x for x in _v_edges(page, sub.top - 8, stop) if x < limit]
    header_words = [w for w in sub.words if w["x0"] < limit]
    anchors = {}
    for key, names in fields.items():
        w = next((w for w in header_words if w["text"].upper() in names), None)
        if w is not None:
            anchors[key] = w
    if len(bounds) >= 3:
        bounds = [0.0] + bounds + [limit]
        col_of = {k: _col_index(bounds, _center(w)) for k, w in anchors.items()}
        where = lambda w: _col_index(bounds, _center(w))  # noqa: E731
    else:
        # Senza griglia: ogni parola va all'intestazione più vicina; il numero
        # di riga, a sinistra di tutte le intestazioni, non conta.
        centers = {k: _center(w) for k, w in anchors.items()}
        col_of = {k: k for k in anchors}
        first = min((w["x0"] for w in header_words), default=0) - 10

        def where(w):
            if w["x1"] < first:
                return None
            return min(centers, key=lambda k: abs(centers[k] - _center(w)))
    out = []
    for line in lines:
        if not (sub.bottom < line.mid < stop):
            continue
        cells = defaultdict(list)
        for w in line.words:
            if w["x0"] >= limit:
                continue
            cells[where(w)].append(w["text"])
        item = {k: " ".join(cells.get(col_of.get(k), [])) for k in fields}
        if any(item.get(k) for k in required):
            out.append(item)
    return out


def _pdf_images(pdfium_page, header_bottom):
    """Logo e maglie dall'intestazione: le immagini incorporate, non un ritaglio
    della pagina (che si porterebbe dietro linee e testo sovrapposti)."""
    out = {}
    if pdfium_page is None:
        return out
    try:
        import pypdfium2.raw as pdfium_c

        height = pdfium_page.get_height()
        found = []
        for obj in pdfium_page.get_objects(filter=(pdfium_c.FPDF_PAGEOBJ_IMAGE,)):
            left, bottom, right, top = obj.get_bounds()
            if height - top >= header_bottom or right - left < 10 or top - bottom < 10:
                continue
            found.append((left, obj.get_bitmap(render=True).to_pil()))
    except Exception:
        return out
    found.sort(key=lambda t: t[0])
    if not found:
        return out
    out["logo"] = _png(_trim(found[0][1]))
    _store_kits(out, [pil for _x, pil in found[1:]])
    return out


def _trim(pil):
    """Toglie il bianco (o il trasparente) attorno al disegno."""
    from PIL import Image, ImageChops

    rgba = pil.convert("RGBA")
    flat = Image.new("RGB", rgba.size, (255, 255, 255))
    flat.paste(rgba, mask=rgba.getchannel("A"))
    diff = ImageChops.difference(flat, Image.new("RGB", rgba.size, (255, 255, 255))).convert("L")
    box = diff.point(lambda v: 255 if v > 20 else 0).getbbox()
    return rgba.crop(box) if box else rgba


def parse_team_sheet_file(data, filename):
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        return parse_team_sheet_pdf(data, filename)
    if name.endswith((".xlsx", ".xlsm")):
        return parse_team_sheet_xlsx(data, filename)
    raise SheetError(f"{filename}: carica la scheda in PDF o in Excel (.xlsx).")


# --- Scrittura sulla lega ----------------------------------------------------

def _tokens(s):
    return [t for t in _norm(s).split() if t]


def _abbr_compatible(a, b):
    """"S. VIAFONDA" ~ "SLAVIA VIAFONDA", "ATL.CIMITERO" ~ "ATLETICO CIMITERO"."""
    ta, tb = _tokens(a), _tokens(b)
    if not ta or len(ta) != len(tb):
        return False
    return all(x == y or x.startswith(y) or y.startswith(x) for x, y in zip(ta, tb))


def match_team(title, participants):
    """La squadra della lega a cui appartiene la scheda, o None."""
    name, _year = _split_founded(title)
    keys = {_norm(title), _norm(name)}
    for p in participants:
        if _norm(p.display_name) in keys or (p.short_name and _norm(p.short_name) in keys):
            return p
    for p in participants:
        if _split_founded(p.display_name)[0] and _norm(_split_founded(p.display_name)[0]) in keys:
            return p
    hits = [p for p in participants
            if _abbr_compatible(name, p.display_name) or (p.short_name and _abbr_compatible(name, p.short_name))]
    return hits[0] if len(hits) == 1 else None


def _names_compatible(a, b):
    la, sa = _name_parts(a)
    lb, sb = _name_parts(b)
    if _norm(a) == _norm(b):
        return True
    return bool(la & lb) and _shorts_compatible(sa, sb)


def _match_player(row, team, pool, claimed):
    name = row["name"]
    nname = _norm(name)
    free = [p for p in pool if p.pk not in claimed]
    if row.get("ext_id"):
        for p in free:
            if p.ext_id == row["ext_id"] and _names_compatible(name, p.name):
                return p
    if team is not None and team.pk:
        own = [p for p in free if p.owner_id == team.pk]
        for p in own:
            if _norm(p.name) == nname:
                return p
        for p in own:
            if _names_compatible(name, p.name) and (
                    not row.get("club") or _team_code(p.team) == _team_code(row["club"])
                    or _team_code(p.left_club) == _team_code(row["club"])):
                return p
    match = _find_match({"name": name, "role": row.get("role", ""), "team": row.get("club", "")}, free, claimed)
    if match is not None or row.get("status") == "out":
        # Fuori dalla Serie A il club della scheda non è sul listone: niente
        # somiglianze larghe, meglio un giocatore nuovo che quello sbagliato.
        return match
    code = _team_code(row.get("club"))
    best, best_ratio = None, 0.0
    for p in free:
        if code and _team_code(p.team) != code:
            continue
        ratio = SequenceMatcher(None, nname, _norm(p.name)).ratio()
        if ratio > best_ratio:
            best, best_ratio = p, ratio
    return best if best_ratio >= 0.85 else None


def _release(player, note):
    from ..models import RosterLog
    from ..services.loans import return_loan

    owner = player.owner
    if player.loan_from_id:
        return_loan(player, note=note)
        return
    RosterLog.objects.create(
        participant=owner, participant_name=owner.display_name if owner else "",
        player_name=player.name, player_role=player.role,
        action=RosterLog.Action.ADMIN_RELEASE, by_admin=True, note=note,
    )
    player.owner = None
    player.cost = Decimal("0")
    player.contract_years = None
    player.renewal_declared = None
    player.abroad_list = False
    player.abroad_compensation = None
    player.left_serie_a_at = None
    player.left_club, player.left_rank_kind, player.left_rank_pos = "", "", None
    player.loan_sessions_left = None
    player.save()


def _team_by_label(label, participants):
    if not label:
        return None
    return match_team(label, participants)


def apply_team_sheets(sheets, league, *, mapping=None, replace=True, dry_run=False,
                      overwrite_images=False, today=None):
    """Scrive le schede sulla lega. Con ``dry_run`` fa tutto e poi annulla.

    ``mapping`` associa l'indice di una scheda all'id della squadra ("new" per
    crearne una nuova, "skip" per saltarla); senza, la squadra si riconosce
    dal titolo della scheda.
    """
    from .. import team_sheets as ts
    from ..models import Participant, Player
    from ..services.sala import ensure_unlocked
    ensure_unlocked(league)

    mapping = mapping or {}
    now = timezone.now()
    season_start = ts.season_start(today)
    report = {"sheets": [], "dry_run": dry_run}
    with transaction.atomic():
        participants = list(Participant.objects.filter(league=league))
        pool = list(Player.objects.filter(league=league).select_related("owner", "loan_from"))
        claimed = set()
        for index, sheet in enumerate(sheets):
            choice = str(mapping.get(str(index), mapping.get(index, "")) or "")
            entry = _apply_one(sheet, index, choice, league, participants, pool, claimed,
                               replace=replace, dry_run=dry_run, overwrite_images=overwrite_images,
                               now=now, season_start=season_start)
            report["sheets"].append(entry)
        if dry_run:
            transaction.set_rollback(True)
    report["teams"] = [{"id": p.pk, "name": p.display_name}
                       for p in Participant.objects.filter(league=league).order_by("display_name")]
    return report


def _apply_one(sheet, index, choice, league, participants, pool, claimed, *, replace, dry_run,
               overwrite_images, now, season_start):
    from ..models import Participant, Player

    title = sheet.get("title") or ""
    entry = {"index": index, "source": sheet.get("source", ""), "title": title,
             "players": 0, "created_players": [], "moved": [], "released": [], "flagged_out": [],
             "reserved": 0, "loans": 0, "warnings": [], "images": [], "skipped": False,
             "created_team": False, "suggested_id": None}
    suggested = match_team(title, participants)
    entry["suggested_id"] = suggested.pk if suggested else None
    if choice == "skip":
        entry["skipped"] = True
        return entry
    team = None
    if choice.isdigit():
        team = next((p for p in participants if p.pk == int(choice)), None)
    elif choice != "new":
        team = suggested
    if team is None:
        base, year = _split_founded(title)
        team = Participant.objects.create(
            league=league, display_name=(_nice_name(base) or f"Squadra {index + 1}")[:80],
            founded=year, access_code=generate_access_code(), credits=league.budget, is_active=True,
        )
        participants.append(team)
        entry["created_team"] = True
    entry["team"] = team.display_name
    entry["team_id"] = None if entry["created_team"] else team.pk

    # Intestazione.
    fields = []
    for attr, key in (("president_name", "president"), ("coach_name", "coach"), ("stadium", "stadium")):
        value = _text(sheet.get(key))[:80]
        if value and value != getattr(team, attr):
            setattr(team, attr, value)
            fields.append(attr)
    if sheet.get("capacity") and sheet["capacity"] != team.stadium_capacity:
        team.stadium_capacity = sheet["capacity"]
        fields.append("stadium_capacity")
    if sheet.get("honours") and sheet["honours"] != team.honours:
        team.honours = sheet["honours"]
        fields.append("honours")
    _base, year = _split_founded(title)
    if year and not team.founded:
        team.founded = year
        fields.append("founded")
    if fields:
        team.save(update_fields=fields)
    entry["info"] = {"president": team.president_name, "coach": team.coach_name,
                     "stadium": team.stadium, "capacity": team.stadium_capacity,
                     "honours": team.honours}

    touched = set()
    for row in sheet.get("players", []):
        player, created = _place(row, team, league, pool, claimed, entry, now=now)
        touched.add(player.pk)
        entry["players"] += 1

    for item in sheet.get("reserved", []):
        row = {"name": item.get("name", ""), "club": item.get("club", ""), "status": "out", "role": ""}
        player = _match_player(row, team, pool, claimed)
        cost = parse_cost(item.get("cost"))
        years = parse_expiry_years(item.get("expiry"), season_start)
        if player is None:
            player = Player.objects.create(
                league=league, name=_nice_name(row["name"])[:120], role="A", team=_nice_name(row["club"])[:80],
                initial_price=max(cost or Decimal("1"), Decimal("1")))
            pool.append(player)
            entry["warnings"].append(f"{row['name']} (lista ceduti) non era nel listone: creato come attaccante.")
        _take(player, team, entry)
        claimed.add(player.pk)
        touched.add(player.pk)
        player.owner = team
        if cost is not None:
            player.cost = cost
        player.contract_years = years
        player.abroad_list = True
        player.left_club = _nice_name(row["club"])[:120]
        player.left_serie_a_at = None
        player.save()
        entry["reserved"] += 1

    for item in sheet.get("loans", []):
        other = _team_by_label(item.get("team"), participants)
        sessions = parse_sessions(item.get("expiry"))
        if item.get("in"):
            player = next((p for p in pool if p.pk in touched and _names_compatible(item["in"], p.name)), None)
            if player is None or other is None:
                entry["warnings"].append(f"Prestito di {item['in']} non importato: giocatore o società non trovati.")
                continue
            player.loan_from = other
            player.loan_sessions_left = sessions
            player.save(update_fields=["loan_from", "loan_sessions_left"])
            entry["loans"] += 1
        if item.get("out"):
            player = next((p for p in pool if _norm(p.name) == _norm(item["out"])), None)
            if player is None or other is None:
                entry["warnings"].append(f"Prestito di {item['out']} non importato: giocatore o società non trovati.")
                continue
            touched.discard(player.pk)
            player.owner = other
            player.loan_from = team
            player.loan_sessions_left = sessions
            player.save(update_fields=["owner", "loan_from", "loan_sessions_left"])
            claimed.add(player.pk)
            entry["loans"] += 1

    if replace:
        for player in [p for p in pool if p.owner_id == team.pk and p.pk not in touched]:
            entry["released"].append(player.name)
            _release(player, "Non presente nella scheda squadra importata")

    if sheet.get("substitutes"):
        entry["warnings"].append("Colonna SOSTITUTO non importata: " + ", ".join(sheet["substitutes"][:6]))
    if not dry_run:
        _save_images(team, sheet.get("images") or {}, overwrite_images, entry)
    else:
        entry["images"] = [k for k in (sheet.get("images") or {})
                           if overwrite_images or not getattr(team, k)]
    return entry


def _take(player, team, entry):
    if player.owner_id and player.owner_id != team.pk:
        entry["moved"].append(f"{player.name} (da {player.owner.display_name})")


def _place(row, team, league, pool, claimed, entry, *, now):
    """Assegna la riga della rosa a un giocatore (del listone o nuovo)."""
    from ..models import Player

    player = _match_player(row, team, pool, claimed)
    created = player is None
    club = _nice_name(row.get("club"))
    out = row.get("status") == "out"
    cost = row.get("cost")
    if created:
        player = Player.objects.create(
            league=league, name=_nice_name(row["name"])[:120], role=row.get("role") or "A",
            team=club[:80], initial_price=max(cost or Decimal("1"), Decimal("1")),
        )
        pool.append(player)
        entry["created_players"].append(player.name)
    else:
        _take(player, team, entry)
    claimed.add(player.pk)
    if cost is None and row.get("cost_raw"):
        entry["warnings"].append(f"{row['name']}: spesa «{row['cost_raw']}» non numerica, "
                                 + ("tenuta quella registrata." if player.owner_id == team.pk else "messa a 0."))
    if cost is not None:
        player.cost = cost
    elif player.owner_id != team.pk:
        player.cost = Decimal("0")
    player.owner = team
    if row.get("years_raw") is not None or row.get("years") is not None:
        player.contract_years = row.get("years")
    player.abroad_list = False
    if out and player.left_serie_a_at is None:
        player.left_serie_a_at = now
        if club.upper().startswith("SVINCOL"):
            player.left_rank_kind = "free"
        elif _team_code(club) != _team_code(player.team) or created:
            player.left_club = club[:120]
        entry["flagged_out"].append(player.name)
    player.save()
    return player, created


def _save_images(team, images, overwrite, entry):
    changed = []
    for key in ("logo", "kit_home", "kit_away"):
        data = images.get(key)
        if not data or (getattr(team, key) and not overwrite):
            continue
        try:
            image = clean_image_bytes(data)
        except UploadRejected:
            continue
        getattr(team, key).save(image.name, image, save=False)
        changed.append(key)
    if changed:
        team.save(update_fields=changed)
    entry["images"] = changed
